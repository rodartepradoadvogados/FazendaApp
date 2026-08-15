"""
Biblioteca de alimentos (Formulação de Dietas > aba "Biblioteca de
referência") — biblioteca mestre CowData + cópia por fazenda (copy-on-write),
CRUD e importação/exportação de planilha em lote.

Reúne num só lugar a regra "editar ou excluir um item mestre NUNCA grava na
linha mestre, sempre cria ou reaproveita a cópia da fazenda" — usada tanto
pela tela (criar/editar/excluir um alimento por vez) quanto pela importação
em lote, que precisa da mesma regra ao reimportar uma planilha com um nome
que já bate com um item mestre (vira cópia editada, não um alimento duplicado
com o mesmo nome). Ver docstring de `AlimentoNutricional` em
fazenda/models/formulacao.py para o desenho completo do dado.
"""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from datetime import datetime

from fastapi import HTTPException
from openpyxl import load_workbook
from sqlmodel import Session, select

from fazenda.models import AlimentoNutricional
from fazenda.models.formulacao import CAMPOS_CNCPS_CARBOIDRATO, CAMPOS_CNCPS_FRACIONAMENTO, CAMPOS_CNCPS_PROTEINA
from fazenda.rules.busca import normalizar_busca
from fazenda.rules.nutricao.biblioteca import biblioteca_semente
from fazenda.rules.nutricao.tipos import CATEGORIAS_NASEM
from fazenda.rules.planilha_modelo import gerar_modelo_xlsx

MAXIMO_LINHAS_IMPORTACAO = 500

# ±1 ponto percentual — folga de arredondamento de laudo de laboratório (a
# soma de frações medidas separadamente raramente fecha em 100,00% exato,
# mesmo quando o laudo está correto). Acima disso, o fechamento vira aviso.
TOLERANCIA_FECHAMENTO_FRACOES_PCT = 1.0


# ---------------------------------------------------------------------------
# Seed da biblioteca mestre — idempotente, mesmo padrão preguiçoso de
# `_seed_tabela_nutricional` em alimentacao.py (checa se já existe antes de
# inserir), pra rodar em todo GET sem duplicar e sem exigir uma migração de
# DADO à parte (os testes sobem o schema com create_all puro, sem alembic).
# ---------------------------------------------------------------------------
def semear_biblioteca_mestre(session: Session) -> None:
    ja_semeada = session.exec(
        select(AlimentoNutricional).where(AlimentoNutricional.fazenda_id == None)  # noqa: E711
    ).first()
    if ja_semeada:
        return
    for bruto in biblioteca_semente():
        campos = {c: bruto.get(c) for c in _CAMPOS_SEMENTE_MODELO}
        session.add(AlimentoNutricional(
            fazenda_id=None, origem_mestre_id=None, nome=bruto["nome"], categoria_nasem=bruto["categoria_nasem"],
            conc_pct=bruto["conc_pct"], fonte=bruto.get("fonte"), ativo=True,
            inclusao_min_pct=bruto.get("inclusao_min_pct"), inclusao_max_pct=bruto.get("inclusao_max_pct"),
            **campos,
        ))
    session.commit()


# Os 10 campos bromatológicos que `biblioteca_semente()` de fato preenche
# (ver fazenda.rules.nutricao.biblioteca._CAMPOS_SEMENTE) — os demais ~26
# campos de AlimentoNutricional ficam None nas linhas mestre e são
# completados pelo template da categoria no cálculo, igual a qualquer outro
# ingrediente sem laudo completo.
_CAMPOS_SEMENTE_MODELO: tuple[str, ...] = (
    "ms_pct", "pb_pct", "fdn_pct", "fda_pct", "lignina_pct", "amido_pct", "ee_pct", "cinzas_pct", "ca_pct", "p_pct",
)


# ---------------------------------------------------------------------------
# Listagem mesclada (mestre não sobrescrita/oculta + linhas da fazenda)
# ---------------------------------------------------------------------------
def listar_biblioteca(session: Session, fazenda_id: int | None) -> list[AlimentoNutricional]:
    """Biblioteca EFETIVA de uma fazenda: os itens mestre que ela não
    sobrescreveu nem ocultou + os itens dela mesma (cópia editada, item
    próprio, ou tombstone de item oculto — tombstone fica de fora do
    resultado, é só o que faz o item mestre correspondente sumir da lista).

    Filtro OBRIGATORIAMENTE tolerante a fazenda_id nulo (ver docstring de
    `AlimentoNutricional`) — usar `== fazenda_id` seco faria a biblioteca
    mestre inteira sumir da tela de toda fazenda."""
    semear_biblioteca_mestre(session)
    if fazenda_id is None:
        # Sem fazenda selecionada (ex.: acesso administrativo amplo) — mesma
        # convenção frouxa do resto do router: sem filtro, mostra tudo.
        return list(session.exec(select(AlimentoNutricional)).all())

    query = select(AlimentoNutricional).where(
        (AlimentoNutricional.fazenda_id == fazenda_id) | (AlimentoNutricional.fazenda_id == None)  # noqa: E711
    )
    linhas = session.exec(query).all()
    da_fazenda = [l for l in linhas if l.fazenda_id == fazenda_id]
    mestres = [l for l in linhas if l.fazenda_id is None]
    sobrescritas = {l.origem_mestre_id for l in da_fazenda if l.origem_mestre_id is not None}
    return [l for l in da_fazenda if l.ativo] + [m for m in mestres if m.id not in sobrescritas]


# ---------------------------------------------------------------------------
# Copy-on-write — editar/excluir um item mestre nunca grava na linha mestre
# ---------------------------------------------------------------------------
def _clonar_para_fazenda(mestre: AlimentoNutricional, fazenda_id: int, *, ativo: bool) -> AlimentoNutricional:
    from fazenda.rules.nutricao.tipos import CAMPOS_NUTRICIONAIS
    campos = {c: getattr(mestre, c) for c in CAMPOS_NUTRICIONAIS}
    # Frações CNCPS (fora de CAMPOS_NUTRICIONAIS de propósito — ver docstring
    # de CAMPOS_CNCPS_FRACIONAMENTO) também precisam ser copiadas: sem isso,
    # editar um item mestre que já tivesse fracionamento cadastrado apagaria
    # silenciosamente esse dado na cópia da fazenda.
    campos_cncps = {c: getattr(mestre, c) for c in CAMPOS_CNCPS_FRACIONAMENTO}
    return AlimentoNutricional(
        fazenda_id=fazenda_id, origem_mestre_id=mestre.id, alimento_id=None, nome=mestre.nome,
        categoria_nasem=mestre.categoria_nasem, conc_pct=mestre.conc_pct, fonte=mestre.fonte,
        observacao=mestre.observacao, ativo=ativo, extras_json=mestre.extras_json,
        inclusao_min_pct=mestre.inclusao_min_pct, inclusao_max_pct=mestre.inclusao_max_pct,
        # campos_editados_json NÃO é copiado: a mestre nunca tem nada
        # marcado como editado (é sempre 100% CowData, por definição) — a
        # cópia nasce sem marcação, e o router grava a marcação de verdade
        # (o que o USUÁRIO editou nesta operação) logo em seguida.
        **campos, **campos_cncps,
    )


def obter_para_editar(session: Session, fazenda_id: int | None, item_id: int) -> AlimentoNutricional:
    """Devolve a linha que o PUT deve gravar: a própria (se já é da fazenda)
    ou uma cópia nova/reaproveitada (se `item_id` aponta pra biblioteca
    mestre) — quem chama só precisa aplicar os campos editados e commitar,
    nunca decide sozinho se está mexendo na mestre ou numa cópia."""
    item = session.get(AlimentoNutricional, item_id)
    if not item:
        raise HTTPException(status_code=404, detail="Alimento não encontrado na biblioteca")
    if fazenda_id is None:
        return item
    if item.fazenda_id == fazenda_id:
        item.ativo = True  # reativa um tombstone que o usuário decidiu editar
        return item
    if item.fazenda_id is not None:
        raise HTTPException(status_code=404, detail="Alimento não encontrado na biblioteca")

    existente = session.exec(
        select(AlimentoNutricional).where(
            AlimentoNutricional.fazenda_id == fazenda_id, AlimentoNutricional.origem_mestre_id == item.id,
        )
    ).first()
    if existente:
        existente.ativo = True
        return existente
    return _clonar_para_fazenda(item, fazenda_id, ativo=True)


def excluir_ou_restaurar(session: Session, fazenda_id: int | None, item_id: int) -> dict:
    """Um único botão "Excluir" na tela cobre os três casos do requisito de
    CRUD: item próprio some de vez; cópia editada de um item mestre "volta ao
    padrão CowData" (a linha mestre nunca foi tocada, só reaparece); item
    mestre nunca editado é ocultado por um tombstone só desta fazenda."""
    item = session.get(AlimentoNutricional, item_id)
    if not item:
        raise HTTPException(status_code=404, detail="Alimento não encontrado na biblioteca")
    if fazenda_id is not None and item.fazenda_id is not None and item.fazenda_id != fazenda_id:
        raise HTTPException(status_code=404, detail="Alimento não encontrado na biblioteca")

    if item.fazenda_id is not None:
        if item.origem_mestre_id is not None:
            session.delete(item)
            session.commit()
            return {"acao": "restaurado", "mensagem": "Alimento restaurado ao padrão CowData."}
        session.delete(item)
        session.commit()
        return {"acao": "excluido", "mensagem": "Alimento excluído da sua biblioteca."}

    # item é a própria linha mestre — nunca apaga/edita ela; oculta só para
    # esta fazenda via tombstone (as outras fazendas continuam vendo-a).
    if fazenda_id is None:
        raise HTTPException(status_code=400, detail="Selecione a fazenda antes de remover um item da biblioteca padrão")
    existente = session.exec(
        select(AlimentoNutricional).where(
            AlimentoNutricional.fazenda_id == fazenda_id, AlimentoNutricional.origem_mestre_id == item.id,
        )
    ).first()
    if existente:
        if existente.ativo:
            session.delete(existente)
            session.commit()
            return {"acao": "restaurado", "mensagem": "Alimento restaurado ao padrão CowData."}
        return {"acao": "ja_oculto", "mensagem": "Este item já estava oculto na sua biblioteca."}
    oculto = _clonar_para_fazenda(item, fazenda_id, ativo=False)
    session.add(oculto)
    session.commit()
    return {"acao": "oculto", "mensagem": "Item padrão CowData removido da sua biblioteca (as outras fazendas continuam vendo-o)."}


# ---------------------------------------------------------------------------
# Checagem de fechamento das frações CNCPS (não bloqueante — ver docstring)
# ---------------------------------------------------------------------------
def _soma_se_alguma_preenchida(item: AlimentoNutricional, campos: tuple[str, ...]) -> float | None:
    """None quando NENHUM campo do grupo foi preenchido (nada a conferir —
    o usuário simplesmente ainda não começou a fracionar este alimento, o
    que não é um erro). Caso contrário, soma tratando os campos vazios como
    zero — é exatamente essa soma parcial que avisa "faltam X pontos" quando
    o preenchimento é incompleto."""
    valores = [getattr(item, c) for c in campos]
    if all(v is None for v in valores):
        return None
    return sum(v or 0.0 for v in valores)


def avisos_fechamento_fracoes(item: AlimentoNutricional) -> list[str]:
    """Avisos (nunca erros que bloqueiam) de que as frações CNCPS não fecham
    100% dentro de `TOLERANCIA_FECHAMENTO_FRACOES_PCT` — chamado tanto ao
    salvar um alimento pela tela quanto linha a linha na importação de
    planilha. Deliberadamente NÃO É validação bloqueante: a maioria dos
    alimentos vai chegar com fracionamento parcial (só o que o laudo do
    nutricionista trouxe), e travar o cadastro por causa disso empurraria a
    pessoa a inventar um número só para fechar a conta — pior que deixar o
    campo vazio. Cada mensagem já diz QUANTO falta (déficit) ou sobra
    (excesso), não só "não fecha".

    Proteína: PA1+PA2+PB1+PB2+PC precisa somar 100% da PB — os campos já são
    definidos em % da PB (ver docstring de AlimentoNutricional), então o
    alvo é sempre 100, sem cálculo extra.

    Carboidrato: os campos são em % da MATÉRIA SECA (não % do carboidrato),
    então o alvo não é 100 — é o carboidrato TOTAL do alimento, que se
    calcula por diferença (100 − PB − EE − Cinzas, todos em % da MS; mesma
    fórmula do CNCPS para "CHO"). Sem PB/EE/Cinzas preenchidos não dá pra
    calcular esse alvo — o aviso então explica a causa em vez de arriscar
    comparar com o número errado."""
    avisos: list[str] = []

    soma_proteina = _soma_se_alguma_preenchida(item, CAMPOS_CNCPS_PROTEINA)
    if soma_proteina is not None:
        diferenca = 100.0 - soma_proteina
        if abs(diferenca) > TOLERANCIA_FECHAMENTO_FRACOES_PCT:
            if diferenca > 0:
                avisos.append(
                    f"Frações de proteína (PA1+PA2+PB1+PB2+PC) somam {soma_proteina:.1f}% da PB — "
                    f"faltam {diferenca:.1f} pontos percentuais para fechar 100%."
                )
            else:
                avisos.append(
                    f"Frações de proteína (PA1+PA2+PB1+PB2+PC) somam {soma_proteina:.1f}% da PB — "
                    f"{-diferenca:.1f} pontos percentuais além de 100%."
                )

    soma_carboidrato = _soma_se_alguma_preenchida(item, CAMPOS_CNCPS_CARBOIDRATO)
    if soma_carboidrato is not None:
        if item.pb_pct is None or item.ee_pct is None or item.cinzas_pct is None:
            avisos.append(
                "Frações de carboidrato preenchidas, mas PB, EE e Cinzas precisam estar todos informados "
                "para calcular quanto de carboidrato este alimento tem e conferir se as frações fecham 100% dele."
            )
        else:
            cho_total = 100.0 - item.pb_pct - item.ee_pct - item.cinzas_pct
            diferenca = cho_total - soma_carboidrato
            if abs(diferenca) > TOLERANCIA_FECHAMENTO_FRACOES_PCT:
                base = (
                    f"Frações de carboidrato (CA1+CA2+CA3+CA4+CB1+CB2+CB3+CC) somam {soma_carboidrato:.1f}% da MS — "
                    f"o carboidrato deste alimento é {cho_total:.1f}% da MS (100% − PB − EE − Cinzas)"
                )
                if diferenca > 0:
                    avisos.append(f"{base}; faltam {diferenca:.1f} pontos percentuais para fechar.")
                else:
                    avisos.append(f"{base}; {-diferenca:.1f} pontos percentuais além disso.")

    return avisos


# ---------------------------------------------------------------------------
# Planilha-modelo (.xlsx) e importação em lote (.xlsx/.csv)
# ---------------------------------------------------------------------------
@dataclass(frozen=True)
class _Coluna:
    campo: str
    rotulo: str
    aliases: tuple[str, ...] = ()
    minimo: float | None = None
    maximo: float | None = None


# Só o subconjunto de campos que um laudo/planilha de nutricionista costuma
# trazer, MAIS o fracionamento CNCPS de carboidrato/proteína (frações e kd,
# ver CAMPOS_CNCPS_FRACIONAMENTO/AlimentoNutricional) — o resto
# (digestibilidades de referência, coeficientes de absorção de mineral,
# "Fase 2/3") não entra na planilha e continua sendo completado pelo
# template da categoria no cálculo, exatamente como um alimento cadastrado
# manualmente com campo em branco.
COLUNAS_IMPORTACAO: tuple[_Coluna, ...] = (
    _Coluna("nome", "Nome do alimento", ("alimento", "ingrediente", "nome do ingrediente")),
    _Coluna("categoria_nasem", "Categoria NASEM", ("categoria", "grupo", "categoria nutricional")),
    _Coluna("conc_pct", "Concentrado (% da MS do próprio alimento)", ("conc", "concentrado", "conc pct"), 0, 100),
    _Coluna("ms_pct", "MS - matéria seca (% da matéria NATURAL)", ("ms", "ms pct", "materia seca"), 0, 100),
    _Coluna("pb_pct", "PB - proteína bruta (% da MS)", ("pb", "pb pct", "proteina bruta"), 0, 300),
    _Coluna("fdn_pct", "FDN - fibra em detergente neutro (% da MS)", ("fdn", "fdn pct"), 0, 100),
    _Coluna("fda_pct", "FDA - fibra em detergente ácido (% da MS)", ("fda", "fda pct"), 0, 100),
    _Coluna("lignina_pct", "Lignina (% da MS)", ("lignina", "lignina pct"), 0, 100),
    _Coluna("amido_pct", "Amido (% da MS)", ("amido", "amido pct"), 0, 100),
    _Coluna("acucares_pct", "Açúcares (% da MS)", ("acucares", "acucar", "acucares pct"), 0, 100),
    _Coluna("ee_pct", "EE - extrato etéreo/gordura (% da MS)", ("ee", "ee pct", "gordura"), 0, 100),
    _Coluna("cinzas_pct", "Cinzas/matéria mineral (% da MS)", ("cinzas", "cinzas pct", "materia mineral", "mm"), 0, 100),
    # ---- Fracionamento CNCPS de carboidrato (% da MS) e proteína (% da PB)
    # — ver docstring completa de AlimentoNutricional/CAMPOS_CNCPS_*. Todas
    # opcionais; kd só existe para a fração que de fato tem taxa própria.
    _Coluna("cncps_ca1_pct", "CNCPS CA1 - ácidos orgânicos (% da MS)", ("ca1", "cncps ca1"), 0, 100),
    _Coluna("cncps_ca2_pct", "CNCPS CA2 - ácido lático (% da MS)", ("ca2", "cncps ca2"), 0, 100),
    _Coluna("cncps_kd_ca2_pct_h", "CNCPS kd CA2 (%/h)", ("kd ca2",), 0, 500),
    _Coluna("cncps_ca3_pct", "CNCPS CA3 - outros solúveis (% da MS)", ("ca3", "cncps ca3"), 0, 100),
    _Coluna("cncps_kd_ca3_pct_h", "CNCPS kd CA3 (%/h)", ("kd ca3",), 0, 500),
    _Coluna("cncps_ca4_pct", "CNCPS CA4 - açúcares (% da MS)", ("ca4", "cncps ca4"), 0, 100),
    _Coluna("cncps_kd_ca4_pct_h", "CNCPS kd CA4 (%/h)", ("kd ca4",), 0, 500),
    _Coluna("cncps_cb1_pct", "CNCPS CB1 - amido (% da MS)", ("cb1", "cncps cb1"), 0, 100),
    _Coluna("cncps_kd_cb1_pct_h", "CNCPS kd CB1 (%/h)", ("kd cb1",), 0, 500),
    _Coluna("cncps_cb2_pct", "CNCPS CB2 - fibra solúvel (% da MS)", ("cb2", "cncps cb2"), 0, 100),
    _Coluna("cncps_kd_cb2_pct_h", "CNCPS kd CB2 (%/h)", ("kd cb2",), 0, 500),
    _Coluna("cncps_cb3_pct", "CNCPS CB3 - FDN digestível (% da MS)", ("cb3", "cncps cb3"), 0, 100),
    _Coluna("cncps_kd_cb3_pct_h", "CNCPS kd CB3 (%/h)", ("kd cb3",), 0, 500),
    _Coluna("cncps_cc_pct", "CNCPS CC - FDN indigestível (% da MS)", ("cc", "cncps cc", "undf"), 0, 100),
    _Coluna("cncps_pa1_pct", "CNCPS PA1 - amônia (% da PB)", ("pa1", "cncps pa1"), 0, 100),
    _Coluna("cncps_pa2_pct", "CNCPS PA2 - peptídeos solúveis (% da PB)", ("pa2", "cncps pa2"), 0, 100),
    _Coluna("cncps_kd_pa2_pct_h", "CNCPS kd PA2 (%/h)", ("kd pa2",), 0, 500),
    _Coluna("cncps_pb1_pct", "CNCPS PB1 - proteína rapidamente degradável (% da PB)", ("pb1", "cncps pb1"), 0, 100),
    _Coluna("cncps_kd_pb1_pct_h", "CNCPS kd PB1 (%/h)", ("kd pb1",), 0, 500),
    _Coluna("cncps_pb2_pct", "CNCPS PB2 - proteína lentamente degradável (% da PB)", ("pb2", "cncps pb2"), 0, 100),
    _Coluna("cncps_kd_pb2_pct_h", "CNCPS kd PB2 (%/h)", ("kd pb2",), 0, 500),
    _Coluna("cncps_pc_pct", "CNCPS PC - proteína indisponível (% da PB)", ("pc", "cncps pc"), 0, 100),
    _Coluna("ca_pct", "Ca - cálcio (% da MS)", ("ca", "ca pct", "calcio"), 0, 50),
    _Coluna("p_pct", "P - fósforo (% da MS)", ("p", "p pct", "fosforo"), 0, 50),
    _Coluna("mg_pct", "Mg - magnésio (% da MS)", ("mg", "mg pct", "magnesio"), 0, 50),
    _Coluna("k_pct", "K - potássio (% da MS)", ("k", "k pct", "potassio"), 0, 50),
    _Coluna("na_pct", "Na - sódio (% da MS)", ("na", "na pct", "sodio"), 0, 50),
    _Coluna("cl_pct", "Cl - cloro (% da MS)", ("cl", "cl pct", "cloro"), 0, 50),
    _Coluna("s_pct", "S - enxofre (% da MS)", ("s", "s pct", "enxofre"), 0, 50),
    _Coluna("custo_kg_mn", "Custo (R$ por kg de matéria natural)", ("custo", "preco", "valor", "custo kg mn"), 0, 1000),
    _Coluna("inclusao_min_pct", "Inclusão mínima sugerida (% da MS da DIETA)", ("inclusao min", "inclusao minima"), 0, 100),
    _Coluna("inclusao_max_pct", "Inclusão máxima sugerida (% da MS da DIETA)", ("inclusao max", "inclusao maxima"), 0, 100),
    _Coluna("fonte", "Fonte/observação", ("fonte", "observacao", "obs")),
)

_CATEGORIAS_NORMALIZADAS: dict[str, str] = {normalizar_busca(c): c for c in CATEGORIAS_NASEM}
_ALIASES_CATEGORIA: dict[str, str] = {
    normalizar_busca(k): v for k, v in {
        "volumoso": "Forragem", "forragens": "Forragem",
        "concentrado energetico": "Concentrado energetico", "energetico": "Concentrado energetico",
        "concentrado proteico": "Concentrado proteico", "proteico": "Concentrado proteico",
        "proteina animal": "Proteina animal", "farinha de carne": "Proteina animal",
        "suplemento de gordura": "Suplemento de gordura", "gordura protegida": "Suplemento de gordura",
        "suplemento de acido graxo": "Suplemento de acido graxo", "acido graxo": "Suplemento de acido graxo",
        "acucar": "Acucar/alcool de acucar", "acucar/alcool de acucar": "Acucar/alcool de acucar",
        "mineral": "Vitaminico/mineral", "vitaminico": "Vitaminico/mineral", "nucleo": "Vitaminico/mineral",
        "leite": "Leite/sucedaneo", "sucedaneo": "Leite/sucedaneo",
        "outro": "Outros",
    }.items()
}


def _normalizar_cabecalho(s: str) -> str:
    return normalizar_busca(s)


def _resolver_categoria(texto: str) -> tuple[str, bool]:
    """Devolve (categoria_canônica, reconhecida) — cai em "Outros" (com
    aviso) quando o texto da planilha não bate com nenhum valor nem apelido
    conhecido, pra nunca travar a linha inteira por causa da categoria."""
    norm = normalizar_busca(texto)
    if norm in _CATEGORIAS_NORMALIZADAS:
        return _CATEGORIAS_NORMALIZADAS[norm], True
    if norm in _ALIASES_CATEGORIA:
        return _ALIASES_CATEGORIA[norm], True
    return "Outros", False


def gerar_modelo_planilha() -> bytes:
    exemplo_fonte = biblioteca_semente()[0]  # "Silagem de milho" — exemplo realista
    exemplo = []
    for col in COLUNAS_IMPORTACAO:
        if col.campo == "fonte":
            exemplo.append("Laudo do laboratório X, agosto/2026")
        else:
            valor = exemplo_fonte.get(col.campo)
            exemplo.append(valor if valor is not None else "")
    return gerar_modelo_xlsx([c.rotulo for c in COLUNAS_IMPORTACAO], exemplo, aba="Biblioteca de alimentos")


def _parse_numero(valor: object) -> float | None:
    if valor is None:
        return None
    if isinstance(valor, (int, float)):
        return float(valor)
    texto = str(valor).strip()
    if not texto:
        return None
    # "33,5" (decimal brasileiro sem separador de milhar) -> "33.5"; qualquer
    # outra coisa ("33.5", "33", "-2") tenta direto — cobre tanto exportação
    # brasileira quanto americana sem adivinhar separador de milhar.
    if re.fullmatch(r"-?\d+,\d+", texto):
        texto = texto.replace(",", ".")
    try:
        return float(texto)
    except ValueError:
        return None


def _linhas_da_planilha(nome_arquivo: str, conteudo: bytes) -> list[dict[str, str]]:
    """Lê .xlsx/.xlsm (openpyxl) ou .csv (delimitador `,` ou `;`, UTF-8 ou
    Latin-1 — as duas codificações que um Excel brasileiro salva, dependendo
    da versão) e devolve uma lista de {cabeçalho da coluna: valor}."""
    nome = (nome_arquivo or "").lower()
    if nome.endswith((".xlsx", ".xlsm")):
        try:
            wb = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
        except Exception:
            raise HTTPException(status_code=400, detail="Não foi possível ler o arquivo — envie um .xlsx válido")
        try:
            ws = wb.active
            linhas_brutas = list(ws.iter_rows(values_only=True))
        finally:
            wb.close()
        if not linhas_brutas:
            return []
        cabecalho = [str(c).strip() if c is not None else "" for c in linhas_brutas[0]]
        linhas = []
        for valores in linhas_brutas[1:]:
            if valores is None or all(v is None or str(v).strip() == "" for v in valores):
                continue
            linhas.append({cabecalho[i]: valores[i] for i in range(len(cabecalho)) if i < len(valores) and cabecalho[i]})
        return linhas

    try:
        texto = conteudo.decode("utf-8-sig")
    except UnicodeDecodeError:
        texto = conteudo.decode("latin-1", errors="replace")
    primeira_linha = texto.splitlines()[0] if texto else ""
    try:
        delimitador = csv.Sniffer().sniff(primeira_linha, delimiters=",;\t").delimiter
    except csv.Error:
        delimitador = ";" if primeira_linha.count(";") > primeira_linha.count(",") else ","
    leitor = csv.DictReader(io.StringIO(texto), delimiter=delimitador)
    return [
        {(k or "").strip(): v for k, v in row.items() if k}
        for row in leitor
        if any((v or "").strip() for v in row.values())
    ]


def importar_planilha(session: Session, fazenda_id: int | None, nome_arquivo: str, conteudo: bytes) -> dict:
    if fazenda_id is None:
        raise HTTPException(status_code=400, detail="Selecione a fazenda antes de importar a biblioteca")
    semear_biblioteca_mestre(session)  # garante a mestre semeada mesmo se ninguém abriu a biblioteca ainda

    linhas = _linhas_da_planilha(nome_arquivo, conteudo)
    if not linhas:
        raise HTTPException(status_code=400, detail="Planilha vazia ou sem linhas de dados")
    if len(linhas) > MAXIMO_LINHAS_IMPORTACAO:
        raise HTTPException(status_code=422, detail=f"No máximo {MAXIMO_LINHAS_IMPORTACAO} linhas por importação (recebidas: {len(linhas)}).")

    # Cabeçalho -> campo canônico, tolerante a acento/caixa/espaço (mesma
    # normalização usada em toda busca do site — ver fazenda.rules.busca).
    campo_por_coluna: dict[str, str] = {}
    for coluna in linhas[0].keys():
        norm = _normalizar_cabecalho(coluna)
        for col in COLUNAS_IMPORTACAO:
            candidatos = {_normalizar_cabecalho(col.rotulo), _normalizar_cabecalho(col.campo)}
            candidatos.update(_normalizar_cabecalho(a) for a in col.aliases)
            if norm in candidatos:
                campo_por_coluna[coluna] = col.campo
                break
    if "nome" not in campo_por_coluna.values():
        raise HTTPException(
            status_code=400,
            detail='A planilha precisa de uma coluna com o nome do alimento (ex.: "Nome do alimento") — baixe o modelo do sistema se não tiver certeza dos cabeçalhos aceitos.',
        )

    # Todas as linhas JÁ da fazenda por nome — inclusive tombstones (ativo=
    # False) — pra reimportar o nome de um item que a fazenda tinha ocultado
    # reaproveitar/reativar a mesma linha em vez de criar uma duplicata (ver
    # `listar_biblioteca`, que some com os tombstones do resultado).
    da_fazenda_por_nome = {
        normalizar_busca(l.nome): l
        for l in session.exec(select(AlimentoNutricional).where(AlimentoNutricional.fazenda_id == fazenda_id)).all()
    }
    mestres_por_nome = {
        normalizar_busca(m.nome): m
        for m in session.exec(select(AlimentoNutricional).where(AlimentoNutricional.fazenda_id == None)).all()  # noqa: E711
    }

    criados = atualizados = 0
    avisos: list[str] = []
    erros: list[str] = []

    for i, linha_bruta in enumerate(linhas, start=2):  # linha 1 da planilha = cabeçalho
        dados: dict[str, object] = {}
        for coluna, valor in linha_bruta.items():
            campo = campo_por_coluna.get(coluna)
            if campo:
                dados[campo] = valor

        nome = str(dados.get("nome") or "").strip()
        if not nome:
            erros.append(f"Linha {i}: nome do alimento é obrigatório — linha ignorada.")
            continue

        categoria_bruta = str(dados.get("categoria_nasem") or "").strip()
        if categoria_bruta:
            categoria, reconhecida = _resolver_categoria(categoria_bruta)
            if not reconhecida:
                avisos.append(f'Linha {i} ("{nome}"): categoria "{categoria_bruta}" não reconhecida — usando "Outros".')
        else:
            categoria = "Outros"
            avisos.append(f'Linha {i} ("{nome}"): categoria NASEM não informada — usando "Outros".')

        valores: dict[str, float | None] = {}
        for col in COLUNAS_IMPORTACAO:
            if col.campo in ("nome", "categoria_nasem", "fonte") or col.campo not in dados:
                continue
            bruto = dados[col.campo]
            numero = _parse_numero(bruto)
            if bruto not in (None, "") and numero is None:
                avisos.append(f'Linha {i} ("{nome}"): valor "{bruto}" de "{col.rotulo}" não é um número — ignorado.')
                continue
            if numero is not None and col.minimo is not None and not (col.minimo <= numero <= col.maximo):
                avisos.append(
                    f'Linha {i} ("{nome}"): "{col.rotulo}"={numero} fora da faixa esperada '
                    f'({col.minimo}–{col.maximo}) — ignorado; a categoria "{categoria}" preenche um valor padrão no cálculo.'
                )
                continue
            if numero is not None:
                valores[col.campo] = numero

        fonte = str(dados.get("fonte") or "").strip() or None
        chave = normalizar_busca(nome)
        existente = da_fazenda_por_nome.get(chave)

        if existente is not None:
            item = existente
            atualizados += 1
        else:
            mestre = mestres_por_nome.get(chave)
            if mestre is not None:
                # Nome bate com um item mestre que esta fazenda ainda não
                # sobrescreveu — copy-on-write, não cria um alimento duplicado.
                item = _clonar_para_fazenda(mestre, fazenda_id, ativo=True)
                atualizados += 1
            else:
                item = AlimentoNutricional(fazenda_id=fazenda_id, nome=nome, categoria_nasem=categoria)
                criados += 1

        item.nome = nome
        item.categoria_nasem = categoria
        item.ativo = True  # reimportar reativa um item que estivesse oculto (tombstone)
        if "conc_pct" in valores:
            item.conc_pct = valores.pop("conc_pct") or 0.0
        if "inclusao_min_pct" in valores:
            item.inclusao_min_pct = valores.pop("inclusao_min_pct")
        if "inclusao_max_pct" in valores:
            item.inclusao_max_pct = valores.pop("inclusao_max_pct")
        if fonte:
            item.fonte = fonte
        for campo, valor in valores.items():
            setattr(item, campo, valor)
        item.atualizado_em = datetime.utcnow()
        # Fechamento das frações CNCPS — mesma checagem não bloqueante da
        # tela (ver avisos_fechamento_fracoes), aqui linha a linha: a
        # importação não pode travar por causa de um laudo parcial, então
        # isto vira aviso, nunca erro.
        for aviso_fechamento in avisos_fechamento_fracoes(item):
            avisos.append(f'Linha {i} ("{nome}"): {aviso_fechamento}')
        session.add(item)
        da_fazenda_por_nome[chave] = item

    session.commit()
    return {"criados": criados, "atualizados": atualizados, "avisos": avisos, "erros": erros}


__all__ = [
    "semear_biblioteca_mestre",
    "listar_biblioteca",
    "obter_para_editar",
    "excluir_ou_restaurar",
    "avisos_fechamento_fracoes",
    "TOLERANCIA_FECHAMENTO_FRACOES_PCT",
    "gerar_modelo_planilha",
    "importar_planilha",
    "COLUNAS_IMPORTACAO",
]
