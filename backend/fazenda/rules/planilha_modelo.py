"""Geração de planilhas-modelo (.xlsx) para download — cabeçalho em negrito
mais uma linha de exemplo, para o usuário só apagar o exemplo e preencher."""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def gerar_modelo_xlsx(colunas: list[str], exemplo: list[str] | None = None, aba: str = "Modelo") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = aba[:31]  # limite do Excel para nome de aba

    ws.append(colunas)
    for cel in ws[1]:
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor="4A6B3A")

    if exemplo:
        ws.append(exemplo)

    for idx, coluna in enumerate(colunas, start=1):
        largura = max(len(coluna), len(str(exemplo[idx - 1])) if exemplo and idx - 1 < len(exemplo) else 0)
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = max(12, min(40, largura + 4))

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
