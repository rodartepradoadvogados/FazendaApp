"""
Pacote do contador (Fase C5 — docs/financeiro-fechamento-conciliacao.md).

Um PDF e um ZIP com os arquivos de apoio (CSV e uma planilha Excel com
números reais) para o período escolhido, montados pelas MESMAS funções dos
relatórios — nenhum número é recalculado aqui:

- DRE de competência e de caixa: `calcular_dre` (a cascata única do servidor),
  em CSV pelo mesmo `_dre_para_csv` do e-mail do Portal;
- Livro caixa da atividade rural: o que passou pelo caixa no período (data de
  caixa e valor pago da regra do saldo, fazenda/rules/saldo_conta.py);
- Não classificados e pendências: o balde "não classificado" da DRE e o
  checklist do Fechamento do mês (fazenda/rules/fechamento_mes.py), por mês;
- Patrimônio e depreciação do período (rules/depreciacao_periodo.py);
- Conciliação bancária e fechamentos do período, quando existirem;
- Nota de método (flag `financeiro_regras_v2` ligada ou não) e um arquivo de
  APOIO no leiaute do LCDPR (registros Q100 e Q200 com o que o sistema sabe;
  o que falta fica em branco e listado no LEIA-ME — nada é inventado).

LGPD: só o necessário para a escrituração. Vai o nome da fazenda e o CPF/CNPJ
do cadastro dela (o contador precisa); dos participantes vai só o nome como
está no lançamento (sem CPF/CNPJ, que o sistema não guarda); da folha, só
totais. Nenhum dado de saúde, endereço ou contato.
"""
from __future__ import annotations

import csv
import html
import io
import zipfile
from datetime import date, datetime, timedelta

from sqlmodel import Session, select

from fazenda.models import ContaCorrente, ContaGerencial, Fazenda, LancamentoItem, Patrimonio
from fazenda.rules import saldo_conta
from fazenda.rules.datas import agora_local, hoje_local
from fazenda.rules.parametros import regras_v2_ativas

ROTULO_NATUREZA = {
    "OPERACIONAL": "Operacional", "INVESTIMENTO": "Investimento", "FINANCIAMENTO": "Financiamento",
    "CAPITAL": "Capital (aporte/retirada)", "TRANSFERENCIA": "Transferência", "ADIANTAMENTO": "Adiantamento",
    "OBRIGACAO": "Obrigação (retenção/guia)", "NAO_INFORMADA": "A classificar",
}
# Entra no resultado da atividade rural do livro caixa (custeio e investimento
# são dedutíveis; financiamento, capital, transferência, adiantamento e
# obrigação não são receita nem despesa da atividade).
ENTRA_NO_LIVRO = {"OPERACIONAL", "INVESTIMENTO"}
# Código do tipo de documento no leiaute do LCDPR (Q100.TIPO_DOC) — conferir no
# manual vigente da Receita antes de transmitir.
TIPO_DOC_LCDPR = {"nota fiscal": "1", "fatura": "2", "recibo": "3", "contrato": "4", "folha de pagamento": "5"}


def _pl(n: int, um: str, varios: str) -> str:
    return f"{n} {um if n == 1 else varios}"


def _f(v) -> float:
    return round(float(v or 0.0), 2)


def _br(v: float | None, casas: int = 2) -> str:
    if v is None:
        return "—"
    s = f"{abs(v):,.{casas}f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"-R$ {s}" if v < 0 else f"R$ {s}"


def _dmy(d: date | str | None) -> str:
    if not d:
        return "—"
    s = d.isoformat() if isinstance(d, date) else str(d)
    return f"{s[8:10]}/{s[5:7]}/{s[0:4]}"


def _meses(ini: date, fim: date) -> list[str]:
    saida, a, m = [], ini.year, ini.month
    while (a, m) <= (fim.year, fim.month):
        saida.append(f"{a:04d}-{m:02d}")
        a, m = (a + 1, 1) if m == 12 else (a, m + 1)
    return saida


# ── dados ─────────────────────────────────────────────────────────────────
def cabecalho(session: Session, fazenda_id: int, ini: date, fim: date, usuario_nome: str | None) -> dict:
    faz = session.get(Fazenda, fazenda_id)
    doc = None
    if faz and faz.documento:
        doc = f"{(faz.tipo_documento or 'documento').upper()} {faz.documento}"
    return {
        "fazenda": faz.nome if faz else "", "documento": doc,
        "periodo": {"inicio": ini.isoformat(), "fim": fim.isoformat()},
        "regime": "Competência (mês do gasto) e caixa (dia do pagamento)", "centro": "Todos os centros",
        "gerado_em": agora_local().strftime("%d/%m/%Y %H:%M"), "gerado_por": usuario_nome or "",
        "regras_v2": regras_v2_ativas(session, fazenda_id),
    }


def livro_caixa(session: Session, fazenda_id: int, ini: date, fim: date, regras_v2: bool) -> dict:
    """Lançamentos que passaram pelo caixa no período (data de caixa, valor pago)."""
    from fazenda.api.routers.financeiro import _contexto_natureza, saldos_contas_v2
    from fazenda.rules.dre import resolver_linha_dre
    from fazenda.rules.natureza import resolver_natureza

    contexto = _contexto_natureza(session, fazenda_id) if regras_v2 else None
    q = select(ContaGerencial).where(ContaGerencial.fazenda_id == fazenda_id, ContaGerencial.data_pagamento != None)  # noqa: E711
    pagos = []
    for c in session.exec(q).all():
        dc = saldo_conta.data_caixa(c)
        if dc is None or not (ini <= dc <= fim) or not saldo_conta.conta_entra_no_caixa(c):
            continue
        valor = saldo_conta.valor_pago_efetivo(c)
        if not valor:
            continue
        pagos.append((dc, c, valor))
    numeros = {c.numero_lancamento for _d, c, _v in pagos if c.numero_lancamento}
    codigo_item: dict[str, str] = {}
    if numeros:
        for it in session.exec(select(LancamentoItem).where(
                LancamentoItem.fazenda_id == fazenda_id, LancamentoItem.numero_lancamento.in_(list(numeros)))).all():
            if it.codigo_conta_gerencial and it.numero_lancamento not in codigo_item:
                codigo_item[it.numero_lancamento] = it.codigo_conta_gerencial
    contas = list(session.exec(select(ContaCorrente).where(ContaCorrente.fazenda_id == fazenda_id).order_by(ContaCorrente.id)).all())
    saldos_ini = saldos_contas_v2(session, contas, fazenda_id, ini - timedelta(days=1)) if contas else {}
    saldo_inicial = round(sum(s.saldo for s in saldos_ini.values()), 2)
    pendente_abertura = any(s.pendente_abertura for s in saldos_ini.values()) or not contas
    linhas, saldo = [], saldo_inicial
    for dc, c, valor in sorted(pagos, key=lambda x: (x[0], x[1].id)):
        codigo = c.codigo_conta or codigo_item.get(c.numero_lancamento or "")
        natureza = None
        if contexto is not None:
            natureza = resolver_natureza(
                natureza_conta=c.natureza_fin, tipo=c.tipo, patrimonio_id=c.patrimonio_id,
                patrimonios_vendidos=contexto.patrimonios_vendidos, codigo_conta=codigo,
                mapa_natureza_plano=contexto.mapa_natureza_plano,
                linha_dre=resolver_linha_dre(codigo, contexto.mapa_linha_dre),
            )
        entrada = valor if c.tipo == "receita" else 0.0
        saida = valor if c.tipo != "receita" else 0.0
        saldo = round(saldo + entrada - saida, 2)
        linhas.append({
            "data": dc.isoformat(), "lancamento": c.numero_lancamento or f"id {c.id}",
            "documento": c.numero_nota or "", "tipo_documento": c.tipo_documento or "",
            "participante": c.fornecedor_cliente or "", "historico": c.descricao or "",
            "conta_gerencial": codigo or "", "conta_corrente": c.conta_bancaria or "",
            "conta_corrente_id": c.conta_corrente_id,
            "natureza": ROTULO_NATUREZA.get(natureza or "", "") if natureza else "",
            "entra_no_resultado": ("sim" if natureza in ENTRA_NO_LIVRO else "a classificar" if natureza == "NAO_INFORMADA" else "não") if natureza else "",
            "tipo": c.tipo, "entrada": _f(entrada), "saida": _f(saida), "saldo": saldo,
        })
    meses = []
    acumulado = saldo_inicial
    for m in _meses(ini, fim):
        doms = [l for l in linhas if l["data"][:7] == m]
        e, s = _f(sum(l["entrada"] for l in doms)), _f(sum(l["saida"] for l in doms))
        acumulado = round(acumulado + e - s, 2)
        r_e = _f(sum(l["entrada"] for l in doms if l["entra_no_resultado"] == "sim"))
        r_s = _f(sum(l["saida"] for l in doms if l["entra_no_resultado"] == "sim"))
        meses.append({"mes": m, "entradas": e, "saidas": s, "saldo": acumulado, "lancamentos": len(doms),
                      "receitas_atividade": r_e if regras_v2 else None, "despesas_atividade": r_s if regras_v2 else None})
    tot_e, tot_s = _f(sum(l["entrada"] for l in linhas)), _f(sum(l["saida"] for l in linhas))
    return {
        "linhas": linhas, "meses": meses, "saldo_inicial": saldo_inicial, "saldo_final": round(saldo_inicial + tot_e - tot_s, 2),
        "saldo_inicial_pendente": pendente_abertura, "entradas": tot_e, "saidas": tot_s,
        "receitas_atividade": _f(sum(l["entrada"] for l in linhas if l["entra_no_resultado"] == "sim")) if regras_v2 else None,
        "despesas_atividade": _f(sum(l["saida"] for l in linhas if l["entra_no_resultado"] == "sim")) if regras_v2 else None,
    }


def patrimonio(session: Session, fazenda_id: int, ini: date, fim: date) -> dict:
    from fazenda.rules.depreciacao_periodo import calcular_depreciacao_periodo
    from fazenda.rules.patrimonio import calcular_depreciacao

    bens = list(session.exec(select(Patrimonio).where(Patrimonio.fazenda_id == fazenda_id).order_by(Patrimonio.id)).all())
    dumps = [b.model_dump() for b in bens]
    per = calcular_depreciacao_periodo(dumps, ini, fim)
    no_periodo = {i.get("patrimonio_id"): i.get("valor") for i in per.get("itens") or []}
    linhas = []
    for b, d in zip(bens, dumps):
        acum = calcular_depreciacao(d, fim)
        linhas.append({
            "codigo": b.codigo or "", "nome": b.nome, "tipo": b.tipo or "", "centro_custo": b.centro_custo or "",
            "data_imobilizacao": b.data_imobilizacao.isoformat() if b.data_imobilizacao else "",
            "valor_aquisicao": _f(b.valor_total), "metodo": b.metodo_depreciacao or "",
            "vida_util_anos": b.vida_util_anos, "depreciavel": bool(b.depreciavel),
            "depreciacao_periodo": _f(no_periodo.get(b.id)), "depreciacao_acumulada": acum.get("depreciacao_acumulada"),
            "valor_atual": acum.get("valor_atual"), "inconsistencia": acum.get("inconsistencia") or "",
            "data_baixa": b.data_baixa.isoformat() if b.data_baixa else "", "valor_baixa": b.valor_baixa,
            "adquirido_no_periodo": bool(b.data_imobilizacao and ini <= b.data_imobilizacao <= fim),
        })
    return {"bens": linhas, "depreciacao_periodo": _f(per.get("total")), "inconsistencias": per.get("inconsistencias") or []}


def pendencias_por_mes(session: Session, fazenda_id: int, ini: date, fim: date) -> list[dict]:
    from fazenda.rules.fechamento_mes import checklist

    saida = []
    hoje = hoje_local()
    for m in _meses(ini, fim):
        if m > f"{hoje.year:04d}-{hoje.month:02d}":
            break
        for c in checklist(session, fazenda_id, m):
            if c["ok"]:
                continue
            detalhe = "; ".join(_resumo_item(c["id"], i) for i in c["itens"][:12])
            if len(c["itens"]) > 12:
                detalhe += f"; … e mais {len(c['itens']) - 12}"
            saida.append({"mes": m, "conferencia": c["nome"], "quantidade": c["quantidade"], "bloqueia": c["bloqueia"], "detalhe": detalhe})
    return saida


def _resumo_item(id_: str, i: dict) -> str:
    if id_ == "sem_conta":
        return f"{i.get('codigo') or ''} {i.get('nome') or ''}: receita {_br(i.get('receita'))}, despesa {_br(i.get('despesa'))}".strip()
    if id_ == "natureza":
        return f"{i.get('numero_lancamento') or ''} {i.get('descricao') or ''} {_br(i.get('valor'))}".strip()
    if id_ == "contas_automaticas":
        return f"{i.get('rotulo')} ({_br(i.get('valor'))})"
    if id_ == "saldo_abertura":
        return str(i.get("conta"))
    if id_ == "cartao":
        return f"{i.get('cartao')} {i.get('competencia')}"
    if id_ == "folha":
        return f"folha {i.get('folha_id')} ({i.get('status')})"  # sem nome do funcionário (LGPD)
    if id_ == "conciliacao":
        return f"{i.get('conta')}: {'extrato não importado' if not i.get('importado') else str(i.get('pendentes')) + ' linha(s) pendente(s)'}"
    if id_ == "depreciacao":
        return f"{i.get('item')}: {i.get('motivo')}"
    return str(i)


def fechamentos(session: Session, fazenda_id: int, ini: date, fim: date) -> dict:
    from fazenda.rules.fechamento_mes import dump_evento, eventos

    meses = set(_meses(ini, fim))
    evs = [e for e in eventos(session, fazenda_id) if e.mes in meses]
    ultimo = {}
    for e in evs:
        ultimo[e.mes] = e
    return {
        "meses": [{"mes": m, "status": "fechado" if ultimo.get(m) and ultimo[m].acao == "fechar" else "aberto",
                   "retrato_sha256": ultimo[m].retrato_sha256 if ultimo.get(m) and ultimo[m].acao == "fechar" else None}
                  for m in sorted(meses)],
        "trilha": [dump_evento(e) for e in evs],
    }


def montar(session: Session, fazenda_id: int, ini: date, fim: date, usuario_nome: str | None) -> dict:
    from fazenda.api.routers.financeiro import calcular_dre
    from fazenda.rules import conciliacao_banco

    cab = cabecalho(session, fazenda_id, ini, fim, usuario_nome)
    v2 = cab["regras_v2"]
    dre_c = calcular_dre(session, fazenda_id, ini, fim, None, "competencia", regras_v2=v2)
    dre_k = calcular_dre(session, fazenda_id, ini, fim, None, "caixa", regras_v2=v2)
    dre_c.pop("_registros", None)
    dre_k.pop("_registros", None)
    livro = livro_caixa(session, fazenda_id, ini, fim, v2)
    pat = patrimonio(session, fazenda_id, ini, fim)
    pend = pendencias_por_mes(session, fazenda_id, ini, fim)
    conc = conciliacao_banco.resumo_contas(session, fazenda_id, ini, fim)
    fech = fechamentos(session, fazenda_id, ini, fim)
    nao = dre_c.get("nao_classificado") or {}
    dados = {
        "cabecalho": cab, "dre_competencia": dre_c, "dre_caixa": dre_k, "livro": livro, "patrimonio": pat,
        "nao_classificados": nao.get("contas") or [], "pendencias": pend, "conciliacao": conc, "fechamentos": fech,
    }
    dados["itens"] = itens_do_pacote(dados)
    dados["metodo"] = nota_de_metodo(dados)
    return dados


def _linha(dre: dict, chave: str) -> float:
    return next((l["valor"] for l in dre.get("cascata") or [] if l["chave"] == chave), 0.0)


def itens_do_pacote(d: dict) -> list[dict]:
    """O que vai no pacote, com ok/pendência — a lista da tela."""
    livro, pat = d["livro"], d["patrimonio"]
    nao = d["dre_competencia"].get("nao_classificado") or {}
    conc_pend = [c for c in d["conciliacao"] if not c["conciliada"]]
    fech = d["fechamentos"]["meses"]
    abertos = [m["mes"] for m in fech if m["status"] != "fechado" and m["mes"] <= f"{hoje_local():%Y-%m}"]
    pend_bloq = [p for p in d["pendencias"] if p["bloqueia"]]
    return [
        {"id": "dre_competencia", "nome": "DRE gerencial — mês do gasto (competência)", "formato": "PDF + CSV + Excel", "ok": True,
         "detalhe": f"resultado {_br(_linha(d['dre_competencia'], 'RESULTADO_LIQUIDO'))}"},
        {"id": "dre_caixa", "nome": "DRE gerencial — dia do pagamento (caixa)", "formato": "PDF + CSV + Excel", "ok": True,
         "detalhe": f"resultado {_br(_linha(d['dre_caixa'], 'RESULTADO_LIQUIDO'))}"},
        {"id": "livro", "nome": "Livro caixa da atividade rural", "formato": "PDF + CSV + Excel", "ok": not livro["saldo_inicial_pendente"],
         "detalhe": f"{_pl(len(livro['linhas']), 'lançamento', 'lançamentos')} · entradas {_br(livro['entradas'])} · saídas {_br(livro['saidas'])}"
                    + (" · saldo inicial sem abertura conferida" if livro["saldo_inicial_pendente"] else "")},
        {"id": "nao_classificados", "nome": "Lançamentos sem conta ou sem linha da DRE", "formato": "CSV + Excel",
         "ok": not d["nao_classificados"],
         "detalhe": f"{_pl(len(d['nao_classificados']), 'conta', 'contas')} · receita {_br(nao.get('total_receita', 0.0))} · despesa {_br(nao.get('total_despesa', nao.get('total', 0.0)))}"
                    if d["nao_classificados"] else "nenhum"},
        {"id": "pendencias", "nome": "Pendências do fechamento (por mês)", "formato": "CSV + Excel", "ok": not pend_bloq,
         "detalhe": f"{_pl(len(pend_bloq), 'pendência', 'pendências')} em {_pl(len({p['mes'] for p in pend_bloq}), 'mês', 'meses')}" if pend_bloq else "nenhuma"},
        {"id": "patrimonio", "nome": "Patrimônio e depreciação do período", "formato": "PDF + CSV + Excel", "ok": not pat["inconsistencias"],
         "detalhe": f"{_pl(len(pat['bens']), 'bem', 'bens')} · depreciação {_br(pat['depreciacao_periodo'])}"
                    + (f" · {len(pat['inconsistencias'])} sem dados para depreciar" if pat["inconsistencias"] else "")},
        {"id": "conciliacao", "nome": "Conciliação bancária", "formato": "PDF + CSV + Excel", "ok": not conc_pend,
         "detalhe": "todas as contas conciliadas" if not conc_pend else f"{_pl(len(conc_pend), 'conta', 'contas')} a conciliar"},
        {"id": "fechamento", "nome": "Fechamento dos meses (trilha)", "formato": "PDF + CSV", "ok": not abertos,
         "detalhe": "todos os meses fechados" if not abertos else f"{_pl(len(abertos), 'mês', 'meses')} em aberto"},
        {"id": "lcdpr", "nome": "Apoio ao LCDPR (Q100 e Q200)", "formato": "CSV", "ok": d["cabecalho"]["regras_v2"],
         "detalhe": "leiaute de apoio — imóvel e CPF/CNPJ dos participantes ficam em branco"
                    if d["cabecalho"]["regras_v2"] else "sem a natureza dos lançamentos (regras novas desligadas)"},
    ]


def nota_de_metodo(d: dict) -> list[str]:
    v2 = d["cabecalho"]["regras_v2"]
    notas = [
        "Os números saem das mesmas funções dos relatórios do sistema (DRE do servidor, saldo das contas, depreciação do Patrimônio): são idênticos aos das telas.",
        "DRE: cascata gerencial de 15 linhas. Competência = mês do gasto (pago ou não); caixa = dia do pagamento, pelo valor pago.",
        "Livro caixa: o que passou pelo caixa no período, pela data de caixa (cartão avulso: vencimento do cartão) e pelo valor pago. Transferências entre contas próprias não entram.",
        "Fora da DRE por regra: compra de bem (investimento), principal de financiamento, aporte/retirada, transferência e adiantamento.",
        "Não classificado: contas sem linha da DRE — ficam fora do resultado até serem classificadas (a DRE prefere mostrar o buraco a fechar com número errado).",
        "Depreciação: linha reta pelo cadastro de cada bem; não é saída de caixa.",
    ]
    if v2:
        notas.append("Regras novas dos relatórios LIGADAS nesta fazenda (financeiro_regras_v2): natureza do lançamento, folha pelo bruto, juros e descontos da baixa em Outras receitas e despesas, cartão por compra, saldo de abertura.")
        notas.append("No livro caixa, 'entra no resultado' = natureza operacional ou investimento (custeio e investimento da atividade rural); confira com a legislação vigente.")
    else:
        notas.append("Regras novas dos relatórios DESLIGADAS nesta fazenda: a DRE é a das regras antigas e o livro caixa não traz a natureza de cada lançamento (coluna em branco).")
    if d["livro"]["saldo_inicial_pendente"]:
        notas.append("Saldo inicial do livro: há conta corrente sem saldo de abertura conferido com o extrato — o saldo é só a soma dos lançamentos.")
    return notas


# ── arquivos ──────────────────────────────────────────────────────────────
def _csv(cabecalhos: list[str], linhas: list[list], topo: list[list] | None = None) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    for t in topo or []:
        w.writerow(t)
    if topo:
        w.writerow([])
    w.writerow(cabecalhos)
    for l in linhas:
        w.writerow(["" if v is None else (f"{v:.2f}" if isinstance(v, float) else v) for v in l])
    return buf.getvalue().encode("utf-8-sig")


def _topo(cab: dict, titulo: str) -> list[list]:
    return [[titulo], ["Fazenda", cab["fazenda"]], ["CPF/CNPJ", cab["documento"] or "não cadastrado"],
            ["Período", cab["periodo"]["inicio"], cab["periodo"]["fim"]], ["Centro de custo", cab["centro"]],
            ["Gerado em", cab["gerado_em"], "por", cab["gerado_por"]]]


COLS_LIVRO = [("data", "Data"), ("lancamento", "Lançamento"), ("documento", "Documento"), ("tipo_documento", "Tipo de documento"),
              ("participante", "Participante"), ("historico", "Histórico"), ("conta_gerencial", "Conta gerencial"),
              ("conta_corrente", "Conta corrente"), ("natureza", "Natureza"), ("entra_no_resultado", "Entra no resultado da atividade"),
              ("entrada", "Entrada"), ("saida", "Saída"), ("saldo", "Saldo")]
COLS_PAT = [("codigo", "Código"), ("nome", "Bem"), ("tipo", "Tipo"), ("centro_custo", "Centro"), ("data_imobilizacao", "Imobilização"),
            ("valor_aquisicao", "Valor de aquisição"), ("metodo", "Método"), ("vida_util_anos", "Vida útil (anos)"),
            ("depreciacao_periodo", "Depreciação no período"), ("depreciacao_acumulada", "Depreciação acumulada no fim"),
            ("valor_atual", "Valor contábil no fim"), ("data_baixa", "Baixa"), ("valor_baixa", "Valor da baixa"),
            ("inconsistencia", "Pendência")]


def lcdpr_q100(livro: dict) -> tuple[list[str], list[list]]:
    """Apoio no leiaute do registro Q100 do LCDPR (demonstrativo do resultado).
    Só lançamentos que ENTRAM no resultado (natureza conhecida). COD_IMOVEL e
    ID_PARTIC ficam em branco: o sistema não tem o imóvel (CAFIR/CIB) nem o
    CPF/CNPJ do participante."""
    cab = ["DATA", "COD_IMOVEL", "COD_CONTA", "NUM_DOC", "TIPO_DOC", "HIST", "ID_PARTIC", "TIPO_LANC",
           "VL_ENTRADA", "VL_SAIDA", "SLD_FIN", "NAT_SLD_FIN"]
    linhas, saldo = [], 0.0
    for l in livro["linhas"]:
        if l["entra_no_resultado"] != "sim":
            continue
        saldo = round(saldo + l["entrada"] - l["saida"], 2)
        d = l["data"]
        linhas.append([
            f"{d[8:10]}{d[5:7]}{d[0:4]}", "", l["conta_corrente_id"] or "", l["documento"] or l["lancamento"],
            TIPO_DOC_LCDPR.get((l["tipo_documento"] or "").lower(), "6"), l["historico"][:255], "",
            "1" if l["tipo"] == "receita" else "2",
            f"{l['entrada']:.2f}".replace(".", ","), f"{l['saida']:.2f}".replace(".", ","),
            f"{abs(saldo):.2f}".replace(".", ","), "P" if saldo >= 0 else "N",
        ])
    return cab, linhas


def lcdpr_q200(livro: dict) -> tuple[list[str], list[list]]:
    cab = ["MES", "VL_ENTRADA", "VL_SAIDA", "SLD_FIN", "NAT_SLD_FIN"]
    linhas, saldo = [], 0.0
    for m in livro["meses"]:
        e, s = m.get("receitas_atividade") or 0.0, m.get("despesas_atividade") or 0.0
        saldo = round(saldo + e - s, 2)
        linhas.append([f"{m['mes'][5:7]}{m['mes'][0:4]}", f"{e:.2f}".replace(".", ","), f"{s:.2f}".replace(".", ","),
                       f"{abs(saldo):.2f}".replace(".", ","), "P" if saldo >= 0 else "N"])
    return cab, linhas


LEIA_ME_LCDPR = """APOIO AO LCDPR (Livro Caixa Digital do Produtor Rural) — NÃO É O ARQUIVO DE TRANSMISSÃO

O que o sistema já preenche (registros Q100 e Q200, a partir do livro caixa do período):
- DATA, NUM_DOC (nº do documento ou do lançamento), TIPO_DOC (1 nota fiscal, 2 fatura, 3 recibo,
  4 contrato, 5 folha de pagamento, 6 outros — conferir os códigos no manual vigente), HIST,
  TIPO_LANC (1 receita, 2 despesa de custeio/investimento), VL_ENTRADA, VL_SAIDA, SLD_FIN, NAT_SLD_FIN.
- COD_CONTA: o número interno da conta corrente do sistema (cadastre o mesmo código no registro 0050).
- Entram só os lançamentos de natureza operacional ou investimento (exige as regras novas ligadas).

O que FALTA e o contador completa (o sistema não guarda; nada foi inventado):
- Registro 0000/0010/0030: CPF do produtor, forma de apuração, endereço e dados cadastrais.
- Registro 0040/0045: imóveis rurais (CAFIR/CIB, participação) — COD_IMOVEL fica em branco no Q100.
- Registro 0050: contas bancárias (banco, agência, conta) — confira com o cadastro de contas correntes.
- ID_PARTIC: CPF/CNPJ de cada participante (o sistema guarda só o nome).
- Q100 de "receita de produtos entregues por adiantamento" e demais tipos especiais de lançamento.
- Registro 9999 (encerramento) e a assinatura/transmissão pelo programa da Receita.
"""


def gerar_xlsx(d: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    cab = d["cabecalho"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Cabeçalho"
    for linha in _topo(cab, "Pacote do contador"):
        ws.append(linha)
    ws.append([])
    ws.append(["Regras novas dos relatórios", "ligadas" if cab["regras_v2"] else "desligadas"])
    ws.append([])
    ws.append(["Nota de método"])
    for n in d["metodo"]:
        ws.append([n])
    ws["A1"].font = Font(bold=True, size=14)
    num = "#,##0.00"

    def aba(nome: str, cols: list[str], linhas: list[list]):
        w = wb.create_sheet(nome[:31])
        w.append(cols)
        for c in w[1]:
            c.font = Font(bold=True)
        for l in linhas:
            w.append(l)
        for row in w.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value, float):
                    c.number_format = num
        return w

    for chave, nome in (("dre_competencia", "DRE competência"), ("dre_caixa", "DRE caixa")):
        dre = d[chave]
        linhas = []
        for l in dre.get("cascata") or []:
            linhas.append([l["rotulo"], _f(l["valor"]), "subtotal" if l.get("eh_subtotal") else ""])
            for c in l.get("contas") or []:
                linhas.append([f"    {c.get('codigo') or ''} {c.get('nome') or ''}".rstrip(), _f(c.get("valor")), ""])
        aba(nome, ["Linha da DRE", "Valor", ""], linhas)
    livro = d["livro"]
    aba("Livro caixa", [t for _k, t in COLS_LIVRO], [[l[k] for k, _t in COLS_LIVRO] for l in livro["linhas"]])
    aba("Livro caixa (mês)", ["Mês", "Entradas", "Saídas", "Saldo", "Lançamentos", "Receitas da atividade", "Despesas da atividade"],
        [[m["mes"], m["entradas"], m["saidas"], m["saldo"], m["lancamentos"], m["receitas_atividade"], m["despesas_atividade"]] for m in livro["meses"]])
    aba("Não classificados", ["Código", "Conta", "Receita", "Despesa"],
        [[c.get("codigo"), c.get("nome"), _f(c.get("receita", 0.0)), _f(c.get("despesa", c.get("valor")))] for c in d["nao_classificados"]])
    aba("Pendências", ["Mês", "Conferência", "Quantidade", "Impede fechar", "Detalhe"],
        [[p["mes"], p["conferencia"], p["quantidade"], "sim" if p["bloqueia"] else "não", p["detalhe"]] for p in d["pendencias"]])
    aba("Patrimônio", [t for _k, t in COLS_PAT], [[b[k] for k, _t in COLS_PAT] for b in d["patrimonio"]["bens"]])
    aba("Conciliação", ["Conta", "Extrato importado", "Linhas pendentes", "Só no sistema", "Saldo do extrato", "Saldo do sistema", "Diferença", "Data do saldo"],
        [[c["conta"], "sim" if c["importado"] else "não", c["pendentes"], c["so_no_sistema"], c["saldos"].get("saldo_extrato"),
          c["saldos"].get("saldo_sistema"), c["saldos"].get("diferenca"), c["saldos"].get("data")] for c in d["conciliacao"]])
    aba("Fechamentos", ["Mês", "Situação", "SHA-256 do retrato"],
        [[m["mes"], m["status"], m["retrato_sha256"] or ""] for m in d["fechamentos"]["meses"]])
    q100c, q100 = lcdpr_q100(livro)
    aba("LCDPR Q100 (apoio)", q100c, q100)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _esc(v) -> str:
    return "—" if v is None or v == "" else html.escape(str(v))


def _tabela(cols: list[str], linhas: list[list], direita: set[int] = frozenset()) -> str:
    th = "".join(f"<th style='text-align:{'right' if i in direita else 'left'}'>{_esc(c)}</th>" for i, c in enumerate(cols))
    trs = []
    for l in linhas:
        tds = "".join(f"<td style='text-align:{'right' if i in direita else 'left'}'>{v if isinstance(v, _Html) else _esc(v)}</td>" for i, v in enumerate(l))
        trs.append(f"<tr>{tds}</tr>")
    return f"<table><thead><tr>{th}</tr></thead><tbody>{''.join(trs)}</tbody></table>"


class _Html(str):
    """Texto já escapado (não escapa de novo)."""


def gerar_pdf(d: dict) -> bytes:
    from xhtml2pdf import pisa

    from fazenda.rules.manual_fazenda_pdf import _bloquear_recursos_externos

    cab, livro, pat = d["cabecalho"], d["livro"], d["patrimonio"]
    partes = []

    def dre_html(dre: dict, titulo: str) -> str:
        linhas = []
        for l in dre.get("cascata") or []:
            nome = _Html(f"<b>= {_esc(l['rotulo'])}</b>") if l.get("eh_subtotal") else _Html(f"&nbsp;&nbsp;{_esc(l['rotulo'])}")
            linhas.append([nome, _br(l["valor"])])
        nao = dre.get("nao_classificado") or {}
        fora = dre.get("fora_da_dre") or {}
        rod = f"<p class='m'>Fora da DRE: {_esc(_br(fora.get('total')))} · Não classificado: receita {_esc(_br(nao.get('total_receita', 0.0)))}, despesa {_esc(_br(nao.get('total_despesa', nao.get('total', 0.0))))}</p>"
        return f"<h2>{_esc(titulo)}</h2>{_tabela(['Linha', 'Valor'], linhas, {1})}{rod}"

    partes.append(dre_html(d["dre_competencia"], "1. DRE gerencial — mês do gasto (competência)"))
    partes.append(dre_html(d["dre_caixa"], "2. DRE gerencial — dia do pagamento (caixa)"))
    partes.append("<h2>3. Livro caixa da atividade rural</h2>")
    partes.append(f"<p class='m'>Saldo inicial {_esc(_br(livro['saldo_inicial']))}{' (sem saldo de abertura conferido)' if livro['saldo_inicial_pendente'] else ''} · entradas {_esc(_br(livro['entradas']))} · saídas {_esc(_br(livro['saidas']))} · saldo final {_esc(_br(livro['saldo_final']))}</p>")
    partes.append(_tabela(["Mês", "Entradas", "Saídas", "Saldo", "Lançamentos"],
                          [[m["mes"][5:7] + "/" + m["mes"][:4], _br(m["entradas"]), _br(m["saidas"]), _br(m["saldo"]), m["lancamentos"]] for m in livro["meses"]], {1, 2, 3, 4}))
    if len(livro["linhas"]) <= 400:
        partes.append(_tabela(["Data", "Doc.", "Participante", "Histórico", "Natureza", "Entrada", "Saída"],
                              [[_dmy(l["data"]), l["documento"] or l["lancamento"], l["participante"], l["historico"], l["natureza"],
                                _br(l["entrada"]) if l["entrada"] else "", _br(l["saida"]) if l["saida"] else ""] for l in livro["linhas"]], {5, 6}))
    else:
        partes.append(f"<p class='m'>{len(livro['linhas'])} lançamentos: a lista completa está em livro_caixa.csv e na planilha.</p>")
    partes.append("<h2>4. Não classificados e pendências</h2>")
    if d["nao_classificados"]:
        partes.append(_tabela(["Código", "Conta", "Receita", "Despesa"],
                              [[c.get("codigo"), c.get("nome"), _br(c.get("receita", 0.0)), _br(c.get("despesa", c.get("valor")))] for c in d["nao_classificados"]], {2, 3}))
    else:
        partes.append("<p class='m'>Nenhuma conta sem classificação no período.</p>")
    if d["pendencias"]:
        partes.append(_tabela(["Mês", "Conferência", "Qtd.", "Detalhe"],
                              [[p["mes"], p["conferencia"], p["quantidade"], p["detalhe"][:300]] for p in d["pendencias"]], {2}))
    else:
        partes.append("<p class='m'>Nenhuma pendência de fechamento no período.</p>")
    partes.append("<h2>5. Patrimônio e depreciação do período</h2>")
    partes.append(f"<p class='m'>Depreciação do período: {_esc(_br(pat['depreciacao_periodo']))}</p>")
    partes.append(_tabela(["Bem", "Imobilização", "Aquisição", "Depr. período", "Valor contábil", "Pendência"],
                          [[b["nome"], _dmy(b["data_imobilizacao"]), _br(b["valor_aquisicao"]), _br(b["depreciacao_periodo"]),
                            _br(b["valor_atual"]) if b["valor_atual"] is not None else "—", b["inconsistencia"] or (f"baixa {_dmy(b['data_baixa'])}" if b["data_baixa"] else "")]
                           for b in pat["bens"]], {2, 3, 4}))
    partes.append("<h2>6. Conciliação bancária e fechamento</h2>")
    if d["conciliacao"]:
        partes.append(_tabela(["Conta", "Extrato", "Pendentes", "Saldo extrato", "Saldo sistema", "Diferença"],
                              [[c["conta"], "importado" if c["importado"] else "não importado", c["pendentes"],
                                _br(c["saldos"].get("saldo_extrato")), _br(c["saldos"].get("saldo_sistema")), _br(c["saldos"].get("diferenca"))]
                               for c in d["conciliacao"]], {2, 3, 4, 5}))
    meses_f = d["fechamentos"]["meses"]
    partes.append(_tabela(["Mês", "Situação", "Retrato (SHA-256)"], [[m["mes"], m["status"], (m["retrato_sha256"] or "")[:16]] for m in meses_f]))
    partes.append("<h2>7. Nota de método</h2><ul>" + "".join(f"<li>{_esc(n)}</li>" for n in d["metodo"]) + "</ul>")

    doc = f"""<html><head><meta charset="utf-8"><style>
      @page {{ size: a4 portrait; margin: 1.4cm; @frame rodape {{ -pdf-frame-content: rodape; bottom: .6cm; height: .6cm; left: 1.4cm; right: 1.4cm; }} }}
      body {{ font-family: Helvetica; font-size: 8.5pt; color: #111; }}
      h1 {{ font-size: 15pt; margin: 0; }}
      h2 {{ font-size: 11pt; margin: 14px 0 4px; border-bottom: 1px solid #999; padding-bottom: 2px; }}
      table {{ width: 100%; border-collapse: collapse; margin: 4px 0; }}
      th {{ background: #eee; font-weight: bold; padding: 3px 4px; border-bottom: 1px solid #999; }}
      td {{ padding: 2px 4px; border-bottom: 0.5px solid #ddd; }}
      .m {{ color: #444; margin: 2px 0; }}
      .cab td {{ border: none; padding: 1px 4px; }}
    </style></head><body>
      <h1>Pacote do contador · {_esc(cab['fazenda'])}</h1>
      <table class="cab">
        <tr><td><b>CPF/CNPJ</b></td><td>{_esc(cab['documento'] or 'não cadastrado')}</td><td><b>Período</b></td><td>{_dmy(cab['periodo']['inicio'])} a {_dmy(cab['periodo']['fim'])}</td></tr>
        <tr><td><b>Regime</b></td><td>{_esc(cab['regime'])}</td><td><b>Centro</b></td><td>{_esc(cab['centro'])}</td></tr>
        <tr><td><b>Gerado em</b></td><td>{_esc(cab['gerado_em'])}</td><td><b>Por</b></td><td>{_esc(cab['gerado_por'])}</td></tr>
        <tr><td><b>Regras dos relatórios</b></td><td colspan="3">{'novas (financeiro_regras_v2 ligada)' if cab['regras_v2'] else 'antigas (financeiro_regras_v2 desligada)'}</td></tr>
      </table>
      {''.join(partes)}
      <div id="rodape" style="font-size:7pt;color:#666;">{_esc(cab['fazenda'])} · Pacote do contador · {_dmy(cab['periodo']['inicio'])} a {_dmy(cab['periodo']['fim'])} · página <pdf:pagenumber/></div>
    </body></html>"""
    buf = io.BytesIO()
    pisa.CreatePDF(doc, dest=buf, link_callback=_bloquear_recursos_externos)
    return buf.getvalue()


def nome_base(d: dict) -> str:
    import re

    faz = re.sub(r"[^a-z0-9]+", "-", (d["cabecalho"]["fazenda"] or "fazenda").lower()).strip("-")[:40] or "fazenda"
    p = d["cabecalho"]["periodo"]
    return f"pacote-contador_{faz}_{p['inicio']}_{p['fim']}"


def gerar_zip(d: dict) -> bytes:
    from fazenda.api.routers.portal import _dre_para_csv

    cab, livro = d["cabecalho"], d["livro"]
    base = nome_base(d)
    arquivos: list[tuple[str, bytes]] = [
        (f"{base}.pdf", gerar_pdf(d)),
        (f"{base}.xlsx", gerar_xlsx(d)),
        ("dre_competencia.csv", _dre_para_csv(d["dre_competencia"], cab["fazenda"]).encode("utf-8-sig")),
        ("dre_caixa.csv", _dre_para_csv(d["dre_caixa"], cab["fazenda"]).encode("utf-8-sig")),
        ("livro_caixa.csv", _csv([t for _k, t in COLS_LIVRO], [[l[k] for k, _t in COLS_LIVRO] for l in livro["linhas"]],
                                 _topo(cab, "Livro caixa da atividade rural") + [["Saldo inicial", f"{livro['saldo_inicial']:.2f}"]])),
        ("livro_caixa_mensal.csv", _csv(["Mês", "Entradas", "Saídas", "Saldo", "Lançamentos", "Receitas da atividade", "Despesas da atividade"],
                                        [[m["mes"], m["entradas"], m["saidas"], m["saldo"], m["lancamentos"], m["receitas_atividade"], m["despesas_atividade"]] for m in livro["meses"]],
                                        _topo(cab, "Livro caixa — resumo mensal"))),
        ("nao_classificados.csv", _csv(["Código", "Conta", "Receita", "Despesa"],
                                       [[c.get("codigo"), c.get("nome"), _f(c.get("receita", 0.0)), _f(c.get("despesa", c.get("valor")))] for c in d["nao_classificados"]],
                                       _topo(cab, "Não classificados (competência)"))),
        ("pendencias.csv", _csv(["Mês", "Conferência", "Quantidade", "Impede fechar", "Detalhe"],
                                [[p["mes"], p["conferencia"], p["quantidade"], "sim" if p["bloqueia"] else "não", p["detalhe"]] for p in d["pendencias"]],
                                _topo(cab, "Pendências do fechamento"))),
        ("patrimonio_depreciacao.csv", _csv([t for _k, t in COLS_PAT], [[b[k] for k, _t in COLS_PAT] for b in d["patrimonio"]["bens"]],
                                            _topo(cab, "Patrimônio e depreciação do período"))),
        ("conciliacao.csv", _csv(["Conta", "Extrato importado", "Linhas pendentes", "Só no sistema", "Saldo do extrato", "Saldo do sistema", "Diferença", "Data do saldo"],
                                 [[c["conta"], "sim" if c["importado"] else "não", c["pendentes"], c["so_no_sistema"], c["saldos"].get("saldo_extrato"),
                                   c["saldos"].get("saldo_sistema"), c["saldos"].get("diferenca"), c["saldos"].get("data")] for c in d["conciliacao"]],
                                 _topo(cab, "Conciliação bancária"))),
        ("fechamentos.csv", _csv(["Mês", "Ação", "Quem", "Quando (UTC)", "Motivo", "SHA-256 do retrato", "Pendências no fechamento"],
                                 [[e["mes"], e["acao"], e["usuario"], e["criado_em"], e["motivo"] or "", e["retrato_sha256"] or "", e["pendencias_no_fechamento"]]
                                  for e in d["fechamentos"]["trilha"]], _topo(cab, "Trilha do fechamento do mês"))),
    ]
    q100c, q100 = lcdpr_q100(livro)
    q200c, q200 = lcdpr_q200(livro)
    buf_q = io.StringIO()
    w = csv.writer(buf_q, delimiter="|")
    w.writerow(q100c)
    w.writerows(q100)
    arquivos.append(("lcdpr_apoio_Q100.csv", buf_q.getvalue().encode("utf-8-sig")))
    buf_q = io.StringIO()
    w = csv.writer(buf_q, delimiter="|")
    w.writerow(q200c)
    w.writerows(q200)
    arquivos.append(("lcdpr_apoio_Q200.csv", buf_q.getvalue().encode("utf-8-sig")))
    leia = [f"PACOTE DO CONTADOR — {cab['fazenda']}", f"CPF/CNPJ: {cab['documento'] or 'não cadastrado'}",
            f"Período: {_dmy(cab['periodo']['inicio'])} a {_dmy(cab['periodo']['fim'])}", f"Gerado em {cab['gerado_em']} por {cab['gerado_por']}", "",
            "O QUE VAI NO PACOTE"]
    leia += [f"- [{'ok' if i['ok'] else 'pendência'}] {i['nome']} ({i['formato']}): {i['detalhe']}" for i in d["itens"]]
    leia += ["", "NOTA DE MÉTODO"] + [f"- {n}" for n in d["metodo"]] + ["", LEIA_ME_LCDPR]
    arquivos.append(("LEIA-ME.txt", "\n".join(leia).encode("utf-8-sig")))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for nome, conteudo in arquivos:
            z.writestr(nome, conteudo)
    return buf.getvalue()


def resumo_para_tela(d: dict) -> dict:
    """O que a tela mostra antes de gerar (sem as linhas do livro)."""
    livro = d["livro"]
    return {
        "cabecalho": d["cabecalho"], "itens": d["itens"], "metodo": d["metodo"],
        "dre": {
            "competencia": {"resultado": _linha(d["dre_competencia"], "RESULTADO_LIQUIDO"), "receita_liquida": _linha(d["dre_competencia"], "RECEITA_LIQUIDA")},
            "caixa": {"resultado": _linha(d["dre_caixa"], "RESULTADO_LIQUIDO"), "receita_liquida": _linha(d["dre_caixa"], "RECEITA_LIQUIDA")},
        },
        "livro": {k: livro[k] for k in ("saldo_inicial", "saldo_final", "entradas", "saidas", "saldo_inicial_pendente", "receitas_atividade", "despesas_atividade")}
                 | {"lancamentos": len(livro["linhas"]), "meses": livro["meses"]},
        "pendencias": d["pendencias"], "conciliacao": d["conciliacao"], "fechamentos": d["fechamentos"],
        "patrimonio": {"bens": len(d["patrimonio"]["bens"]), "depreciacao_periodo": d["patrimonio"]["depreciacao_periodo"],
                       "inconsistencias": d["patrimonio"]["inconsistencias"]},
        "nao_classificados": d["nao_classificados"],
        "arquivos": ["PDF", "Excel (.xlsx)", "DRE competência e caixa (CSV)", "Livro caixa (CSV)", "Não classificados e pendências (CSV)",
                     "Patrimônio e depreciação (CSV)", "Conciliação e fechamentos (CSV)", "Apoio ao LCDPR (Q100/Q200)", "LEIA-ME"],
    }
