"""
Router de alimentação — plano de dieta por lote, consumo diário estimado,
necessidade mensal (com conversão para sacos) e a baixa automática de
estoque (opção A: o sistema recalcula quantos dias se passaram desde a
última baixa e desconta o consumo acumulado de uma vez, sempre que a tela
é aberta — sem botão manual nem cron real).
"""
from __future__ import annotations

import json
from datetime import date, datetime

from fastapi import APIRouter, Depends, HTTPException, UploadFile
from fastapi.responses import Response
from pydantic import BaseModel
from sqlalchemy import text
from sqlalchemy.exc import IntegrityError
from sqlmodel import Session, select

from fazenda.auth import get_current_user, get_fazenda_atual_id, get_fazenda_id_escrita
from fazenda.database import get_session
from fazenda.models import (
    AlimentacaoEstado, Alimento, AlimentoNutricional, AnaliseBromatologica, Animal, CategoriaAlimento,
    ConsumoAlimento, ConsumoSobra, CurvaABC, Dieta, DietaItemProgramado, DietaLancamento, DietaRegistroReal,
    DietaSimulacaoItem, Estoque, IngredienteMS, LancamentoItem, Lote, MovimentoEstoque, Sanidade,
    TabelaNutricionalProduto, Usuario,
)
from fazenda.rules.alimentacao import calcular_consumo, calcular_necessidade_mensal, resolver_kg_por_unidade, _codigo_grupo
from fazenda.rules.auditoria import fazenda_id_seguro, mapa_usuarios
from fazenda.rules.busca import normalizar_busca
from fazenda.rules.dieta_lancamento import criar_lancamento_programado
from fazenda.rules.producao_leiteira import ultimo_controle_por_animal, com_fallback_animal
from fazenda.rules import estoque_baixa
from fazenda.rules.farmacia import pode_baixar_estoque
from fazenda.rules.parametros import get_param
from fazenda.rules.unidades import converte_para_kg, kg_equivalente

# Nº de tratos por dia (fornecimentos). Hoje são 2.
NUM_TRATOS = 2

router = APIRouter(prefix="/alimentacao", tags=["alimentacao"])

# Alimentos do protocolo padrão — sugestão no lançamento; o campo aceita
# qualquer texto (inclusive itens do Estoque não listados aqui).
ALIMENTOS_PADRAO = [
    "Silagem", "Ração Teck Milk 24%", "Milk Proteico", "Ração Pré-parto", "Corte 21",
    "Ração Bezerro 1", "Ração Bezerro 2",
]


def _dietas_e_animais(session: Session, fazenda_id: int | None) -> tuple[list[dict], list[dict]]:
    """Funde as duas fontes de plano de dieta por lote — achado ao investigar
    por que uma dieta lançada em "Lançar nova dieta" não aparecia em Plano
    por Lote / Necessidade mensal / baixa automática: essas telas sempre
    leram só `Dieta` (linha congelada do DIETA.csv importado uma vez),
    nunca `DietaLancamento`/`DietaItemProgramado` (a tela de lançamento
    "de verdade", com histórico e reconstrução por lote).

    Por lote, a fonte é EXCLUSIVA, nunca somada: um lote com dieta ATIVA
    lançada na tela nova usa só ela (senão o plano novo somaria com o
    import antigo e dobraria o consumo/baixa); um lote sem lançamento
    nenhum na tela nova continua lendo só `Dieta`, como sempre — zero
    mudança de comportamento pra quem nunca usou "Lançar nova dieta"."""
    query_animal = select(Animal).where(Animal.ativo == True)  # noqa: E712
    if fazenda_id is not None:
        query_animal = query_animal.where(Animal.fazenda_id == fazenda_id)
    animais = [
        a.model_dump() for a in session.exec(query_animal).all()
        if not a.eh_semen and a.sexo != "M"
    ]

    query_ativas = select(DietaLancamento).where(DietaLancamento.data_efetivo_encerramento == None)  # noqa: E711
    if fazenda_id is not None:
        query_ativas = query_ativas.where(DietaLancamento.fazenda_id == fazenda_id)
    ativas = session.exec(query_ativas).all()
    lotes_com_lancamento = {d.lote for d in ativas}

    dietas: list[dict] = []
    n_animais_por_lote: dict[int, int] = {}
    for dieta in ativas:
        if dieta.lote not in n_animais_por_lote:
            n_animais_por_lote[dieta.lote] = len(_animais_do_lote(session, dieta.lote, fazenda_id))
        n = n_animais_por_lote[dieta.lote]
        query_itens = select(DietaItemProgramado).where(DietaItemProgramado.dieta_lancamento_id == dieta.id)
        if fazenda_id is not None:
            query_itens = query_itens.where(DietaItemProgramado.fazenda_id == fazenda_id)
        for it in session.exec(query_itens).all():
            qtd_fisica = _quantidade_fisica(it.quantidade, it.unidade, it.base, it.ms_pct)
            base_efetiva = _base_efetiva(it.base_quantidade, dieta.base_quantidade)
            _, por_cabeca = _totais_item(qtd_fisica, base_efetiva, n)
            # Sem `categoria` (None) de propósito: `_remapear_lote_pela_categoria`
            # só remapeia lote congelado de import antigo — `dieta.lote` aqui já
            # é o lote atual escolhido na tela, não precisa de remapeamento.
            dietas.append({
                "lote": dieta.lote, "categoria": None, "ingrediente": it.alimento,
                "quantidade": por_cabeca, "unidade": it.unidade,
            })

    query_dieta = select(Dieta)
    if fazenda_id is not None:
        query_dieta = query_dieta.where(Dieta.fazenda_id == fazenda_id)
    for d in session.exec(query_dieta).all():
        if d.lote in lotes_com_lancamento:
            continue
        dietas.append(d.model_dump())

    return dietas, animais


def _lotes_cadastro(session: Session, fazenda_id: int | None) -> list[dict]:
    query = select(Lote)
    if fazenda_id is not None:
        query = query.where(Lote.fazenda_id == fazenda_id)
    return [l.model_dump() for l in session.exec(query).all()]


def _estoque_por_alimento(session: Session, fazenda_id: int | None) -> tuple[dict[str, list[dict]], set[str]]:
    """Vínculo Alimento → Estoque (ver `Estoque.alimento_id`), chaveado pelo
    nome do Alimento normalizado (trim + minúsculas) — usado como segunda
    tentativa quando o nome do ingrediente do plano de dieta não casa
    diretamente com nenhum `Estoque.nome` (ver `calcular_necessidade_mensal`
    e `_dar_baixa_automatica`)."""
    query_alimento = select(Alimento)
    query_estoque = select(Estoque).where(Estoque.alimento_id.is_not(None))  # type: ignore[union-attr]
    if fazenda_id is not None:
        query_alimento = query_alimento.where(Alimento.fazenda_id == fazenda_id)
        query_estoque = query_estoque.where(Estoque.fazenda_id == fazenda_id)
    alimentos = {a.id: a for a in session.exec(query_alimento).all()}
    por_alimento: dict[str, list[dict]] = {}
    for e in session.exec(query_estoque).all():
        alimento = alimentos.get(e.alimento_id)
        if not alimento:
            continue
        chave = alimento.nome.strip().lower()
        por_alimento.setdefault(chave, []).append(e.model_dump())
    # Move o item preferido (`Alimento.estoque_preferido_id`, Fase P1) para o
    # INÍCIO da lista de candidatos — nunca remove os outros. Todo call site
    # que pega `candidatos[0]` (`_dar_baixa_automatica`, `_resolver_estoque_item`,
    # `calcular_necessidade_mensal`) passa a usar a escolha deliberada sem
    # precisar mudar uma linha. Sem preferência (o padrão, None), a ordem
    # continua EXATAMENTE a da query acima — ver TestEscolhaArbitrariaDeCandidato
    # (T11) em tests/test_migracao_alimento.py, que trava esse comportamento.
    for alimento in alimentos.values():
        if alimento.estoque_preferido_id is None:
            continue
        candidatos = por_alimento.get(alimento.nome.strip().lower())
        if not candidatos:
            continue
        for i, c in enumerate(candidatos):
            if c["id"] == alimento.estoque_preferido_id:
                if i != 0:
                    candidatos.insert(0, candidatos.pop(i))
                break
    cadastrados = {a.nome.strip().lower() for a in alimentos.values()}
    return por_alimento, cadastrados


def _obter_estado_alimentacao(session: Session, fazenda_id: int | None) -> AlimentacaoEstado | None:
    query = select(AlimentacaoEstado)
    if fazenda_id is not None:
        query = query.where(AlimentacaoEstado.fazenda_id == fazenda_id)
    else:
        query = query.where(AlimentacaoEstado.fazenda_id.is_(None))  # type: ignore[union-attr]
    return session.exec(query).first()


def _consumo_total_modo_automatica(por_lote: list[dict], lotes_cadastro: list[dict]) -> list[dict]:
    """Re-agrega `calcular_consumo(...)["por_lote"]` por ingrediente, só para
    os lotes cujo `Lote.modo_baixa_estoque == "automatica"` — mesma soma que
    `calcular_consumo` já faz para `consumo_total`, mas escopada aos lotes que
    de fato pediram a baixa dia-a-dia pelo plano (ver proposta aceita pelo
    proprietário: "automática"/"consumo real"/"sem baixa" por lote).

    Lote sem cadastro (dado legado — ver comentário equivalente em
    `lancar_consumo` sobre `permitir_fora_da_dieta`/`permitir_sem_estoque`)
    cai no padrão restritivo "consumo_real", ou seja, NÃO entra aqui: ausência
    de cadastro nunca é motivo pra afrouxar uma checagem de segurança, e aqui
    a checagem é "não debita estoque sem o lote ter pedido explicitamente"."""
    modo_por_codigo = {l["codigo"]: (l.get("modo_baixa_estoque") or "consumo_real") for l in lotes_cadastro if l.get("codigo")}
    consumo_total: dict[str, dict] = {}
    for info in por_lote:
        modo = modo_por_codigo.get(f"{info['lote']:02d}", "consumo_real")
        if modo != "automatica":
            continue
        for item in info["itens"]:
            chave = item["ingrediente"]
            acc = consumo_total.setdefault(chave, {"ingrediente": chave, "unidade": item["unidade"], "consumo_dia": 0.0})
            acc["consumo_dia"] = round(acc["consumo_dia"] + item["consumo_dia"], 2)
    return sorted(consumo_total.values(), key=lambda x: -x["consumo_dia"])


def _dar_baixa_automatica(session: Session, fazenda_id: int | None) -> dict:
    """
    Baixa automática de estoque por dias decorridos (opção A). Usa uma trava
    otimista (compare-and-swap) na linha de AlimentacaoEstado da fazenda: só
    quem conseguir avançar `ultima_data_deducao` de fato aplica a baixa — uma
    segunda requisição concorrente vê 0 linhas afetadas e não faz nada,
    evitando baixa duplicada quando dois usuários abrem a tela ao mesmo tempo.
    """
    hoje = date.today()
    estado = _obter_estado_alimentacao(session, fazenda_id)
    if not estado:
        # Primeiro acesso desta fazenda: cria a linha de estado. Duas
        # requisições concorrentes podem cair aqui ao mesmo tempo — a segunda
        # perde a corrida na constraint única de fazenda_id; trata como "já
        # criada" e segue sem tentar deduzir nada agora (não há baseline anterior).
        try:
            session.add(AlimentacaoEstado(fazenda_id=fazenda_id, ultima_data_deducao=hoje))
            session.commit()
        except IntegrityError:
            session.rollback()
        return {"dias_deduzidos": 0, "ultima_data_deducao": hoje.isoformat()}

    dias = (hoje - estado.ultima_data_deducao).days if estado.ultima_data_deducao else 0
    if dias <= 0:
        return {"dias_deduzidos": 0, "ultima_data_deducao": estado.ultima_data_deducao.isoformat()}

    if fazenda_id is not None:
        condicao_fazenda = "fazenda_id = :fazenda_id"
    else:
        condicao_fazenda = "fazenda_id IS NULL"
    resultado = session.execute(
        text(
            f"UPDATE alimentacao_estado SET ultima_data_deducao = :novo "
            f"WHERE id = :id AND {condicao_fazenda} AND ultima_data_deducao = :antigo"
        ),
        {"novo": hoje.isoformat(), "id": estado.id, "fazenda_id": fazenda_id, "antigo": estado.ultima_data_deducao.isoformat()},
    )
    session.commit()
    if resultado.rowcount == 0:
        # Outra requisição venceu a corrida e já processou essa janela de dias.
        atualizado = _obter_estado_alimentacao(session, fazenda_id)
        return {"dias_deduzidos": 0, "ultima_data_deducao": atualizado.ultima_data_deducao.isoformat()}

    dietas, animais = _dietas_e_animais(session, fazenda_id)
    lotes_cadastro = _lotes_cadastro(session, fazenda_id)
    por_lote = calcular_consumo(dietas, animais, lotes_cadastro)["por_lote"]
    # Só os lotes em modo "automatica" entram na baixa por dias decorridos —
    # "consumo_real" (padrão) e "sem_baixa" só são tocados (ou nunca são,
    # respectivamente) pelo lançamento manual em `lancar_consumo`.
    consumo_total = _consumo_total_modo_automatica(por_lote, lotes_cadastro)
    estoque_por_alimento, _ = _estoque_por_alimento(session, fazenda_id)

    itens_baixados = []
    avisos: list[str] = []
    for item in consumo_total:
        # Resolução própria (não usa estoque_baixa.resolver_item): além do
        # nome, cai no vínculo via cadastro de Alimento quando o nome do
        # ingrediente do plano não bate direto com nenhum Estoque.nome (mesma
        # resolução usada na necessidade mensal).
        query_estoque_item = select(Estoque).where(Estoque.nome == item["ingrediente"])
        if fazenda_id is not None:
            query_estoque_item = query_estoque_item.where(Estoque.fazenda_id == fazenda_id)
        estoque_item = session.exec(query_estoque_item).first()
        if not estoque_item:
            candidatos = estoque_por_alimento.get((item["ingrediente"] or "").strip().lower())
            if candidatos:
                estoque_item = session.get(Estoque, candidatos[0]["id"])
        # Ingrediente de dieta sem item de Estoque casado (comum — nem todo
        # ingrediente do plano é fisicamente controlado) não gera aviso: essa
        # baixa roda sozinha a cada acesso à tela, sem ação do usuário para
        # reagir a um aviso por ingrediente não cadastrado.
        if not estoque_item or not item["consumo_dia"] or estoque_item.estocavel is False:
            continue
        # Gatilho de comunicação: só deduz insumo cujo estoque inicial/primeira
        # compra já foi registrado (None = insumo legado, mantém comportamento).
        if not pode_baixar_estoque(estoque_item):
            continue
        # O consumo da dieta é sempre em kg — mas o saldo do item de Estoque
        # pode estar em uma unidade ensacada (ex.: "saca 30kg"/"saca 60kg").
        # Sem esta conversão, os kg consumidos eram debitados 1:1 da
        # quantidade em sacas (erro de ~30x/~60x no saldo). Mesma resolução
        # de kg_por_saco já usada em Necessidade Mensal (ver
        # `resolver_kg_por_unidade`, fazenda/rules/alimentacao.py).
        baixa_kg = round(item["consumo_dia"] * dias, 2)
        kg_por_unidade = resolver_kg_por_unidade(estoque_item.model_dump())
        if kg_por_unidade:
            baixa = round(baixa_kg / kg_por_unidade, 4)
            observacao = (
                f"Baixa automática da Alimentação — {dias} dia(s) desde a última baixa "
                f"({baixa_kg} kg ÷ {kg_por_unidade} kg/{estoque_item.unidade} = {baixa} {estoque_item.unidade})"
            )
        else:
            baixa = baixa_kg
            observacao = f"Baixa automática da Alimentação — {dias} dia(s) desde a última baixa"
        avisos.extend(estoque_baixa.movimentar(
            session, item=estoque_item, quantidade=baixa, unidade=estoque_item.unidade, data=hoje,
            fazenda_id=fazenda_id, movimento="Saída de ajuste",
            observacao=observacao,
            origem_tipo="alimentacao", produto=item["ingrediente"],
        ))
        itens_baixados.append({"ingrediente": item["ingrediente"], "baixa": baixa, "baixa_kg": baixa_kg})

    session.commit()
    return {"dias_deduzidos": dias, "ultima_data_deducao": hoje.isoformat(), "itens": itens_baixados, "avisos": avisos}


@router.get("/")
def obter_alimentacao(
    fazenda_id: int | None = Depends(get_fazenda_atual_id), session: Session = Depends(get_session),
) -> dict:
    """Plano de dieta por lote cruzado com o efetivo atual → consumo/dia por ingrediente."""
    fazenda_id = fazenda_id_seguro(fazenda_id)
    _dar_baixa_automatica(session, fazenda_id)
    dietas, animais = _dietas_e_animais(session, fazenda_id)
    return calcular_consumo(dietas, animais, _lotes_cadastro(session, fazenda_id))


@router.get("/necessidade-mensal")
def necessidade_mensal(
    fazenda_id: int | None = Depends(get_fazenda_atual_id), session: Session = Depends(get_session),
) -> dict:
    """Projeção de 30 dias por ingrediente, convertida em sacos quando o item é ensacado."""
    fazenda_id = fazenda_id_seguro(fazenda_id)
    _dar_baixa_automatica(session, fazenda_id)
    dietas, animais = _dietas_e_animais(session, fazenda_id)
    consumo_total = calcular_consumo(dietas, animais, _lotes_cadastro(session, fazenda_id))["consumo_total"]
    query_estoque = select(Estoque)
    if fazenda_id is not None:
        query_estoque = query_estoque.where(Estoque.fazenda_id == fazenda_id)
    estoque_por_nome = {e.nome: e.model_dump() for e in session.exec(query_estoque).all()}
    estoque_por_alimento, alimentos_cadastrados = _estoque_por_alimento(session, fazenda_id)
    return {"itens": calcular_necessidade_mensal(consumo_total, estoque_por_nome, estoque_por_alimento, alimentos_cadastrados)}


@router.get("/estado-baixa")
def estado_baixa(
    fazenda_id: int | None = Depends(get_fazenda_atual_id), session: Session = Depends(get_session),
) -> dict:
    """Última data em que a baixa automática de estoque foi aplicada."""
    fazenda_id = fazenda_id_seguro(fazenda_id)
    estado = _obter_estado_alimentacao(session, fazenda_id)
    return {"ultima_data_deducao": estado.ultima_data_deducao.isoformat() if estado and estado.ultima_data_deducao else None}


# ---------------------------------------------------------------------------
# Categorias de alimento (Configurações > Cadastro > Alimentação >
# Categorias) e cadastro de Alimento (...> Alimentos) — a camada que
# normaliza o vínculo com o Estoque em vez de depender do nome bater
# (ver `calcular_necessidade_mensal`/`_dar_baixa_automatica` acima).
# ---------------------------------------------------------------------------
CATEGORIAS_ALIMENTO_PADRAO = ["Volumoso", "Concentrado", "Mineral"]


def _seed_categorias_alimento(session: Session, fazenda_id: int | None = None) -> None:
    query = select(CategoriaAlimento)
    if fazenda_id is not None:
        query = query.where(CategoriaAlimento.fazenda_id == fazenda_id)
    # (bug pré-existente corrigido) Antes checava só os NOMES padrão que já
    # existiam e reinseria os que faltassem — rodando a cada GET, isso
    # ressuscitava Volumoso/Concentrado/Mineral se o usuário apagasse de
    # propósito. Passa a semear só quando a fazenda não tem NENHUMA
    # categoria (primeira vez de verdade); depois disso apagar uma padrão é
    # definitivo, como em qualquer cadastro editável.
    if session.exec(query).first() is not None:
        return
    novas = [CategoriaAlimento(nome=nome, fazenda_id=fazenda_id) for nome in CATEGORIAS_ALIMENTO_PADRAO]
    session.add_all(novas)
    session.commit()


SUBCATEGORIAS_CONCENTRADO_PADRAO = ["Concentrado Proteico", "Concentrado Energético"]


def _seed_subcategorias_concentrado(session: Session, fazenda_id: int | None = None) -> None:
    """Subcategorias de "Concentrado" nunca tinham sido semeadas (só existiam
    como texto solto num módulo de nutrição sem relação com este cadastro) —
    o select de Subcategoria em Alimentos ficava sempre vazio. Mesma regra de
    "só semeia a primeira vez" do `_seed_categorias_alimento`: dispara apenas
    quando "Concentrado" ainda não tem NENHUMA filha, então apagar as duas
    depois é definitivo, não ressuscita."""
    query = select(CategoriaAlimento)
    if fazenda_id is not None:
        query = query.where(CategoriaAlimento.fazenda_id == fazenda_id)
    categorias = session.exec(query).all()
    concentrado = next((c for c in categorias if c.nome == "Concentrado" and c.categoria_pai_id is None), None)
    if concentrado is None or any(c.categoria_pai_id == concentrado.id for c in categorias):
        return
    novas = [CategoriaAlimento(nome=nome, categoria_pai_id=concentrado.id, fazenda_id=fazenda_id) for nome in SUBCATEGORIAS_CONCENTRADO_PADRAO]
    session.add_all(novas)
    session.commit()


def _ordenar_categorias_hierarquia(categorias: list[CategoriaAlimento]) -> list[CategoriaAlimento]:
    """Agrupa cada subcategoria logo abaixo do próprio pai: raízes por nome
    e, dentro de cada raiz, as filhas por nome — em vez da ordem alfabética
    simples de antes, que espalharia "Proteico" longe de "Concentrado"."""
    por_id = {c.id: c for c in categorias}

    def chave(c: CategoriaAlimento) -> tuple[str, int, str]:
        pai = por_id.get(c.categoria_pai_id) if c.categoria_pai_id is not None else None
        nome_raiz = pai.nome if pai else c.nome
        eh_filha = 1 if c.categoria_pai_id is not None else 0
        return (nome_raiz, eh_filha, c.nome)

    return sorted(categorias, key=chave)


def _validar_categoria_pai(
    session: Session, categoria_pai_id: int | None, fazenda_id: int | None, categoria_id: int | None,
) -> None:
    """Garante o invariante de EXATAMENTE dois níveis (raiz -> subcategoria).
    `categoria_id` é None na criação (nada a comparar ainda) e o id da
    própria categoria na edição — para recusar ela virar pai de si mesma e
    para recusar ela virar subcategoria se já tiver filhas (senão a edição
    criaria um 3º nível por baixo dela, escapando pela porta dos fundos)."""
    if categoria_pai_id is None:
        return
    if categoria_pai_id == categoria_id:
        raise HTTPException(status_code=409, detail="Uma categoria não pode ser pai de si mesma")
    pai = session.get(CategoriaAlimento, categoria_pai_id)
    if not pai or (fazenda_id is not None and pai.fazenda_id != fazenda_id):
        raise HTTPException(status_code=404, detail="Categoria pai não encontrada")
    if pai.categoria_pai_id is not None:
        raise HTTPException(
            status_code=409,
            detail=f'"{pai.nome}" já é uma subcategoria — só dois níveis são permitidos (raiz e subcategoria)',
        )
    if categoria_id is not None:
        tem_filha = session.exec(
            select(CategoriaAlimento).where(CategoriaAlimento.categoria_pai_id == categoria_id)
        ).first()
        if tem_filha:
            raise HTTPException(
                status_code=409,
                detail=f'Categoria tem a subcategoria "{tem_filha.nome}" — não pode também virar subcategoria de outra',
            )


# Alimentos padrão + categoria sugerida — preenche o cadastro na primeira
# vez com os alimentos já usados pela fazenda. Cada entrada é (nome do
# Alimento — conceito NUTRICIONAL genérico, categoria, nome do item de
# Estoque COMERCIAL a vincular automaticamente). Quando o 3º campo é None,
# genérico e comercial são o mesmo termo (ex.: silagens) e o próprio nome
# do Alimento é usado na busca por Estoque; do contrário fica "sem vínculo"
# até o usuário linkar manualmente. Ex.: "Teck Milk 24%" é um produto
# comercial que É um concentrado proteico 24% de proteína — o Alimento
# nasce com o nome nutricional e vinculado ao item de Estoque comercial.
_ALIMENTOS_PADRAO_CATEGORIA: list[tuple[str, str, str | None]] = [
    ("Silagem", "Volumoso", None),
    ("Silagem de milho", "Volumoso", None),
    ("Silagem de sorgo", "Volumoso", None),
    ("Concentrado proteico 24%", "Concentrado", "Teck Milk 24%"),
    ("Milk Proteico", "Concentrado", None),
    ("Corte 21", "Concentrado", None),
    ("Ração Bezerro 1", "Concentrado", None),
    ("Ração Bezerro 2", "Concentrado", None),
    ("Ração Pré-parto", "Mineral", None),
    ("Reprodução 80", "Mineral", None),
]


def seed_alimentos(session: Session, fazenda_id: int | None = None) -> None:
    """Chamado no startup (ver `main.py`), sempre com `fazenda_id=1`: a lista
    `_ALIMENTOS_PADRAO_CATEGORIA` é histórica/grandfathered — nomes de
    produtos comerciais específicos já usados por essa fazenda.

    (bug real corrigido, relato do usuário 01/09/2026: "excluí esses alimentos
    várias vezes e eles sempre voltam") Igual ao que já tinha sido corrigido em
    `_seed_categorias_alimento`/`_seed_subcategorias_concentrado` acima: a
    versão antiga checava nome a nome (`if nome in existentes: continue`) e
    rodava a CADA restart do servidor — apagar "Corte 21" de propósito só
    durava até o próximo deploy, porque o nome já não estava mais em
    `existentes` e o seed o recriava do zero. Bônus: a comparação por string
    exata também nunca reconhecia "Ração Pré-parto" (seed) como igual a
    "Ração Pré-Parto" (já cadastrado pelo usuário, com vínculo de estoque) —
    diferença de maiúscula bastava para duplicar o registro. Agora só semeia
    quando a fazenda está com ZERO alimentos (primeira vez de verdade);
    depois disso, apagar um alimento padrão é definitivo, como em qualquer
    cadastro editável."""
    _seed_categorias_alimento(session, fazenda_id=fazenda_id)
    _seed_subcategorias_concentrado(session, fazenda_id=fazenda_id)
    alimento_query = select(Alimento)
    if fazenda_id is not None:
        alimento_query = alimento_query.where(Alimento.fazenda_id == fazenda_id)
    if session.exec(alimento_query).first() is not None:
        return
    categoria_query = select(CategoriaAlimento)
    estoque_query = select(Estoque)
    if fazenda_id is not None:
        categoria_query = categoria_query.where(CategoriaAlimento.fazenda_id == fazenda_id)
        estoque_query = estoque_query.where(Estoque.fazenda_id == fazenda_id)
    categorias = {c.nome: c.id for c in session.exec(categoria_query).all()}
    estoque_por_nome = {e.nome.strip().lower(): e for e in session.exec(estoque_query).all()}
    novos_com_vinculo = [
        (Alimento(nome=nome, categoria_alimento_id=categorias.get(categoria_nome), fazenda_id=fazenda_id), nome_estoque or nome)
        for nome, categoria_nome, nome_estoque in _ALIMENTOS_PADRAO_CATEGORIA
    ]
    session.add_all([a for a, _ in novos_com_vinculo])
    session.commit()
    for alimento, nome_estoque in novos_com_vinculo:
        session.refresh(alimento)
        item = estoque_por_nome.get(nome_estoque.strip().lower())
        if item and item.alimento_id is None:
            item.alimento_id = alimento.id
            session.add(item)
    session.commit()


@router.get("/categorias")
def listar_categorias_alimento(
    fazenda_id: int | None = Depends(get_fazenda_atual_id), session: Session = Depends(get_session),
) -> list[dict]:
    fazenda_id = fazenda_id_seguro(fazenda_id)
    _seed_categorias_alimento(session, fazenda_id=fazenda_id)
    _seed_subcategorias_concentrado(session, fazenda_id=fazenda_id)
    query = select(CategoriaAlimento)
    if fazenda_id is not None:
        query = query.where(CategoriaAlimento.fazenda_id == fazenda_id)
    # Ordenação em Python (não dá pra expressar "filha logo abaixo do pai"
    # num único ORDER BY simples sem self-join) — ver _ordenar_categorias_hierarquia.
    categorias = session.exec(query).all()
    return [c.model_dump() for c in _ordenar_categorias_hierarquia(categorias)]


class CategoriaAlimentoIn(BaseModel):
    nome: str
    ativo: bool = True
    # None = raiz. Ver decisão de modelagem da sessão: só dois níveis, a
    # validação fica em `_validar_categoria_pai`.
    categoria_pai_id: int | None = None


@router.post("/categorias", status_code=201)
def criar_categoria_alimento(
    dados: CategoriaAlimentoIn, fazenda_id: int = Depends(get_fazenda_id_escrita), session: Session = Depends(get_session),
) -> dict:
    _validar_categoria_pai(session, dados.categoria_pai_id, fazenda_id, categoria_id=None)
    # NULL não colide em UniqueConstraint (nome, categoria_pai_id, fazenda_id)
    # — duas raízes de mesmo nome passariam batido pela constraint do banco.
    # A checagem em código é quem garante nome único DENTRO DO MESMO PAI, e
    # é ela que permite "Proteico" existir tanto sob "Concentrado" quanto
    # sob "Volumoso" (mesmo nome, pais diferentes).
    query_dup = select(CategoriaAlimento).where(
        CategoriaAlimento.nome == dados.nome, CategoriaAlimento.categoria_pai_id == dados.categoria_pai_id,
    )
    if fazenda_id is not None:
        query_dup = query_dup.where(CategoriaAlimento.fazenda_id == fazenda_id)
    if session.exec(query_dup).first():
        raise HTTPException(status_code=409, detail=f'Já existe uma categoria chamada "{dados.nome}"')
    cat = CategoriaAlimento(
        nome=dados.nome, ativo=dados.ativo, fazenda_id=fazenda_id, categoria_pai_id=dados.categoria_pai_id,
    )
    session.add(cat)
    session.commit()
    session.refresh(cat)
    return cat.model_dump()


@router.put("/categorias/{categoria_id}")
def atualizar_categoria_alimento(
    categoria_id: int, dados: CategoriaAlimentoIn, fazenda_id: int | None = Depends(get_fazenda_atual_id), session: Session = Depends(get_session),
) -> dict:
    fazenda_id = fazenda_id_seguro(fazenda_id)
    cat = session.get(CategoriaAlimento, categoria_id)
    # (bug pré-existente corrigido) `session.get` não confere a fazenda do
    # registro encontrado — sem essa checagem, uma fazenda edita categoria
    # de outra só sabendo o id. Mesmo padrão já usado em `atualizar_alimento`.
    if not cat or (fazenda_id is not None and cat.fazenda_id != fazenda_id):
        raise HTTPException(status_code=404, detail="Categoria não encontrada")
    _validar_categoria_pai(session, dados.categoria_pai_id, fazenda_id, categoria_id=categoria_id)
    query_outra = select(CategoriaAlimento).where(
        CategoriaAlimento.nome == dados.nome, CategoriaAlimento.categoria_pai_id == dados.categoria_pai_id,
        CategoriaAlimento.id != categoria_id,
    )
    if fazenda_id is not None:
        query_outra = query_outra.where(CategoriaAlimento.fazenda_id == fazenda_id)
    outra = session.exec(query_outra).first()
    if outra:
        raise HTTPException(status_code=409, detail=f'Já existe uma categoria chamada "{dados.nome}"')
    cat.nome = dados.nome
    cat.ativo = dados.ativo
    cat.categoria_pai_id = dados.categoria_pai_id
    session.add(cat)
    session.commit()
    return cat.model_dump()


@router.delete("/categorias/{categoria_id}")
def excluir_categoria_alimento(
    categoria_id: int, fazenda_id: int | None = Depends(get_fazenda_atual_id), session: Session = Depends(get_session),
) -> dict:
    fazenda_id = fazenda_id_seguro(fazenda_id)
    cat = session.get(CategoriaAlimento, categoria_id)
    # (bug pré-existente corrigido) Era o único endpoint do bloco sem filtro
    # de fazenda_id — uma fazenda conseguia apagar categoria de outra só
    # sabendo o id. Mesmo padrão de checagem pós-`session.get` do PUT acima.
    if not cat or (fazenda_id is not None and cat.fazenda_id != fazenda_id):
        raise HTTPException(status_code=404, detail="Categoria não encontrada")
    tem_filha = session.exec(select(CategoriaAlimento).where(CategoriaAlimento.categoria_pai_id == categoria_id)).first()
    if tem_filha:
        raise HTTPException(
            status_code=409,
            detail=f'Categoria tem a subcategoria "{tem_filha.nome}" — mova ou exclua a(s) subcategoria(s) primeiro',
        )
    em_uso = session.exec(select(Alimento).where(Alimento.categoria_alimento_id == categoria_id)).first()
    if em_uso:
        raise HTTPException(status_code=409, detail=f'Categoria em uso pelo alimento "{em_uso.nome}" — mova ou exclua o(s) alimento(s) primeiro')
    session.delete(cat)
    session.commit()
    return {"ok": True}


def _serializar_alimento(session: Session, a: Alimento, fazenda_id: int | None) -> dict:
    query = select(Estoque).where(Estoque.alimento_id == a.id)
    if fazenda_id is not None:
        query = query.where(Estoque.fazenda_id == fazenda_id)
    vinculados = session.exec(query).all()
    return {**a.model_dump(), "estoque_vinculado": [e.model_dump() for e in vinculados]}


@router.get("/alimentos")
def listar_alimentos(
    fazenda_id: int | None = Depends(get_fazenda_atual_id), session: Session = Depends(get_session),
) -> list[dict]:
    fazenda_id = fazenda_id_seguro(fazenda_id)
    query = select(Alimento)
    if fazenda_id is not None:
        query = query.where(Alimento.fazenda_id == fazenda_id)
    return [_serializar_alimento(session, a, fazenda_id) for a in session.exec(query.order_by(Alimento.nome)).all()]


class AlimentoIn(BaseModel):
    nome: str
    categoria_alimento_id: int | None = None
    observacao: str | None = None
    ativo: bool = True
    # Itens de Estoque a vincular a este alimento — substitui o conjunto
    # anterior (ver `_vincular_estoque_ao_alimento`).
    estoque_ids: list[int] = []


def _vincular_estoque_ao_alimento(session: Session, alimento_id: int, estoque_ids: list[int], fazenda_id: int | None) -> None:
    """Substitui o conjunto de itens de Estoque vinculados a este Alimento
    pelos informados. Um item de Estoque só pode estar vinculado a UM
    alimento por vez (campo escalar `Estoque.alimento_id`) — vincular aqui
    "rouba" o vínculo de qualquer outro alimento que o item estivesse usando."""
    query = select(Estoque).where(Estoque.alimento_id == alimento_id)
    if fazenda_id is not None:
        query = query.where(Estoque.fazenda_id == fazenda_id)
    atuais = session.exec(query).all()
    for e in atuais:
        if e.id not in estoque_ids:
            e.alimento_id = None
            session.add(e)
    for eid in estoque_ids:
        item = session.get(Estoque, eid)
        if item and item.alimento_id != alimento_id:
            item.alimento_id = alimento_id
            session.add(item)
    session.commit()


@router.post("/alimentos", status_code=201)
def criar_alimento(
    dados: AlimentoIn, fazenda_id: int = Depends(get_fazenda_id_escrita), session: Session = Depends(get_session),
) -> dict:
    query_dup = select(Alimento).where(Alimento.nome == dados.nome)
    if fazenda_id is not None:
        query_dup = query_dup.where(Alimento.fazenda_id == fazenda_id)
    if session.exec(query_dup).first():
        raise HTTPException(status_code=409, detail=f'Já existe um alimento chamado "{dados.nome}"')
    alimento = Alimento(
        nome=dados.nome, categoria_alimento_id=dados.categoria_alimento_id,
        observacao=dados.observacao, ativo=dados.ativo, fazenda_id=fazenda_id,
    )
    session.add(alimento)
    session.commit()
    session.refresh(alimento)
    _vincular_estoque_ao_alimento(session, alimento.id, dados.estoque_ids, fazenda_id)
    return _serializar_alimento(session, alimento, fazenda_id)


@router.put("/alimentos/{alimento_id}")
def atualizar_alimento(
    alimento_id: int, dados: AlimentoIn, fazenda_id: int | None = Depends(get_fazenda_atual_id), session: Session = Depends(get_session),
) -> dict:
    fazenda_id = fazenda_id_seguro(fazenda_id)
    alimento = session.get(Alimento, alimento_id)
    if not alimento or (fazenda_id is not None and alimento.fazenda_id != fazenda_id):
        raise HTTPException(status_code=404, detail="Alimento não encontrado")
    query_outro = select(Alimento).where(Alimento.nome == dados.nome, Alimento.id != alimento_id)
    if fazenda_id is not None:
        query_outro = query_outro.where(Alimento.fazenda_id == fazenda_id)
    outro = session.exec(query_outro).first()
    if outro:
        raise HTTPException(status_code=409, detail=f'Já existe um alimento chamado "{dados.nome}"')
    alimento.nome = dados.nome
    alimento.categoria_alimento_id = dados.categoria_alimento_id
    alimento.observacao = dados.observacao
    alimento.ativo = dados.ativo
    alimento.atualizado_em = datetime.utcnow()
    session.add(alimento)
    session.commit()
    _vincular_estoque_ao_alimento(session, alimento_id, dados.estoque_ids, fazenda_id)
    return _serializar_alimento(session, alimento, fazenda_id)


@router.delete("/alimentos/{alimento_id}")
def excluir_alimento(
    alimento_id: int, session: Session = Depends(get_session),
    fazenda_id: int | None = Depends(get_fazenda_atual_id),
) -> dict:
    """(bug real corrigido, relato do usuário 01/09/2026: "Failed to fetch"
    ao excluir) `Alimento.id` é referenciado por FK opcional em outras 6
    tabelas além de `Estoque` (histórico de dieta programada/consumida,
    análise bromatológica, tabela nutricional, biblioteca nutricional da
    Formulação de Dietas e item de simulação) — SQLite (usado nos testes)
    não valida FK por padrão, então o `session.delete()` sempre passava por
    aqui sem erro; o Postgres de produção rejeita a exclusão com uma
    violação de integridade sempre que o Alimento ainda está referenciado
    por qualquer uma delas, e a exceção não tratada aparecia no navegador
    como "Failed to fetch". Cada linha referenciada perde só o vínculo
    (nome/valores continuam intactos, gravados como texto/snapshot à parte
    em todas elas) — nenhuma perde dado histórico."""
    alimento = session.get(Alimento, alimento_id)
    if not alimento or (fazenda_id is not None and alimento.fazenda_id != fazenda_id):
        raise HTTPException(status_code=404, detail="Alimento não encontrado")
    for modelo in (
        Estoque, DietaItemProgramado, ConsumoAlimento, AnaliseBromatologica,
        TabelaNutricionalProduto, AlimentoNutricional, DietaSimulacaoItem,
    ):
        for row in session.exec(select(modelo).where(modelo.alimento_id == alimento_id)).all():
            row.alimento_id = None
            session.add(row)
    # Flush explícito ANTES do delete — garante que os UPDATEs (desvincular)
    # cheguem ao banco antes do DELETE do Alimento em si. Sem isso, nada
    # garante a ordem entre um DELETE avulso e UPDATEs pendentes na mesma
    # sessão (ver o mesmo cuidado em exclusoes.py::_excluir_alvos_em_ordem,
    # achado idêntico no animal 1291 no mesmo dia).
    session.flush()
    session.delete(alimento)
    session.commit()
    return {"ok": True}


class EstoquePreferidoIn(BaseModel):
    # None limpa a preferência (volta a ordem arbitrária de hoje).
    estoque_id: int | None = None


@router.put("/alimentos/{alimento_id}/estoque-preferido")
def atualizar_estoque_preferido(
    alimento_id: int, dados: EstoquePreferidoIn,
    fazenda_id: int | None = Depends(get_fazenda_atual_id), session: Session = Depends(get_session),
) -> dict:
    """Fase P1 do refactor Alimento/Estoque — escolhe QUAL item de Estoque
    vinculado a este Alimento recebe a baixa automática/consumo manual
    quando há 2+ candidatos (ver `Alimento.estoque_preferido_id` e
    `_estoque_por_alimento`). O item escolhido precisa já estar vinculado a
    ESTE alimento (`Estoque.alimento_id == alimento_id`) — escolher um
    candidato que nem é candidato não faz sentido, daí o 400 (e não 404: o
    alimento existe, o problema é a combinação)."""
    fazenda_id = fazenda_id_seguro(fazenda_id)
    alimento = session.get(Alimento, alimento_id)
    if not alimento or (fazenda_id is not None and alimento.fazenda_id != fazenda_id):
        raise HTTPException(status_code=404, detail="Alimento não encontrado")
    if dados.estoque_id is not None:
        query = select(Estoque).where(Estoque.id == dados.estoque_id, Estoque.alimento_id == alimento_id)
        if fazenda_id is not None:
            query = query.where(Estoque.fazenda_id == fazenda_id)
        if not session.exec(query).first():
            raise HTTPException(
                status_code=400,
                detail="O item de estoque informado não está vinculado a este alimento",
            )
    alimento.estoque_preferido_id = dados.estoque_id
    alimento.atualizado_em = datetime.utcnow()
    session.add(alimento)
    session.commit()
    return _serializar_alimento(session, alimento, fazenda_id)


class EstoqueCategoriaIn(BaseModel):
    # None desvincula (o item volta a depender só da categoria indireta via
    # Alimento, quando houver).
    categoria_alimento_id: int | None = None


@router.put("/estoque/{estoque_id}/categoria")
def atualizar_categoria_estoque(
    estoque_id: int, dados: EstoqueCategoriaIn,
    fazenda_id: int | None = Depends(get_fazenda_atual_id), session: Session = Depends(get_session),
) -> dict:
    """Fase P1 do refactor Alimento/Estoque — vincula um item de Estoque
    DIRETO a uma CategoriaAlimento (`Estoque.categoria_alimento_id`), sem
    precisar passar pelo cadastro de Alimento (a única via até aqui — ver
    seção `produtos_sem_categoria` do relatório de conferência). Puramente
    aditivo: não mexe em `Estoque.categoria` (texto livre) nem na
    categorização indireta via `Alimento.categoria_alimento_id`, que
    continua valendo do mesmo jeito para quem não usar este endpoint."""
    fazenda_id = fazenda_id_seguro(fazenda_id)
    item = session.get(Estoque, estoque_id)
    if not item or (fazenda_id is not None and item.fazenda_id != fazenda_id):
        raise HTTPException(status_code=404, detail="Item de estoque não encontrado")
    if dados.categoria_alimento_id is not None:
        cat = session.get(CategoriaAlimento, dados.categoria_alimento_id)
        if not cat or (fazenda_id is not None and cat.fazenda_id != fazenda_id):
            raise HTTPException(status_code=404, detail="Categoria de alimento não encontrada")
    item.categoria_alimento_id = dados.categoria_alimento_id
    session.add(item)
    session.commit()
    session.refresh(item)
    return item.model_dump()


# ---------------------------------------------------------------------------
# Fase P0-B do refactor Alimento/Estoque — relatório de conferência SOMENTE
# LEITURA (nenhum session.add/commit/delete nesta seção), o retrato que o
# usuário confere antes de autorizar as fases seguintes (que vão eliminar a
# camada `Alimento` da interface). Sete seções, cada uma carregando sua(s)
# tabela(s) de origem em UMA consulta e cruzando em Python — nunca query
# dentro de laço (mesmo padrão de `_estoque_por_alimento` acima).
# ---------------------------------------------------------------------------
def _contagem_por_nome_normalizado(valores: list) -> dict[str, int]:
    """Conta ocorrências de cada texto normalizado numa lista (nomes de
    ingrediente/produto vindos de uma fonte) — usado para explicar de onde
    veio cada item fantasma da importação (seção `fantasmas_importacao`)."""
    contagem: dict[str, int] = {}
    for v in valores:
        chave = normalizar_busca(v)
        if not chave:
            continue
        contagem[chave] = contagem.get(chave, 0) + 1
    return contagem


def _candidatos_mesclagem(item: Estoque, todos: list[tuple[int, str, str]]) -> list[dict]:
    """Outros itens de Estoque da mesma fazenda cujo nome normalizado seja
    igual, contenha, ou esteja contido no nome deste item — candidatos a
    mesclagem manual numa fase futura (aqui só identificados, nunca
    mesclados). `todos` é a lista (id, nome, nome_normalizado) de TODO o
    estoque da fazenda, montada uma única vez fora do laço de chamada."""
    nome_norm = normalizar_busca(item.nome)
    if not nome_norm:
        return []
    candidatos = []
    for outro_id, outro_nome, outro_norm in todos:
        if outro_id == item.id or not outro_norm:
            continue
        if nome_norm == outro_norm or nome_norm in outro_norm or outro_norm in nome_norm:
            candidatos.append({"id": outro_id, "nome": outro_nome})
    return candidatos


def _finalidade_indica_alimento(finalidade: str | None) -> bool:
    """"Ração/Alimento" é a finalidade explícita de hoje; qualquer outra cujo
    nome normalizado contenha "aliment", "nutri" ou "racao" também conta —
    cobre finalidades customizadas por fazenda com o mesmo sentido."""
    if not finalidade:
        return False
    norm = normalizar_busca(finalidade)
    return any(termo in norm for termo in ("aliment", "nutri", "racao"))


@router.get("/migracao/relatorio")
def relatorio_migracao(
    fazenda_id: int | None = Depends(get_fazenda_atual_id), session: Session = Depends(get_session),
) -> dict:
    """Relatório de conferência (Fase P0-B) — NUNCA escreve no banco. Mostra
    o estado real dos dados de Alimento/Estoque/Dieta antes do refactor que
    vai eliminar a camada `Alimento` da interface: itens fantasma deixados
    pela importação antiga, "pontes" nunca usadas, alimentos sem categoria,
    agrupamentos que precisarão ser desmembrados, nomes que divergem entre
    Alimento e Estoque (com o custo em laudos que isso implicaria), ingredientes
    de dieta que a resolução de hoje não casa (ou casa com ambiguidade), e o
    impacto exato da futura regra de elegibilidade ao RMCA."""
    fazenda_id = fazenda_id_seguro(fazenda_id)

    def _f(query, modelo):
        return query.where(modelo.fazenda_id == fazenda_id) if fazenda_id is not None else query

    estoques = session.exec(_f(select(Estoque), Estoque)).all()
    movimentos = session.exec(_f(select(MovimentoEstoque), MovimentoEstoque)).all()
    alimentos = session.exec(_f(select(Alimento), Alimento)).all()
    alimentos_por_id = {a.id: a for a in alimentos}
    dieta_ingredientes = session.exec(_f(select(Dieta.ingrediente), Dieta)).all()
    curva_abc_produtos = session.exec(_f(select(CurvaABC.produto), CurvaABC)).all()
    lancamento_item_produtos = session.exec(_f(select(LancamentoItem.produto), LancamentoItem)).all()
    sanidade_produtos = session.exec(_f(select(Sanidade.produto), Sanidade)).all()
    analises = session.exec(_f(select(AnaliseBromatologica), AnaliseBromatologica)).all()

    # ── índices montados uma vez, cruzados em Python (nunca query em laço) ──
    contagem_dieta = _contagem_por_nome_normalizado(dieta_ingredientes)
    contagem_curva_abc = _contagem_por_nome_normalizado(curva_abc_produtos)
    contagem_lancamento_item = _contagem_por_nome_normalizado(lancamento_item_produtos)
    contagem_sanidade = _contagem_por_nome_normalizado(sanidade_produtos)

    mov_por_estoque_id: dict[int, list[MovimentoEstoque]] = {}
    mov_sem_vinculo_por_nome: dict[str, list[MovimentoEstoque]] = {}
    for m in movimentos:
        if m.estoque_id is not None:
            mov_por_estoque_id.setdefault(m.estoque_id, []).append(m)
        else:
            mov_sem_vinculo_por_nome.setdefault(normalizar_busca(m.nome_item), []).append(m)

    def _resumo_movimentos(item: Estoque) -> dict:
        relacionados = mov_por_estoque_id.get(item.id, []) + mov_sem_vinculo_por_nome.get(normalizar_busca(item.nome), [])
        if not relacionados:
            return {"quantidade_movimentos": 0, "primeiro_movimento": None, "ultimo_movimento": None}
        datas = sorted(m.data_movimento for m in relacionados if m.data_movimento)
        return {
            "quantidade_movimentos": len(relacionados),
            "primeiro_movimento": datas[0].isoformat() if datas else None,
            "ultimo_movimento": datas[-1].isoformat() if datas else None,
        }

    todos_estoque_norm = [(e.id, e.nome, normalizar_busca(e.nome)) for e in estoques]

    def _linha_fantasma(item: Estoque) -> dict:
        chave = normalizar_busca(item.nome)
        return {
            "id": item.id, "nome": item.nome, "quantidade": item.quantidade, "unidade": item.unidade,
            "finalidade": item.finalidade,
            "fontes": {
                "dieta": contagem_dieta.get(chave, 0), "curva_abc": contagem_curva_abc.get(chave, 0),
                "lancamento_item": contagem_lancamento_item.get(chave, 0), "sanidade": contagem_sanidade.get(chave, 0),
            },
            **_resumo_movimentos(item),
            "candidatos_mesclagem": _candidatos_mesclagem(item, todos_estoque_norm),
        }

    # ── 1) fantasmas da importação: finalidade/alimento_id nulos, quantidade 0 ──
    fantasmas_importacao = [
        _linha_fantasma(e) for e in estoques
        if e.finalidade is None and e.alimento_id is None and (e.quantidade or 0) == 0
    ]

    # ── 2) "pontes" (finalidade Ração/Alimento + alimento_id) sem NENHUM movimento ──
    # DE PROPÓSITO estrito no literal, ao contrário das seções 3 e 7: a ponte é
    # criada em código com exatamente esse valor (alimentoEstoqueBridge). Alargar
    # para `_finalidade_indica_alimento` varreria para cá todo produto REAL
    # recém-cadastrado que ainda não teve movimento — o próprio "Caroço de
    # Algodão" — e o relatório o ofereceria como fantasma candidato a exclusão.
    fantasmas_ponte = [
        _linha_fantasma(e) for e in estoques
        if e.finalidade == "Ração/Alimento" and e.alimento_id is not None
        and _resumo_movimentos(e)["quantidade_movimentos"] == 0
    ]

    # ── 3) produtos que são alimento mas não têm classificação ──
    # Aditivo (Fase P1): um item com `categoria_alimento_id` PRÓPRIO (ver
    # `Estoque.categoria_alimento_id`) já tem categoria por si só, mesmo sem
    # nenhum vínculo de Alimento — qualquer um dos dois caminhos (direto ou
    # indireto via Alimento) satisfaz "tem categoria", então sai desta lista.
    produtos_sem_categoria = []
    for e in estoques:
        if e.categoria_alimento_id is not None:
            continue
        if e.alimento_id is not None:
            alimento = alimentos_por_id.get(e.alimento_id)
            if alimento is not None and alimento.categoria_alimento_id is None:
                produtos_sem_categoria.append({
                    "id": e.id, "nome": e.nome, "quantidade": e.quantidade, "unidade": e.unidade,
                    "finalidade": e.finalidade,
                    "motivo": f'Vinculado ao alimento "{alimento.nome}", que não tem categoria cadastrada',
                })
        elif _finalidade_indica_alimento(e.finalidade):
            produtos_sem_categoria.append({
                "id": e.id, "nome": e.nome, "quantidade": e.quantidade, "unidade": e.unidade,
                "finalidade": e.finalidade,
                "motivo": f'Finalidade "{e.finalidade}" indica alimento, mas o item não está vinculado a nenhum Alimento cadastrado',
            })

    # ── 4) Alimento com 2+ itens de Estoque vinculados (desmembramento futuro) ──
    estoque_por_alimento_id: dict[int, list[Estoque]] = {}
    for e in estoques:
        if e.alimento_id is not None:
            estoque_por_alimento_id.setdefault(e.alimento_id, []).append(e)
    alimento_ids_desmembrar = [aid for aid, itens in estoque_por_alimento_id.items() if len(itens) >= 2]
    alimento_ids_com_nutricional: set[int] = set()
    if alimento_ids_desmembrar:
        query_nutri = select(AlimentoNutricional.alimento_id).where(
            AlimentoNutricional.alimento_id.in_(alimento_ids_desmembrar)
        )
        alimento_ids_com_nutricional = {aid for aid in session.exec(query_nutri).all() if aid is not None}
    desmembramentos = [
        {
            "alimento_id": aid,
            "alimento_nome": alimentos_por_id[aid].nome if aid in alimentos_por_id else None,
            # Fase P1: qual dos candidatos (se algum) já foi escolhido
            # deliberadamente como o item preferido — ver
            # `Alimento.estoque_preferido_id` e PUT
            # /alimentacao/alimentos/{alimento_id}/estoque-preferido.
            "estoque_preferido_id": alimentos_por_id[aid].estoque_preferido_id if aid in alimentos_por_id else None,
            "produtos": [
                {"id": e.id, "nome": e.nome, "quantidade": e.quantidade, "unidade": e.unidade}
                for e in estoque_por_alimento_id[aid]
            ],
            "tem_alimento_nutricional": aid in alimento_ids_com_nutricional,
        }
        for aid in sorted(alimento_ids_desmembrar)
    ]

    # ── 5) Alimento.nome != Estoque.nome do produto vinculado ──
    contagem_laudos_por_nome_exato: dict[str, int] = {}
    contagem_laudos_por_alimento_id: dict[int, int] = {}
    for a in analises:
        if a.alimento:
            contagem_laudos_por_nome_exato[a.alimento] = contagem_laudos_por_nome_exato.get(a.alimento, 0) + 1
        if a.alimento_id is not None:
            contagem_laudos_por_alimento_id[a.alimento_id] = contagem_laudos_por_alimento_id.get(a.alimento_id, 0) + 1
    divergencia_nome = []
    for e in estoques:
        if e.alimento_id is None:
            continue
        alimento = alimentos_por_id.get(e.alimento_id)
        if alimento is not None and alimento.nome != e.nome:
            divergencia_nome.append({
                "alimento_id": alimento.id, "alimento_nome": alimento.nome,
                "estoque_id": e.id, "estoque_nome": e.nome,
                # Histórico: até a Fase P1, AnaliseBromatologica.alimento_id nunca
                # era preenchido — o vínculo real era só por igualdade exata de
                # string com o nome ATUAL do alimento. Renomear custaria exatamente
                # esta quantidade de laudos. Mantido como está (a API agora escreve
                # `alimento_id` na criação e um backfill cobriu os já existentes,
                # mas nomes podem ter divergido de novo desde então).
                "quantidade_laudos_pelo_nome_atual": contagem_laudos_por_nome_exato.get(alimento.nome, 0),
                # Fase P1: quantos laudos já estão ligados por id (imunes a um
                # futuro rename) — quanto maior, menor o risco que a linha acima
                # descreve. Informativo apenas; não muda nenhuma resolução.
                "quantidade_laudos_pelo_id": contagem_laudos_por_alimento_id.get(alimento.id, 0),
            })

    # ── 6) Dieta.ingrediente que a cascata de resolução de hoje não resolve ──
    # (nome exato de Estoque; senão, via Alimento.nome — ver `_estoque_por_alimento`
    # e `calcular_necessidade_mensal`, ambos usados como referência aqui). Usa
    # `.strip().lower()` (não `normalizar_busca`, usado no resto desta seção)
    # DE PROPÓSITO — é a MESMA normalização exata que `calcular_necessidade_mensal`
    # aplica em produção; usar `normalizar_busca` aqui faria este relatório
    # "resolver" ingredientes que a rotina real de hoje não resolve.
    estoque_por_nome_exato = {e.nome: e for e in estoques}
    estoque_por_alimento, alimentos_cadastrados = _estoque_por_alimento(session, fazenda_id)
    ingredientes_nao_resolviveis = []
    for ingrediente in sorted({i for i in dieta_ingredientes if i}):
        if ingrediente in estoque_por_nome_exato:
            continue  # resolve_1 — não entra no relatório
        nome_norm = ingrediente.strip().lower()
        candidatos = estoque_por_alimento.get(nome_norm, [])
        if len(candidatos) == 1:
            continue  # resolve_1 — não entra no relatório
        if len(candidatos) > 1:
            classificacao, motivo = "ambiguo", "Vários itens de Estoque vinculados ao mesmo Alimento — a rotina de hoje usa o primeiro em silêncio"
        elif nome_norm in alimentos_cadastrados:
            classificacao, motivo = "resolve_0", "Alimento cadastrado, mas sem nenhum item de Estoque vinculado"
        else:
            classificacao, motivo = "resolve_0", "Nenhum item de Estoque nem Alimento com este nome"
        ingredientes_nao_resolviveis.append({
            "ingrediente": ingrediente, "classificacao": classificacao, "motivo": motivo,
            "candidatos": [{"id": c["id"], "nome": c["nome"]} for c in candidatos],
        })

    # ── 7) RMCA: regra atual (conta gerencial) × regra futura (finalidade) ──
    # "Nutrição" NÃO está em FINALIDADES_ESTOQUE (ver rules/categorias.py) — mas
    # isso não a torna irrelevante: a lista é só a semente, e o cadastro de
    # finalidade é livre (routers/cadastro/estoque.py), então a fazenda já criou
    # "Nutrição" à mão. Foi exatamente esse valor que sumiu o "Caroço de Algodão"
    # dos seletores e originou este refactor. Testar aqui pelo literal
    # "Ração/Alimento" faria a seção mais importante do relatório — a que diz
    # quais itens entram no indicador de custo — ignorar justamente os itens
    # reais da fazenda. Por isso usa `_finalidade_indica_alimento`, o mesmo
    # predicado da seção 3, que casa "aliment"/"nutri"/"racao" normalizados.
    def _elegivel_conta(e: Estoque) -> bool:
        return bool(e.conta_gerencial_despesa_padrao and e.conta_gerencial_despesa_padrao.startswith("3.01.01"))

    def _elegivel_finalidade(e: Estoque) -> bool:
        return _finalidade_indica_alimento(e.finalidade)

    def _linha_rmca(e: Estoque) -> dict:
        return {
            "id": e.id, "nome": e.nome,
            "conta_gerencial_despesa_padrao": e.conta_gerencial_despesa_padrao, "finalidade": e.finalidade,
        }

    rmca = {
        "so_pela_conta": [_linha_rmca(e) for e in estoques if _elegivel_conta(e) and not _elegivel_finalidade(e)],
        "so_pela_finalidade": [_linha_rmca(e) for e in estoques if _elegivel_finalidade(e) and not _elegivel_conta(e)],
        "por_ambas": [_linha_rmca(e) for e in estoques if _elegivel_conta(e) and _elegivel_finalidade(e)],
    }

    return {
        "fantasmas_importacao": fantasmas_importacao,
        "fantasmas_ponte": fantasmas_ponte,
        "produtos_sem_categoria": produtos_sem_categoria,
        "desmembramentos": desmembramentos,
        "divergencia_nome": divergencia_nome,
        "ingredientes_nao_resolviveis": ingredientes_nao_resolviveis,
        "rmca": rmca,
    }


# ---------------------------------------------------------------------------
# Lançamento de dieta (Lançamentos > Alimentação) — histórico por lote, com
# comparação entre o programado (nutricionista) e o real oferecido. Só uma
# dieta pode estar ativa (sem encerramento efetivo) por lote de cada vez.
# ---------------------------------------------------------------------------
@router.get("/alimentos-padrao")
def alimentos_padrao() -> list[str]:
    return ALIMENTOS_PADRAO


# % de matéria seca padrão por ingrediente (volumosos ~33%, concentrados ~88%);
# o usuário edita na aba Matéria seca.
_MS_PADRAO = {
    "Silagem": 33.24, "Ração Teck Milk 24%": 88.0, "Milk Proteico": 88.0, "Ração Pré-parto": 88.0,
    "Corte 21": 88.0, "Ração Bezerro 1": 88.0, "Ração Bezerro 2": 88.0,
}


def _seed_materia_seca(session: Session) -> None:
    from fazenda.models import IngredienteMS
    existentes = {i.nome for i in session.exec(select(IngredienteMS)).all()}
    novos = [IngredienteMS(nome=nome, ms_pct=_MS_PADRAO.get(nome)) for nome in ALIMENTOS_PADRAO if nome not in existentes]
    if novos:
        for n in novos:
            session.add(n)
        session.commit()


@router.get("/materia-seca")
def listar_materia_seca(
    fazenda_id: int | None = Depends(get_fazenda_atual_id), session: Session = Depends(get_session),
) -> list[dict]:
    """Lista de ingredientes padrão com seu % de matéria seca (editável)."""
    from fazenda.models import IngredienteMS
    _seed_materia_seca(session)
    fazenda_id = fazenda_id_seguro(fazenda_id)
    query = select(IngredienteMS)
    if fazenda_id is not None:
        query = query.where(IngredienteMS.fazenda_id == fazenda_id)
    itens = session.exec(query.order_by(IngredienteMS.nome)).all()
    return [i.model_dump() for i in itens]


class IngredienteMSIn(BaseModel):
    nome: str
    ms_pct: float | None = None


@router.put("/materia-seca")
def salvar_materia_seca(
    dados: IngredienteMSIn, fazenda_id: int = Depends(get_fazenda_id_escrita), session: Session = Depends(get_session),
) -> dict:
    """Upsert do % de matéria seca de um ingrediente (cadastro/edição)."""
    from fazenda.models import IngredienteMS
    nome = dados.nome.strip()
    if not nome:
        raise HTTPException(status_code=400, detail="Nome do ingrediente é obrigatório")
    if dados.ms_pct is not None and not (0 < dados.ms_pct <= 100):
        raise HTTPException(status_code=400, detail="% de matéria seca deve ser maior que 0 e no máximo 100.")
    query = select(IngredienteMS).where(IngredienteMS.nome == nome)
    if fazenda_id is not None:
        query = query.where(IngredienteMS.fazenda_id == fazenda_id)
    item = session.exec(query).first()
    if not item:
        item = IngredienteMS(nome=nome, fazenda_id=fazenda_id)
    item.ms_pct = dados.ms_pct
    item.atualizado_em = datetime.utcnow()
    session.add(item)
    session.commit()
    session.refresh(item)
    return item.model_dump()


def _seed_tabela_nutricional(session: Session) -> None:
    from fazenda.models import TabelaNutricionalProduto, TabelaNutricionalValor
    from fazenda.rules.tabela_nutricional import ALIMENTOS, LINHAS
    if session.exec(select(TabelaNutricionalProduto)).first():
        return
    produtos = [TabelaNutricionalProduto(nome=nome, ordem=i) for i, nome in enumerate(ALIMENTOS)]
    for p in produtos:
        session.add(p)
    session.commit()
    for p in produtos:
        session.refresh(p)
    for linha in LINHAS:
        nutriente = linha[0]
        for p, valor in zip(produtos, linha[1:]):
            if valor:
                session.add(TabelaNutricionalValor(produto_id=p.id, nutriente=nutriente, valor=valor))
    session.commit()


def _tabela_nutricional_montada(session: Session, fazenda_id: int | None = None):
    from fazenda.models import TabelaNutricionalProduto, TabelaNutricionalValor
    query_produto = select(TabelaNutricionalProduto)
    query_valor = select(TabelaNutricionalValor)
    if fazenda_id is not None:
        query_produto = query_produto.where(TabelaNutricionalProduto.fazenda_id == fazenda_id)
        query_valor = query_valor.where(TabelaNutricionalValor.fazenda_id == fazenda_id)
    produtos = session.exec(query_produto.order_by(TabelaNutricionalProduto.ordem, TabelaNutricionalProduto.nome)).all()
    valores = session.exec(query_valor.order_by(TabelaNutricionalValor.id)).all()
    por_produto: dict[int, dict[str, str]] = {}
    nutrientes_ordem: list[str] = []
    for v in valores:
        por_produto.setdefault(v.produto_id, {})[v.nutriente] = v.valor
        if v.nutriente not in nutrientes_ordem:
            nutrientes_ordem.append(v.nutriente)
    return produtos, nutrientes_ordem, por_produto


@router.get("/tabela-nutricional")
def obter_tabela_nutricional(
    fazenda_id: int | None = Depends(get_fazenda_atual_id), session: Session = Depends(get_session),
) -> dict:
    """Tabela nutricional (nutriente × produto), cadastrável em Alimentação >
    Tabela nutricional — consulta rápida (modal + calculadora) e edição."""
    _seed_tabela_nutricional(session)
    fazenda_id = fazenda_id_seguro(fazenda_id)
    produtos, nutrientes_ordem, por_produto = _tabela_nutricional_montada(session, fazenda_id)
    linhas = [[nutriente] + [por_produto.get(p.id, {}).get(nutriente, "") for p in produtos] for nutriente in nutrientes_ordem]
    return {
        "alimentos": [p.nome for p in produtos], "produto_ids": [p.id for p in produtos],
        "estoque_ids": [p.estoque_id for p in produtos], "linhas": linhas,
    }


class TabelaNutricionalProdutoIn(BaseModel):
    nome: str = ""
    # Produto do cadastro fechado de Estoque (finalidade Ração/Alimento) —
    # quando informado, `nome` é derivado do Estoque se não vier preenchido;
    # None = texto livre (flag "outro produto" na tela).
    estoque_id: int | None = None


@router.post("/tabela-nutricional/produtos", status_code=201)
def criar_produto_tabela_nutricional(
    dados: TabelaNutricionalProdutoIn, fazenda_id: int = Depends(get_fazenda_id_escrita), session: Session = Depends(get_session),
) -> dict:
    from fazenda.models import TabelaNutricionalProduto
    nome = dados.nome.strip()
    if dados.estoque_id is not None:
        produto_estoque = session.get(Estoque, dados.estoque_id)
        if not produto_estoque or (fazenda_id is not None and produto_estoque.fazenda_id != fazenda_id):
            raise HTTPException(status_code=404, detail="Produto de estoque não encontrado")
        nome = nome or produto_estoque.nome
        query_dup_estoque = select(TabelaNutricionalProduto).where(TabelaNutricionalProduto.estoque_id == dados.estoque_id)
        if fazenda_id is not None:
            query_dup_estoque = query_dup_estoque.where(TabelaNutricionalProduto.fazenda_id == fazenda_id)
        if session.exec(query_dup_estoque).first():
            raise HTTPException(status_code=409, detail=f'O produto "{produto_estoque.nome}" já está na tabela nutricional')
    if not nome:
        raise HTTPException(status_code=400, detail="Nome do produto é obrigatório")
    query_dup = select(TabelaNutricionalProduto).where(TabelaNutricionalProduto.nome == nome)
    query_ordem = select(TabelaNutricionalProduto).order_by(TabelaNutricionalProduto.ordem.desc())
    if fazenda_id is not None:
        query_dup = query_dup.where(TabelaNutricionalProduto.fazenda_id == fazenda_id)
        query_ordem = query_ordem.where(TabelaNutricionalProduto.fazenda_id == fazenda_id)
    if session.exec(query_dup).first():
        raise HTTPException(status_code=409, detail=f'Já existe um produto chamado "{nome}" na tabela nutricional')
    maior_ordem = session.exec(query_ordem).first()
    produto = TabelaNutricionalProduto(
        nome=nome, ordem=(maior_ordem.ordem + 1) if maior_ordem else 0, fazenda_id=fazenda_id, estoque_id=dados.estoque_id,
    )
    session.add(produto)
    session.commit()
    session.refresh(produto)
    return produto.model_dump()


@router.put("/tabela-nutricional/produtos/{produto_id}")
def renomear_produto_tabela_nutricional(
    produto_id: int, dados: TabelaNutricionalProdutoIn, session: Session = Depends(get_session),
    fazenda_id: int | None = Depends(get_fazenda_atual_id),
) -> dict:
    from fazenda.models import TabelaNutricionalProduto
    fazenda_id = fazenda_id_seguro(fazenda_id)
    produto = session.get(TabelaNutricionalProduto, produto_id)
    # BUG DE SEGURANÇA CORRIGIDO: sem esta checagem, qualquer usuário podia
    # renomear o produto de tabela nutricional de outra fazenda.
    if not produto or (fazenda_id is not None and produto.fazenda_id != fazenda_id):
        raise HTTPException(status_code=404, detail="Produto não encontrado")
    nome = dados.nome.strip()
    if not nome:
        raise HTTPException(status_code=400, detail="Nome do produto é obrigatório")
    produto.nome = nome
    session.add(produto)
    session.commit()
    session.refresh(produto)
    return produto.model_dump()


@router.delete("/tabela-nutricional/produtos/{produto_id}")
def excluir_produto_tabela_nutricional(
    produto_id: int, session: Session = Depends(get_session),
    fazenda_id: int | None = Depends(get_fazenda_atual_id),
) -> dict:
    from fazenda.models import TabelaNutricionalProduto, TabelaNutricionalValor
    fazenda_id = fazenda_id_seguro(fazenda_id)
    produto = session.get(TabelaNutricionalProduto, produto_id)
    # BUG DE SEGURANÇA CORRIGIDO: sem esta checagem, qualquer usuário podia
    # excluir permanentemente o produto de tabela nutricional de outra
    # fazenda.
    if not produto or (fazenda_id is not None and produto.fazenda_id != fazenda_id):
        raise HTTPException(status_code=404, detail="Produto não encontrado")
    for v in session.exec(select(TabelaNutricionalValor).where(TabelaNutricionalValor.produto_id == produto_id)).all():
        session.delete(v)
    session.delete(produto)
    session.commit()
    return {"excluido": True, "id": produto_id}


@router.post("/tabela-nutricional/produtos/{produto_id}/gerar-composicao")
def gerar_composicao_de_tabela_nutricional(
    produto_id: int, categoria_nasem: str | None = None,
    fazenda_id: int = Depends(get_fazenda_id_escrita), session: Session = Depends(get_session),
    user: Usuario = Depends(get_current_user),
) -> dict:
    """Usa os valores já digitados na Tabela Nutricional deste produto (texto
    livre, unidades mistas) para preencher/atualizar a composição dele na
    Biblioteca de Referência (AlimentoNutricional) — em vez de deixar a
    importação de dieta cair sempre no template genérico da categoria.

    Requer o produto já vinculado a um item de Estoque (que por sua vez
    aponta pra um Alimento) — é essa cadeia que decide em qual `alimento_id`/
    `estoque_id` a composição gerada é gravada. Ver
    `fazenda.rules.tabela_nutricional.compor_alimento_nutricional_de_tabela`
    pra regra de conversão de unidade."""
    from fazenda.models import TabelaNutricionalValor
    from fazenda.rules.tabela_nutricional import compor_alimento_nutricional_de_tabela

    produto = session.get(TabelaNutricionalProduto, produto_id)
    if not produto or (produto.fazenda_id is not None and produto.fazenda_id != fazenda_id):
        raise HTTPException(status_code=404, detail="Produto não encontrado")

    alimento_id = produto.alimento_id
    if alimento_id is None and produto.estoque_id is not None:
        produto_estoque = session.get(Estoque, produto.estoque_id)
        alimento_id = produto_estoque.alimento_id if produto_estoque else None
    if alimento_id is None:
        raise HTTPException(
            status_code=400,
            detail="Vincule este produto a um item de Estoque (que aponte pra um Alimento) antes de gerar a composição",
        )

    valores = session.exec(select(TabelaNutricionalValor).where(TabelaNutricionalValor.produto_id == produto_id)).all()
    nutriente_valores = {v.nutriente: v.valor for v in valores}
    convertidos, nao_convertidos = compor_alimento_nutricional_de_tabela(nutriente_valores)
    if not convertidos:
        raise HTTPException(status_code=400, detail="Nenhum valor desta tabela nutricional pôde ser convertido em composição")

    query_existente = select(AlimentoNutricional).where(
        AlimentoNutricional.alimento_id == alimento_id, AlimentoNutricional.estoque_id == produto.estoque_id,
        AlimentoNutricional.fazenda_id == fazenda_id,
    )
    item = session.exec(query_existente).first()
    criado = item is None
    if item is None:
        item = AlimentoNutricional(
            alimento_id=alimento_id, estoque_id=produto.estoque_id, fazenda_id=fazenda_id,
            nome=produto.nome, categoria_nasem=categoria_nasem or "Outros", usuario_id=user.id,
        )
    for campo, valor in convertidos.items():
        setattr(item, campo, valor)
    item.fonte = "Tabela nutricional (rótulo do produto)"
    extras = json.loads(item.extras_json) if item.extras_json else {}
    extras.update(nao_convertidos)
    item.extras_json = json.dumps(extras) if extras else None
    item.atualizado_em = datetime.utcnow()
    session.add(item)
    session.commit()
    session.refresh(item)
    return {
        "criado": criado, "alimento_nutricional_id": item.id,
        "convertidos": convertidos, "nao_convertidos": nao_convertidos,
    }


class ValorTabelaNutricionalIn(BaseModel):
    produto_id: int
    nutriente: str
    valor: str = ""


class SalvarValoresTabelaNutricionalIn(BaseModel):
    itens: list[ValorTabelaNutricionalIn]


@router.put("/tabela-nutricional/valores")
def salvar_valores_tabela_nutricional(
    dados: SalvarValoresTabelaNutricionalIn, fazenda_id: int = Depends(get_fazenda_id_escrita), session: Session = Depends(get_session),
) -> dict:
    """Upsert em lote — salva a grade inteira (nutriente × produto) de uma vez."""
    from fazenda.models import TabelaNutricionalValor
    query = select(TabelaNutricionalValor)
    if fazenda_id is not None:
        query = query.where(TabelaNutricionalValor.fazenda_id == fazenda_id)
    existentes = {(v.produto_id, v.nutriente): v for v in session.exec(query).all()}
    salvos = 0
    for item in dados.itens:
        nutriente = item.nutriente.strip()
        if not nutriente:
            continue
        chave = (item.produto_id, nutriente)
        v = existentes.get(chave)
        if v:
            v.valor = item.valor
            session.add(v)
        elif item.valor.strip():
            session.add(TabelaNutricionalValor(produto_id=item.produto_id, nutriente=nutriente, valor=item.valor, fazenda_id=fazenda_id))
        salvos += 1
    session.commit()
    return {"salvos": salvos}


@router.get("/tabela-nutricional/modelo")
def baixar_modelo_tabela_nutricional(
    fazenda_id: int | None = Depends(get_fazenda_atual_id), session: Session = Depends(get_session),
) -> Response:
    """Planilha (.xlsx) com os produtos e nutrientes já cadastrados — baixe,
    edite/complete e reimporte em POST /tabela-nutricional/importar."""
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    _seed_tabela_nutricional(session)
    fazenda_id = fazenda_id_seguro(fazenda_id)
    produtos, nutrientes_ordem, por_produto = _tabela_nutricional_montada(session, fazenda_id)

    wb = Workbook()
    ws = wb.active
    ws.title = "Tabela nutricional"
    cabecalho = ["Nutriente"] + [p.nome for p in produtos]
    ws.append(cabecalho)
    for cel in ws[1]:
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor="4A6B3A")
    for nutriente in nutrientes_ordem:
        ws.append([nutriente] + [por_produto.get(p.id, {}).get(nutriente, "") for p in produtos])
    for idx in range(1, len(cabecalho) + 1):
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = 22

    buf = io.BytesIO()
    wb.save(buf)
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=tabela_nutricional_modelo.xlsx"},
    )


@router.post("/tabela-nutricional/importar")
async def importar_tabela_nutricional(
    file: UploadFile, fazenda_id: int = Depends(get_fazenda_id_escrita), session: Session = Depends(get_session),
) -> dict:
    """
    Importa a planilha no mesmo formato do modelo baixado: 1ª coluna =
    nutriente, demais colunas = um produto cada (nome no cabeçalho). Produtos
    com nome novo são criados; existentes (mesmo nome) têm os valores
    atualizados/completados.
    """
    import io
    from openpyxl import load_workbook
    from fazenda.models import TabelaNutricionalProduto, TabelaNutricionalValor

    content = await file.read()
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="Não foi possível ler o arquivo — envie um .xlsx válido")
    ws = wb.active
    linhas = list(ws.iter_rows(values_only=True))
    if not linhas:
        raise HTTPException(status_code=400, detail="Planilha vazia")
    cabecalho = [str(c).strip() if c is not None else "" for c in linhas[0]]
    if len(cabecalho) < 2 or not any(cabecalho[1:]):
        raise HTTPException(status_code=400, detail="A planilha precisa de ao menos uma coluna de produto (além de 'Nutriente')")

    query_produtos = select(TabelaNutricionalProduto)
    if fazenda_id is not None:
        query_produtos = query_produtos.where(TabelaNutricionalProduto.fazenda_id == fazenda_id)
    existentes = {p.nome: p for p in session.exec(query_produtos).all()}
    maior_ordem = max([p.ordem for p in existentes.values()], default=-1)
    produto_por_coluna: dict[int, "TabelaNutricionalProduto"] = {}
    for idx, nome in enumerate(cabecalho[1:], start=1):
        if not nome:
            continue
        produto = existentes.get(nome)
        if not produto:
            maior_ordem += 1
            produto = TabelaNutricionalProduto(nome=nome, ordem=maior_ordem, fazenda_id=fazenda_id)
            session.add(produto)
            session.commit()
            session.refresh(produto)
            existentes[nome] = produto
        produto_por_coluna[idx] = produto

    query_valores = select(TabelaNutricionalValor)
    if fazenda_id is not None:
        query_valores = query_valores.where(TabelaNutricionalValor.fazenda_id == fazenda_id)
    valores_existentes = {(v.produto_id, v.nutriente): v for v in session.exec(query_valores).all()}
    nutrientes_importados = 0
    for row in linhas[1:]:
        if not row or not row[0]:
            continue
        nutriente = str(row[0]).strip()
        teve_valor = False
        for idx, produto in produto_por_coluna.items():
            valor_bruto = row[idx] if idx < len(row) else None
            valor = "" if valor_bruto is None else str(valor_bruto).strip()
            if not valor:
                continue
            teve_valor = True
            chave = (produto.id, nutriente)
            v = valores_existentes.get(chave)
            if v:
                v.valor = valor
                session.add(v)
            else:
                v = TabelaNutricionalValor(produto_id=produto.id, nutriente=nutriente, valor=valor, fazenda_id=fazenda_id)
                session.add(v)
                valores_existentes[chave] = v
        if teve_valor:
            nutrientes_importados += 1
    session.commit()
    return {"produtos": len(produto_por_coluna), "nutrientes": nutrientes_importados}


# ─────────────────────────── Análise bromatológica ───────────────────────────
# Laudo de laboratório de um lote/silo de alimento da própria fazenda — não
# confundir com Tabela Nutricional (referência padrão) ou Matéria seca (só o
# %MS por ingrediente genérico). Registro pontual, sem edição/exclusão (mesmo
# padrão de Qualidade do leite).
class AnaliseBromatologicaIn(BaseModel):
    data: date
    alimento: str
    # Fase P1 do refactor Alimento/Estoque — vínculo explícito com o cadastro
    # de Alimento. Quando omitido, `criar_analise_bromatologica` tenta
    # resolver sozinho por igualdade exata de nome (ver docstring do
    # endpoint); passar aqui é só para quando o usuário já escolheu o
    # Alimento num seletor, em vez de digitar o nome livre.
    alimento_id: int | None = None
    ms_pct: float | None = None
    pb_pct: float | None = None
    fdn_pct: float | None = None
    fda_pct: float | None = None
    ndt_pct: float | None = None
    ee_pct: float | None = None
    cinzas_pct: float | None = None
    ca_pct: float | None = None
    p_pct: float | None = None
    observacao: str | None = None


@router.get("/analise-bromatologica")
def listar_analise_bromatologica(
    fazenda_id: int | None = Depends(get_fazenda_atual_id), session: Session = Depends(get_session),
) -> dict:
    from fazenda.models import AnaliseBromatologica
    fazenda_id = fazenda_id_seguro(fazenda_id)
    query = select(AnaliseBromatologica)
    if fazenda_id is not None:
        query = query.where(AnaliseBromatologica.fazenda_id == fazenda_id)
    registros = session.exec(query.order_by(AnaliseBromatologica.data.desc())).all()
    nomes = mapa_usuarios(session, {r.usuario_id for r in registros})
    linhas = []
    for r in registros:
        linha = r.model_dump()
        linha["usuario_nome"] = nomes.get(linha.pop("usuario_id"))
        linhas.append(linha)
    return {"registros": linhas, "total": len(registros)}


@router.post("/analise-bromatologica", status_code=201)
def criar_analise_bromatologica(
    dados: AnaliseBromatologicaIn, fazenda_id: int = Depends(get_fazenda_id_escrita),
    session: Session = Depends(get_session), user: Usuario = Depends(get_current_user),
) -> dict:
    """Grava um laudo. Gap 1 da Fase P1 (ver relatório de conferência §
    divergência de nome): até aqui `AnaliseBromatologica.alimento_id` nunca
    era escrito, e o vínculo real com o Alimento era só por igualdade exata
    de string — renomear o Alimento "perdia" o laudo em silêncio (ver
    `/formulacao/alimentos/{id}/resolver`, que já prioriza `alimento_id`
    quando presente). Se `dados.alimento_id` vier explícito (usuário
    escolheu num seletor), valida que pertence a esta fazenda e usa direto;
    senão, tenta resolver sozinho por nome exato — melhor esforço, não
    bloqueia a criação do laudo se não achar."""
    from fazenda.models import AnaliseBromatologica
    if not dados.alimento.strip():
        raise HTTPException(status_code=400, detail="Alimento é obrigatório")

    alimento_id = dados.alimento_id
    if alimento_id is not None:
        query = select(Alimento).where(Alimento.id == alimento_id)
        if fazenda_id is not None:
            query = query.where(Alimento.fazenda_id == fazenda_id)
        if not session.exec(query).first():
            raise HTTPException(status_code=404, detail="Alimento não encontrado")
    else:
        query = select(Alimento).where(Alimento.nome == dados.alimento)
        if fazenda_id is not None:
            query = query.where(Alimento.fazenda_id == fazenda_id)
        encontrado = session.exec(query).first()
        alimento_id = encontrado.id if encontrado else None

    registro = AnaliseBromatologica(
        **dados.model_dump(exclude={"alimento_id"}), alimento_id=alimento_id,
        usuario_id=user.id, fazenda_id=fazenda_id,
    )
    session.add(registro)
    session.commit()
    session.refresh(registro)
    return registro.model_dump()


class ItemProgramadoIn(BaseModel):
    alimento: str
    quantidade: float
    unidade: str
    base: str | None = None       # "MN" (matéria natural) | "MS" (matéria seca)
    ms_pct: float | None = None   # % de matéria seca do alimento
    # Override por item de `DietaLancamento.base_quantidade` — "total" ou
    # "animal"; None (padrão) herda a base da dieta. Ver `_base_efetiva`.
    base_quantidade: str | None = None


def _quantidade_fisica(quantidade: float, unidade: str | None, base: str | None, ms_pct: float | None) -> float:
    """Converte a quantidade lançada em matéria seca (MS) para o físico
    (matéria natural, MN) a oferecer de fato no vagão — só se aplica a
    unidades de massa (kg/g); nas demais (L, dose, unidade...) a conversão de
    MS não faz sentido e a quantidade é usada como está."""
    if base == "MS" and ms_pct and (unidade or "").lower() in ("kg", "g"):
        return quantidade / (ms_pct / 100)
    return quantidade


class DietaLancamentoIn(BaseModel):
    lote: int
    responsavel: str | None = None
    data_abertura: date
    data_prevista_encerramento: date | None = None
    observacao: str | None = None
    base_quantidade: str | None = None  # "total" (padrão) | "animal" (por cabeça/dia)
    leite_bezerros_kg_dia: float | None = None  # kg/dia de leite p/ bezerros (relatório Controle × Entregue)
    itens: list[ItemProgramadoIn]
    # Se True e já houver dieta ativa no lote, encerra-a na data de início
    # desta (o veterinário responde "sim" ao salvar). Se False e houver ativa,
    # devolve 409 como antes.
    encerrar_anterior: bool = False


def _serializar_dieta(session: Session, d: DietaLancamento, fazenda_id: int | None) -> dict:
    query = select(DietaItemProgramado).where(DietaItemProgramado.dieta_lancamento_id == d.id)
    if fazenda_id is not None:
        query = query.where(DietaItemProgramado.fazenda_id == fazenda_id)
    itens = session.exec(query).all()
    return {**d.model_dump(), "ativa": d.data_efetivo_encerramento is None, "itens_programados": [i.model_dump() for i in itens]}


@router.get("/dietas")
def listar_dietas(
    lote: int | None = None, ativo: bool | None = None,
    fazenda_id: int | None = Depends(get_fazenda_atual_id), session: Session = Depends(get_session),
) -> list[dict]:
    fazenda_id = fazenda_id_seguro(fazenda_id)
    query = select(DietaLancamento)
    if fazenda_id is not None:
        query = query.where(DietaLancamento.fazenda_id == fazenda_id)
    dietas = session.exec(query).all()
    saida = [_serializar_dieta(session, d, fazenda_id) for d in dietas]
    nomes = mapa_usuarios(session, {s["usuario_id"] for s in saida})
    for s in saida:
        s["usuario_nome"] = nomes.get(s["usuario_id"])
    if lote is not None:
        saida = [s for s in saida if s["lote"] == lote]
    if ativo is not None:
        saida = [s for s in saida if s["ativa"] == ativo]
    return sorted(saida, key=lambda s: s["data_abertura"], reverse=True)


@router.post("/dietas", status_code=201)
def criar_dieta(
    dados: DietaLancamentoIn, fazenda_id: int = Depends(get_fazenda_id_escrita),
    session: Session = Depends(get_session), user: Usuario = Depends(get_current_user),
) -> dict:
    dieta = criar_lancamento_programado(
        session, fazenda_id, user.id,
        lote=dados.lote, responsavel=dados.responsavel, data_abertura=dados.data_abertura,
        data_prevista_encerramento=dados.data_prevista_encerramento, observacao=dados.observacao,
        base_quantidade=dados.base_quantidade, leite_bezerros_kg_dia=dados.leite_bezerros_kg_dia,
        itens=[item.model_dump() for item in dados.itens], encerrar_anterior=dados.encerrar_anterior,
    )
    return _serializar_dieta(session, dieta, fazenda_id)


def _animais_do_lote(session: Session, lote: int, fazenda_id: int | None = None) -> list[Animal]:
    query = select(Animal).where(Animal.ativo == True)  # noqa: E712
    if fazenda_id is not None:
        query = query.where(Animal.fazenda_id == fazenda_id)
    ativos = session.exec(query).all()
    return [
        a for a in ativos
        if not a.eh_semen and a.sexo != "M" and (_codigo_grupo(a.grupo_primario) or "") == f"{lote:02d}"
    ]


@router.get("/dietas/contexto/{lote}")
def contexto_dieta(
    lote: int, fazenda_id: int | None = Depends(get_fazenda_atual_id), session: Session = Depends(get_session),
) -> dict:
    """Contexto do lote para o veterinário formular a dieta: nome/nº do lote,
    nº de animais, última dieta (produtos e qtd/cabeça/dia), último controle
    leiteiro de cada animal e um resumo (DEL médio, média do CL, data do CL)."""
    fazenda_id = fazenda_id_seguro(fazenda_id)
    query_lote = select(Lote).where(Lote.codigo == f"{lote:02d}")
    if fazenda_id is not None:
        query_lote = query_lote.where(Lote.fazenda_id == fazenda_id)
    lote_cad = session.exec(query_lote).first()
    animais = _animais_do_lote(session, lote, fazenda_id)
    n = len(animais)

    # Último controle leiteiro AO VIVO (tabela controle_leiteiro), com
    # fallback ao campo congelado só para quem nunca teve controle lançado
    # pelo app — ver rules/producao_leiteira.py. Sem isso, lançar um
    # controle novo não refletia aqui (a tela continuava mostrando a
    # produção/data do último CSV importado, potencialmente meses velha).
    controles_ao_vivo = ultimo_controle_por_animal(session, {a.numero for a in animais}, fazenda_id)
    producao_data_por_animal = {a.numero: com_fallback_animal(a.numero, controles_ao_vivo, a) for a in animais}

    dels = [a.del_dias for a in animais if a.del_dias is not None]
    cls = [p for p, _ in producao_data_por_animal.values() if p is not None]
    datas_cl = [d for _, d in producao_data_por_animal.values() if d is not None]

    # Última dieta ativa do lote (produtos + qtd total/dia → por cabeça/dia).
    query_ativa = select(DietaLancamento).where(
        DietaLancamento.lote == lote, DietaLancamento.data_efetivo_encerramento == None  # noqa: E711
    )
    if fazenda_id is not None:
        query_ativa = query_ativa.where(DietaLancamento.fazenda_id == fazenda_id)
    ativa = session.exec(query_ativa).first()
    ultima_dieta = None
    if ativa:
        query_itens = select(DietaItemProgramado).where(DietaItemProgramado.dieta_lancamento_id == ativa.id)
        if fazenda_id is not None:
            query_itens = query_itens.where(DietaItemProgramado.fazenda_id == fazenda_id)
        itens = session.exec(query_itens).all()
        ultima_dieta = {
            "data_abertura": ativa.data_abertura.isoformat(),
            "data_prevista_encerramento": ativa.data_prevista_encerramento.isoformat() if ativa.data_prevista_encerramento else None,
            "responsavel": ativa.responsavel,
            "base_quantidade": ativa.base_quantidade,
            "leite_bezerros_kg_dia": ativa.leite_bezerros_kg_dia,
            # Leite/bezerro — mesma divisão simples que `por_cabeca` já faz
            # pros itens da dieta abaixo (total do lote / nº de animais do
            # lote); só faz sentido exibir num lote de bezerras.
            "leite_por_bezerro_kg_dia": (
                round(ativa.leite_bezerros_kg_dia / n, 2) if ativa.leite_bezerros_kg_dia and n else None
            ),
            "itens": [
                _linha_item_dieta(it, ativa.base_quantidade, n)
                for it in itens
            ],
        }

    return {
        "lote": lote,
        "nome": lote_cad.nome if lote_cad else None,
        "qtd_animais": n,
        "del_medio": round(sum(dels) / len(dels)) if dels else None,
        "media_cl": round(sum(cls) / len(cls), 1) if cls else None,
        "data_ult_cl": max(datas_cl).isoformat() if datas_cl else None,
        "animais": sorted([
            {
                "numero": a.numero, "del_dias": a.del_dias,
                "ult_cl_kg": producao_data_por_animal[a.numero][0],
                "data_ult_leite": producao_data_por_animal[a.numero][1].isoformat() if producao_data_por_animal[a.numero][1] else None,
            }
            for a in animais
        ], key=lambda x: (x["ult_cl_kg"] is None, -(x["ult_cl_kg"] or 0))),
        "ultima_dieta": ultima_dieta,
    }


@router.get("/dietas/{dieta_id}/apresentacao")
def apresentacao_dieta(
    dieta_id: int, fazenda_id: int | None = Depends(get_fazenda_atual_id), session: Session = Depends(get_session),
) -> dict:
    """Como o funcionário vê a dieta para conferir no vagão: por produto — qtd/
    cabeça, total/dia, total/trato; e o somatório de kg no vagão do lote."""
    fazenda_id = fazenda_id_seguro(fazenda_id)
    dieta = session.get(DietaLancamento, dieta_id)
    # BUG DE SEGURANÇA CORRIGIDO: o cabeçalho da dieta (lote, datas) de outra
    # fazenda vazava mesmo com os itens já corretamente filtrados abaixo.
    if not dieta or (fazenda_id is not None and dieta.fazenda_id != fazenda_id):
        raise HTTPException(status_code=404, detail="Dieta não encontrada")
    n = len(_animais_do_lote(session, dieta.lote, fazenda_id))
    query_lote = select(Lote).where(Lote.codigo == f"{dieta.lote:02d}")
    if fazenda_id is not None:
        query_lote = query_lote.where(Lote.fazenda_id == fazenda_id)
    lote_cad = session.exec(query_lote).first()
    query_itens = select(DietaItemProgramado).where(DietaItemProgramado.dieta_lancamento_id == dieta_id)
    if fazenda_id is not None:
        query_itens = query_itens.where(DietaItemProgramado.fazenda_id == fazenda_id)
    itens = session.exec(query_itens).all()
    linhas = []
    total_dia = 0.0
    for it in itens:
        # Total do lote/dia deste item, respeitando a base EFETIVA (override
        # do item, senão a da dieta) — não o valor de `it.quantidade` cru, que
        # só é o total do lote quando a base é "total" (ver `_linha_item_dieta`
        # e `_totais_item`; era aqui que a conta ficava errada por um fator do
        # tamanho do lote quando a dieta era lançada "por animal").
        linha = _linha_item_dieta(it, dieta.base_quantidade, n)
        total_lote_item = linha["total_dia"]
        linha["total_trato"] = round(total_lote_item / NUM_TRATOS, 2) if total_lote_item is not None else None
        linhas.append(linha)
        if (it.unidade or "").lower() in ("kg", "g") and total_lote_item is not None:
            total_dia += total_lote_item
    return {
        "lote": dieta.lote, "nome": lote_cad.nome if lote_cad else None, "qtd_animais": n,
        "data_abertura": dieta.data_abertura.isoformat(),
        "data_prevista_encerramento": dieta.data_prevista_encerramento.isoformat() if dieta.data_prevista_encerramento else None,
        "num_tratos": NUM_TRATOS, "itens": linhas,
        # Somatório de quilos que deve estar no vagão (só itens em kg).
        "vagao_kg_dia": round(total_dia, 2),
        "vagao_kg_trato": round(total_dia / NUM_TRATOS, 2),
    }


class EncerrarDietaIn(BaseModel):
    data_efetivo_encerramento: date


@router.put("/dietas/{dieta_id}/encerrar")
def encerrar_dieta(
    dieta_id: int, dados: EncerrarDietaIn, session: Session = Depends(get_session),
    fazenda_id: int | None = Depends(get_fazenda_atual_id),
) -> dict:
    fazenda_id = fazenda_id_seguro(fazenda_id)
    dieta = session.get(DietaLancamento, dieta_id)
    # BUG DE SEGURANÇA CORRIGIDO: sem esta checagem, qualquer usuário
    # autenticado podia encerrar a dieta ATIVA de outra fazenda só
    # adivinhando dieta_id.
    if not dieta or (fazenda_id is not None and dieta.fazenda_id != fazenda_id):
        raise HTTPException(status_code=404, detail="Dieta não encontrada")
    if dieta.data_efetivo_encerramento is not None:
        raise HTTPException(status_code=400, detail="Esta dieta já está encerrada")
    dieta.data_efetivo_encerramento = dados.data_efetivo_encerramento
    session.add(dieta)
    session.commit()
    return {"encerrada": True, "lote": dieta.lote}


class RegistroRealIn(BaseModel):
    data: date
    itens: list[ItemProgramadoIn]


@router.post("/dietas/{dieta_id}/real", status_code=201)
def registrar_real(
    dieta_id: int, dados: RegistroRealIn,
    fazenda_id: int = Depends(get_fazenda_id_escrita),
    session: Session = Depends(get_session), user: Usuario = Depends(get_current_user),
) -> dict:
    dieta = session.get(DietaLancamento, dieta_id)
    # BUG DE SEGURANÇA CORRIGIDO: sem esta checagem, qualquer usuário podia
    # anexar registros de "consumo real" à dieta de outra fazenda.
    if not dieta or (fazenda_id is not None and dieta.fazenda_id != fazenda_id):
        raise HTTPException(status_code=404, detail="Dieta não encontrada")
    if not dados.itens:
        raise HTTPException(status_code=400, detail="Informe ao menos um alimento oferecido")
    for item in dados.itens:
        session.add(DietaRegistroReal(
            dieta_lancamento_id=dieta_id, data=dados.data, usuario_id=user.id, fazenda_id=fazenda_id,
            **item.model_dump(),
        ))
    session.commit()
    return {"registrados": len(dados.itens)}


@router.get("/dietas/{dieta_id}/comparativo")
def comparativo_dieta(
    dieta_id: int, fazenda_id: int | None = Depends(get_fazenda_atual_id), session: Session = Depends(get_session),
) -> dict:
    """Programado × real por alimento — soma total real e média por dia distinto registrado."""
    fazenda_id = fazenda_id_seguro(fazenda_id)
    dieta = session.get(DietaLancamento, dieta_id)
    if not dieta or (fazenda_id is not None and dieta.fazenda_id != fazenda_id):
        raise HTTPException(status_code=404, detail="Dieta não encontrada")
    query_prog = select(DietaItemProgramado).where(DietaItemProgramado.dieta_lancamento_id == dieta_id)
    query_reais = select(DietaRegistroReal).where(DietaRegistroReal.dieta_lancamento_id == dieta_id)
    if fazenda_id is not None:
        query_prog = query_prog.where(DietaItemProgramado.fazenda_id == fazenda_id)
        query_reais = query_reais.where(DietaRegistroReal.fazenda_id == fazenda_id)
    programados = session.exec(query_prog).all()
    reais = session.exec(query_reais).all()

    por_alimento: dict[str, dict] = {}
    for p in programados:
        acc = por_alimento.setdefault(p.alimento, {"alimento": p.alimento, "unidade": p.unidade, "programado": 0.0, "real_total": 0.0, "real_dias": 0, "real_media_dia": None})
        acc["programado"] += p.quantidade

    dias_por_alimento: dict[str, set] = {}
    for r in reais:
        acc = por_alimento.setdefault(r.alimento, {"alimento": r.alimento, "unidade": r.unidade, "programado": 0.0, "real_total": 0.0, "real_dias": 0, "real_media_dia": None})
        acc["real_total"] = round(acc["real_total"] + r.quantidade, 2)
        dias_por_alimento.setdefault(r.alimento, set()).add(r.data.isoformat())

    for alimento, dias in dias_por_alimento.items():
        por_alimento[alimento]["real_dias"] = len(dias)
        por_alimento[alimento]["real_media_dia"] = round(por_alimento[alimento]["real_total"] / len(dias), 2) if dias else None

    return {"dieta": _serializar_dieta(session, dieta, fazenda_id), "itens": sorted(por_alimento.values(), key=lambda x: x["alimento"])}


# ---------------------------------------------------------------------------
# Consumo diário e sobra de cocho (Lançamentos > Alimentação) — o funcionário
# lança o que foi de fato FORNECIDO a cada lote (dá baixa em estoque, ver
# `rules/estoque_baixa.movimentar`) e, num card separado, a sobra do cocho em
# kg totais (só medição — nunca baixa nada). Distinto do bloco de
# "Lançamento de dieta" acima: aquele é o PLANO (nutricionista), este é o
# REALIZADO físico que efetivamente sai do silo/depósito.
# ---------------------------------------------------------------------------
def _dieta_ativa_do_lote(session: Session, lote: int, fazenda_id: int | None) -> DietaLancamento | None:
    """A dieta em vigor de um lote agora — mesma resolução (`data_efetivo_
    encerramento IS NULL`) já usada por `contexto_dieta`/`apresentacao_dieta`.
    Só uma pode estar ativa por lote, então não há ambiguidade de qual usar."""
    query = select(DietaLancamento).where(
        DietaLancamento.lote == lote, DietaLancamento.data_efetivo_encerramento == None,  # noqa: E711
    )
    if fazenda_id is not None:
        query = query.where(DietaLancamento.fazenda_id == fazenda_id)
    return session.exec(query).first()


def _itens_dieta_lancamento(session: Session, dieta_lancamento_id: int, fazenda_id: int | None) -> list[DietaItemProgramado]:
    query = select(DietaItemProgramado).where(DietaItemProgramado.dieta_lancamento_id == dieta_lancamento_id)
    if fazenda_id is not None:
        query = query.where(DietaItemProgramado.fazenda_id == fazenda_id)
    return session.exec(query).all()


def _item_da_dieta(itens: list[DietaItemProgramado], alimento: str, alimento_id: int | None) -> DietaItemProgramado | None:
    """Casa um alimento lançado com o item programado da dieta ativa — por
    `alimento_id` (vínculo de cadastro, preferido, sem ambiguidade de nome) e,
    na falta dele, por nome (trim + minúsculas, mesmo padrão usado em todo o
    módulo para casar Alimento×Estoque). None = fora da dieta (ver B5)."""
    if alimento_id is not None:
        for it in itens:
            if it.alimento_id == alimento_id:
                return it
    alvo = (alimento or "").strip().lower()
    for it in itens:
        if (it.alimento or "").strip().lower() == alvo:
            return it
    return None


def _base_efetiva(item_base: str | None, dieta_base: str | None) -> str:
    """Resolve qual base ("total" ou "animal") vale para UM item: o override
    do próprio item, se informado, senão a base da dieta, senão "total" — o
    legado de todo lançamento anterior a este override existir. Nunca ler
    `DietaItemProgramado.quantidade` sem passar antes por aqui: é o que
    permite misturar bases dentro da mesma dieta (silagem em total do lote,
    concentrado por cabeça) sem quebrar quem nunca usou a base por item."""
    return item_base or dieta_base or "total"


def _por_cabeca(qtd_fisica: float, base_quantidade: str | None, n_animais: int) -> float | None:
    """Quantidade por cabeça de um item programado, dada a base EFETIVA já
    resolvida (ver `_base_efetiva`) — a causa mais provável de dobrar a conta
    (ver spec da sessão): quando o item já foi lançado "por animal", o
    `quantidade` do item JÁ É por cabeça (não dividir de novo); quando foi
    lançado "total" (padrão), é o total do lote e só vira por-cabeça dividindo
    pelo efetivo atual. `None` quando não há efetivo para dividir — melhor não
    responder do que inventar um valor com denominador zero."""
    if base_quantidade == "animal":
        return qtd_fisica
    return qtd_fisica / n_animais if n_animais else None


def _totais_item(qtd_fisica: float, base_efetiva: str, n_animais: int) -> tuple[float | None, float | None]:
    """Retorna (total_lote_dia, por_cabeca_dia) de um item a partir da sua
    quantidade física já convertida (MN) e da base EFETIVA (ver
    `_base_efetiva`) — um dos dois vem direto do valor lançado, o outro é
    derivado multiplicando/dividindo por `n_animais`. `None` no lado derivado
    quando não há efetivo para multiplicar/dividir."""
    if base_efetiva == "animal":
        return (qtd_fisica * n_animais if n_animais else None), qtd_fisica
    return qtd_fisica, (qtd_fisica / n_animais if n_animais else None)


def _linha_item_dieta(it: DietaItemProgramado, dieta_base_quantidade: str | None, n_animais: int) -> dict:
    """Um item de dieta pronto pra exibição (total/dia + por cabeça/dia), já
    resolvendo a base EFETIVA (override do item, senão da dieta, senão
    "total") — usado por `contexto_dieta` e `apresentacao_dieta` para nunca
    tratar `it.quantidade` como total do lote sem checar a base."""
    qtd_fisica = _quantidade_fisica(it.quantidade, it.unidade, it.base, it.ms_pct)
    base_efetiva = _base_efetiva(it.base_quantidade, dieta_base_quantidade)
    total_dia, por_cabeca = _totais_item(qtd_fisica, base_efetiva, n_animais)
    return {
        "alimento": it.alimento, "unidade": it.unidade,
        "total_dia": round(total_dia, 2) if total_dia is not None else None,
        "por_cabeca": round(por_cabeca, 3) if por_cabeca is not None else None,
    }


def _resolver_estoque_item(
    session: Session, fazenda_id: int | None, alimento: str, estoque_por_alimento: dict[str, list[dict]],
) -> Estoque | None:
    """Mesma resolução de `_dar_baixa_automatica`: nome do alimento bate
    direto com `Estoque.nome` primeiro; na falta, cai no vínculo por Alimento
    (`Estoque.alimento_id`). Reaproveitada aqui em vez de duplicada porque o
    consumo manual precisa resolver o MESMO item que a baixa automática
    resolveria para o mesmo alimento."""
    query = select(Estoque).where(Estoque.nome == alimento)
    if fazenda_id is not None:
        query = query.where(Estoque.fazenda_id == fazenda_id)
    item = session.exec(query).first()
    if item:
        return item
    candidatos = estoque_por_alimento.get((alimento or "").strip().lower())
    if candidatos:
        return session.get(Estoque, candidatos[0]["id"])
    return None


def _kg_fornecido_do_dia(session: Session, lote: int, data: date, fazenda_id: int | None) -> float:
    """Soma em kg (via `kg_equivalente`) de tudo que foi lançado como consumo
    de um lote num dia — o denominador do percentual de sobra (B14). Itens
    cuja unidade não converte (litro, dose...) ficam fora da soma: um `None`
    tratado como zero inflaria a sobra artificialmente (ver docstring de
    `rules/unidades`)."""
    query = select(ConsumoAlimento).where(ConsumoAlimento.lote == lote, ConsumoAlimento.data == data)
    if fazenda_id is not None:
        query = query.where(ConsumoAlimento.fazenda_id == fazenda_id)
    total = 0.0
    for r in session.exec(query).all():
        kg = kg_equivalente(r.quantidade, r.unidade)
        if kg is not None:
            total += kg
    return round(total, 4)


def _faixa_sobra(sobra_pct: float | None) -> bool | None:
    """Se o percentual de sobra está dentro da faixa aceitável configurada em
    Parâmetros (`sobra_min_pct`..`sobra_max_pct`). `None` (sem sobra lançada
    ainda) se propaga — não há faixa a avaliar sem medição."""
    if sobra_pct is None:
        return None
    minimo = get_param("sobra_min_pct", 3) or 3
    maximo = get_param("sobra_max_pct", 7) or 7
    return minimo <= sobra_pct <= maximo


@router.get("/consumo/dieta-do-lote")
def dieta_do_lote_consumo(
    lote: int, fazenda_id: int | None = Depends(get_fazenda_atual_id), session: Session = Depends(get_session),
) -> dict:
    """Alimentos disponíveis para o lançamento de consumo (B3): só os da
    dieta ATIVA do lote — quantidade por cabeça já resolvida (considerando
    `base_quantidade`) para a tela não precisar refazer essa conta (e correr
    o risco de dobrá-la)."""
    fazenda_id = fazenda_id_seguro(fazenda_id)
    dieta = _dieta_ativa_do_lote(session, lote, fazenda_id)
    if not dieta:
        raise HTTPException(status_code=404, detail=f"Lote {lote:02d} não tem dieta ativa")
    n = len(_animais_do_lote(session, lote, fazenda_id))
    itens = _itens_dieta_lancamento(session, dieta.id, fazenda_id)
    linhas = []
    for it in itens:
        qtd_fisica = _quantidade_fisica(it.quantidade, it.unidade, it.base, it.ms_pct)
        base_efetiva = _base_efetiva(it.base_quantidade, dieta.base_quantidade)
        por_cabeca = _por_cabeca(qtd_fisica, base_efetiva, n)
        linhas.append({
            "alimento": it.alimento, "alimento_id": it.alimento_id,
            "quantidade": it.quantidade, "unidade": it.unidade,
            "por_cabeca": round(por_cabeca, 4) if por_cabeca is not None else None,
            "converte_para_kg": converte_para_kg(it.unidade),
        })
    return {"itens": linhas, "base_quantidade": dieta.base_quantidade or "total"}


@router.get("/consumo")
def obter_consumo(
    lote: int, data: date,
    fazenda_id: int | None = Depends(get_fazenda_atual_id), session: Session = Depends(get_session),
) -> dict:
    """O que já foi lançado num lote num dia, somado por alimento (B2) —
    lançamentos do mesmo dia SOMAM (o vagão passa mais de uma vez), então a
    tela precisa ver o acumulado, não a lista de eventos crus."""
    fazenda_id = fazenda_id_seguro(fazenda_id)
    query = select(ConsumoAlimento).where(ConsumoAlimento.lote == lote, ConsumoAlimento.data == data)
    if fazenda_id is not None:
        query = query.where(ConsumoAlimento.fazenda_id == fazenda_id)
    registros = session.exec(query.order_by(ConsumoAlimento.criado_em)).all()

    por_alimento: dict[str, dict] = {}
    for r in registros:
        # Agrupado só por nome do alimento — na prática todo lançamento do
        # mesmo alimento no mesmo lote/dia vem do mesmo item de dieta (mesma
        # unidade); se dois lançamentos discordarem de unidade num caso raro,
        # a soma bruta ainda reflete "quanto foi lançado", só o `kg_equivalente`
        # fica impreciso — aceitável porque o valor em kg de cada linha
        # individual nunca é descartado (ver GET /consumo/dieta-do-lote).
        acc = por_alimento.setdefault(r.alimento, {"alimento": r.alimento, "quantidade": 0.0, "unidade": r.unidade})
        acc["quantidade"] = round(acc["quantidade"] + r.quantidade, 4)

    itens = []
    kg_total = 0.0
    for acc in por_alimento.values():
        kg = kg_equivalente(acc["quantidade"], acc["unidade"])
        if kg is not None:
            kg_total += kg
        itens.append({**acc, "kg_equivalente": round(kg, 4) if kg is not None else None})
    kg_total = round(kg_total, 4)

    # Nº de animais do lançamento mais recente do dia — guardado por
    # auditoria em todo registro (mesmo em modo kg direto, ver docstring de
    # `ConsumoAlimento.num_animais`), então o mais recente é o retrato mais
    # atual do efetivo que passou no cocho hoje.
    num_animais = registros[-1].num_animais if registros else None

    query_sobra = select(ConsumoSobra).where(ConsumoSobra.lote == lote, ConsumoSobra.data == data)
    if fazenda_id is not None:
        query_sobra = query_sobra.where(ConsumoSobra.fazenda_id == fazenda_id)
    sobra = session.exec(query_sobra).first()
    sobra_kg = sobra.kg_sobra if sobra else None
    sobra_pct = round(sobra_kg / kg_total * 100, 2) if sobra_kg is not None and kg_total > 0 else None

    return {
        "lote": lote, "data": data.isoformat(), "num_animais": num_animais,
        "itens": sorted(itens, key=lambda x: x["alimento"]),
        "kg_fornecido_total": kg_total,
        "sobra_kg": sobra_kg, "sobra_pct": sobra_pct, "dentro_da_faixa": _faixa_sobra(sobra_pct),
    }


class ConsumoItemIn(BaseModel):
    alimento: str
    alimento_id: int | None = None
    quantidade: float
    unidade: str


class ConsumoIn(BaseModel):
    lote: int
    data: date
    num_animais: int | None = None
    origem: str  # "animais" (por cabeça × dieta) | "kg" (digitado direto)
    itens: list[ConsumoItemIn]


@router.post("/consumo", status_code=201)
def lancar_consumo(
    dados: ConsumoIn, fazenda_id: int = Depends(get_fazenda_id_escrita),
    session: Session = Depends(get_session), user: Usuario = Depends(get_current_user),
) -> dict:
    """B1/B4/B5/B6/B7/B8: grava um ou mais ConsumoAlimento de um lote num dia,
    dá baixa em estoque pelo motor único e propaga os avisos da baixa."""
    if dados.origem not in ("animais", "kg"):
        raise HTTPException(status_code=400, detail='"origem" deve ser "animais" ou "kg"')
    if not dados.itens:
        raise HTTPException(status_code=400, detail="Informe ao menos um alimento")
    if dados.origem == "animais" and not dados.num_animais:
        raise HTTPException(status_code=400, detail='Informe "num_animais" para lançar por cabeça')

    query_lote = select(Lote).where(Lote.codigo == f"{dados.lote:02d}")
    if fazenda_id is not None:
        query_lote = query_lote.where(Lote.fazenda_id == fazenda_id)
    lote_cad = session.exec(query_lote).first()
    # Lote sem cadastro (não deveria acontecer numa dieta lançada, mas o
    # sistema convive com dado legado — ver "O que o levantamento achou" da
    # spec) cai no padrão restritivo das duas flags, igual a um Lote com as
    # flags nunca marcadas: a ausência de cadastro nunca é motivo pra afrouxar
    # uma checagem de segurança.
    permitir_fora_da_dieta = bool(lote_cad.permitir_fora_da_dieta) if lote_cad else False
    permitir_sem_estoque = bool(lote_cad.permitir_sem_estoque) if lote_cad else False
    # Mesma postura: sem cadastro, cai no padrão "consumo_real" — o único
    # modo em que este endpoint sempre debitou estoque, então um lote legado
    # continua se comportando exatamente como antes desta coluna existir.
    modo_baixa_estoque = (lote_cad.modo_baixa_estoque if lote_cad else None) or "consumo_real"
    deduzir_estoque = modo_baixa_estoque == "consumo_real"

    dieta = _dieta_ativa_do_lote(session, dados.lote, fazenda_id)
    itens_dieta = _itens_dieta_lancamento(session, dieta.id, fazenda_id) if dieta else []
    n_animais = len(_animais_do_lote(session, dados.lote, fazenda_id))
    estoque_por_alimento, _ = _estoque_por_alimento(session, fazenda_id)

    avisos: list[str] = []
    for item_in in dados.itens:
        item_dieta = _item_da_dieta(itens_dieta, item_in.alimento, item_in.alimento_id)
        fora_da_dieta = item_dieta is None
        if fora_da_dieta and not permitir_fora_da_dieta:
            raise HTTPException(
                status_code=409,
                detail=f'"{item_in.alimento}" não está na dieta ativa do lote {dados.lote:02d} — ligue '
                       f'"permitir alimentos fora da dieta" no cadastro do lote para lançar mesmo assim.',
            )

        if dados.origem == "animais" and item_dieta is not None:
            # Recalculado no servidor — não confia no valor que o front
            # mandou. É exatamente aqui que a conta dobra se a base EFETIVA
            # (override do item, senão a da dieta) for ignorada (ver docstring
            # de `_por_cabeca`).
            qtd_fisica = _quantidade_fisica(item_dieta.quantidade, item_dieta.unidade, item_dieta.base, item_dieta.ms_pct)
            base_efetiva = _base_efetiva(item_dieta.base_quantidade, dieta.base_quantidade if dieta else None)
            por_cabeca = _por_cabeca(qtd_fisica, base_efetiva, n_animais)
            if por_cabeca is None:
                raise HTTPException(
                    status_code=400,
                    detail=f'Não foi possível calcular a quantidade por cabeça de "{item_in.alimento}" '
                           f'— lote sem animais ativos.',
                )
            quantidade = round(por_cabeca * dados.num_animais, 4)
            unidade = item_dieta.unidade
        else:
            # "kg" direto, ou item fora da dieta (sem quantidade/cabeça
            # conhecida para multiplicar, mesmo que o modo geral seja "animais").
            quantidade = item_in.quantidade
            unidade = item_in.unidade

        estoque_item = _resolver_estoque_item(session, fazenda_id, item_in.alimento, estoque_por_alimento)
        sem_saldo = estoque_item is None or (estoque_item.quantidade or 0) <= 0
        if sem_saldo and not permitir_sem_estoque:
            raise HTTPException(
                status_code=409,
                detail=f'"{item_in.alimento}" está sem saldo em estoque — ligue "permitir alimentos sem '
                       f'estoque" no cadastro do lote para lançar mesmo assim.',
            )

        registro = ConsumoAlimento(
            fazenda_id=fazenda_id, data=dados.data, lote=dados.lote,
            alimento=item_in.alimento, alimento_id=item_in.alimento_id,
            quantidade=quantidade, unidade=unidade, num_animais=dados.num_animais,
            origem=dados.origem, fora_da_dieta=fora_da_dieta, usuario_id=user.id,
            baixou_estoque=deduzir_estoque,
        )
        session.add(registro)
        session.flush()  # gera o id antes do movimento de estoque (origem_id rastreável — B7)
        # "automatica"/"sem_baixa" registram o consumo (sobra/histórico
        # continuam funcionando) mas NÃO tocam o Estoque aqui — "automatica"
        # já é debitada pelo mecanismo de dias decorridos (`_dar_baixa_automatica`)
        # e dobraria a baixa; "sem_baixa" é só plano/receita, por decisão do lote.
        if deduzir_estoque:
            avisos.extend(estoque_baixa.baixar(
                session, item=estoque_item, quantidade=quantidade, unidade=unidade, data=dados.data,
                fazenda_id=fazenda_id,
                observacao=f"Consumo diário — lote {dados.lote:02d}, {item_in.alimento}",
                usuario_id=user.id, origem_tipo="consumo_alimento", origem_id=registro.id, produto=item_in.alimento,
            ))

    session.commit()
    return {"ok": True, "avisos": avisos}


@router.delete("/consumo/{consumo_id}")
def excluir_consumo(
    consumo_id: int, fazenda_id: int | None = Depends(get_fazenda_atual_id),
    session: Session = Depends(get_session), user: Usuario = Depends(get_current_user),
) -> dict:
    """B9: exclui um lançamento de consumo e devolve o produto ao estoque
    pelo mesmo motor — o estorno espelha a baixa, não uma baixa negativa
    inventada na mão.

    Só devolve quando `baixou_estoque` é True: um registro lançado com o lote
    em modo "automatica"/"sem_baixa" nunca tocou o Estoque, então devolver
    aqui inventaria estoque que nunca saiu. Olha o FLAG DO REGISTRO, não o
    modo ATUAL do lote — o modo pode ter mudado depois do lançamento, e o
    estorno tem de espelhar o que aconteceu quando o consumo foi lançado."""
    fazenda_id = fazenda_id_seguro(fazenda_id)
    registro = session.get(ConsumoAlimento, consumo_id)
    if not registro or (fazenda_id is not None and registro.fazenda_id != fazenda_id):
        raise HTTPException(status_code=404, detail="Lançamento de consumo não encontrado")
    avisos: list[str] = []
    if registro.baixou_estoque:
        estoque_por_alimento, _ = _estoque_por_alimento(session, fazenda_id)
        estoque_item = _resolver_estoque_item(session, fazenda_id, registro.alimento, estoque_por_alimento)
        avisos = estoque_baixa.devolver(
            session, item=estoque_item, quantidade=registro.quantidade, unidade=registro.unidade,
            data=registro.data, fazenda_id=fazenda_id,
            observacao=f"Exclusão do consumo diário — lote {registro.lote:02d}, {registro.alimento}",
            usuario_id=user.id, origem_tipo="consumo_alimento", origem_id=registro.id, produto=registro.alimento,
        )
    session.delete(registro)
    session.commit()
    return {"ok": True, "avisos": avisos}


class SobraIn(BaseModel):
    lote: int
    data: date
    kg_sobra: float


@router.post("/sobra", status_code=201)
def lancar_sobra(
    dados: SobraIn, fazenda_id: int = Depends(get_fazenda_id_escrita),
    session: Session = Depends(get_session), user: Usuario = Depends(get_current_user),
) -> dict:
    """B10: grava a sobra do lote no dia, em kg totais. Relançar no mesmo dia
    SUBSTITUI (upsert por `fazenda_id, lote, data` — a UniqueConstraint do
    modelo trava isso no banco), ao contrário do consumo, que soma: a sobra é
    uma medição única do estado do cocho, não um acúmulo de eventos."""
    if dados.kg_sobra < 0:
        raise HTTPException(status_code=400, detail="Sobra não pode ser negativa")
    query = select(ConsumoSobra).where(ConsumoSobra.lote == dados.lote, ConsumoSobra.data == dados.data)
    if fazenda_id is not None:
        query = query.where(ConsumoSobra.fazenda_id == fazenda_id)
    registro = session.exec(query).first()
    if registro:
        registro.kg_sobra = dados.kg_sobra
        registro.usuario_id = user.id
        registro.atualizado_em = datetime.utcnow()
    else:
        registro = ConsumoSobra(
            fazenda_id=fazenda_id, data=dados.data, lote=dados.lote, kg_sobra=dados.kg_sobra, usuario_id=user.id,
        )
    session.add(registro)
    session.commit()
    session.refresh(registro)

    kg_total = _kg_fornecido_do_dia(session, dados.lote, dados.data, fazenda_id)
    sobra_pct = round(dados.kg_sobra / kg_total * 100, 2) if kg_total > 0 else None
    return {
        "ok": True, "kg_sobra": registro.kg_sobra,
        "sobra_pct": sobra_pct, "dentro_da_faixa": _faixa_sobra(sobra_pct),
    }


@router.get("/sobra/relatorio")
def relatorio_sobra(
    de: date, ate: date, lote: int | None = None,
    fazenda_id: int | None = Depends(get_fazenda_atual_id), session: Session = Depends(get_session),
) -> dict:
    """B11-B14: sobra por alimento e total no período. O rateio por alimento
    usa a proporção de cada alimento NA DIETA (não no consumo realmente
    lançado — B12): o relatório serve também para quem só pesa a sobra e
    confia na dieta cadastrada pra saber o que ela era feita de. A dieta usada
    é a ATIVA de cada lote hoje — mesma simplificação que `apresentacao_dieta`/
    `contexto_dieta` já fazem em todo o arquivo (nenhum dos dois resolve a
    dieta vigente numa data passada); um lote que trocou de dieta dentro do
    período aparece com a proporção da dieta atual."""
    fazenda_id = fazenda_id_seguro(fazenda_id)
    query_sobra = select(ConsumoSobra).where(ConsumoSobra.data >= de, ConsumoSobra.data <= ate)
    if lote is not None:
        query_sobra = query_sobra.where(ConsumoSobra.lote == lote)
    if fazenda_id is not None:
        query_sobra = query_sobra.where(ConsumoSobra.fazenda_id == fazenda_id)
    sobras = session.exec(query_sobra).all()

    query_consumo = select(ConsumoAlimento).where(ConsumoAlimento.data >= de, ConsumoAlimento.data <= ate)
    if lote is not None:
        query_consumo = query_consumo.where(ConsumoAlimento.lote == lote)
    if fazenda_id is not None:
        query_consumo = query_consumo.where(ConsumoAlimento.fazenda_id == fazenda_id)
    consumos = session.exec(query_consumo).all()

    total_kg_sobra = round(sum(s.kg_sobra for s in sobras), 2)

    # Fornecido total do período (B14) — convertido igual ao denominador do
    # percentual por dia (GET /consumo): itens sem conversão ficam fora da
    # soma e entram na lista de excluídos, nunca viram zero.
    itens_sem_conversao: set[str] = set()
    total_kg_fornecido = 0.0
    for c in consumos:
        kg = kg_equivalente(c.quantidade, c.unidade)
        if kg is None:
            itens_sem_conversao.add(c.alimento)
            continue
        total_kg_fornecido += kg
    total_kg_fornecido = round(total_kg_fornecido, 2)
    pct_medio = round(total_kg_sobra / total_kg_fornecido * 100, 2) if total_kg_fornecido > 0 else None

    # Rateio por alimento (B12/B13): para cada dia de sobra lançada, distribui
    # o kg total daquele dia entre os alimentos da dieta ATIVA do lote,
    # proporcional ao peso (kg) de cada um. Um lote sem dieta, ou cuja dieta
    # não tem NENHUM item conversível, não entra no rateio (B13) — o total
    # continua contado, só não tem "por alimento" pra aquele lote/dia.
    dieta_itens_cache: dict[int, list[DietaItemProgramado]] = {}
    acumulado: dict[str, float] = {}
    kg_sobra_rateado = 0.0
    for s in sobras:
        if s.lote not in dieta_itens_cache:
            dieta = _dieta_ativa_do_lote(session, s.lote, fazenda_id)
            dieta_itens_cache[s.lote] = _itens_dieta_lancamento(session, dieta.id, fazenda_id) if dieta else []
        itens = dieta_itens_cache[s.lote]

        pesos: dict[str, float] = {}
        soma_pesos = 0.0
        for it in itens:
            qtd_fisica = _quantidade_fisica(it.quantidade, it.unidade, it.base, it.ms_pct)
            kg = kg_equivalente(qtd_fisica, it.unidade)
            if kg is None:
                itens_sem_conversao.add(it.alimento)
                continue
            pesos[it.alimento] = pesos.get(it.alimento, 0.0) + kg
            soma_pesos += kg
        if soma_pesos <= 0:
            continue  # B13 — nada conversível nesta dieta, sem rateio possível pra este dia
        for alimento, kg in pesos.items():
            acumulado[alimento] = acumulado.get(alimento, 0.0) + s.kg_sobra * (kg / soma_pesos)
        kg_sobra_rateado += s.kg_sobra

    por_alimento = [
        {
            "alimento": alimento, "kg_sobra": round(kg, 2),
            # Fatia do total de sobra RATEADA no período que este alimento
            # respondeu — não a proporção física dele na dieta (essa já foi
            # usada internamente, dia a dia, pra fazer o rateio acima; expor
            # ela agregada exigiria normalizar dietas diferentes ao longo do
            # período, o que essa única métrica não consegue carregar sem
            # enganar). É a pergunta que o relatório existe pra responder:
            # de onde veio a sobra. O campo se chamava `pct_da_dieta`, nome
            # que prometia outra coisa — quem fosse montar a tela leria
            # "% da dieta" e exibiria o número errado com toda a confiança.
            "pct_do_total": round(kg / kg_sobra_rateado * 100, 2) if kg_sobra_rateado > 0 else 0.0,
        }
        for alimento, kg in sorted(acumulado.items())
    ]

    return {
        "total_kg_sobra": total_kg_sobra, "total_kg_fornecido": total_kg_fornecido, "pct_medio": pct_medio,
        "por_alimento": por_alimento, "itens_sem_conversao": sorted(itens_sem_conversao),
    }
