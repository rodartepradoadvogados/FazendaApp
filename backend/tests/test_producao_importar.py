"""
Testes de importação de planilha (Excel/.xlsx) de controle leiteiro (por
animal ou por lote, identificado sozinho pelo cabeçalho) e de qualidade do
leite, atalho embutido na tela de Lançamentos.
"""
from __future__ import annotations

import io

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy.pool import StaticPool
from sqlmodel import Session, SQLModel, create_engine

import fazenda.database as database
from fazenda.models import Animal, Lote


def _xlsx_bytes(header: list[str], rows: list[list]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(header)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
def client():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    SQLModel.metadata.create_all(engine)

    def _get_session_override():
        with Session(engine) as session:
            yield session

    import main
    from fazenda.auth import get_current_user

    class _FakeUser:
        id = 1
        papel = "admin"
        ativo = True
        username = "teste"

    main.app.dependency_overrides[database.get_session] = _get_session_override
    main.app.dependency_overrides[get_current_user] = lambda: _FakeUser()

    with TestClient(main.app) as c:
        with Session(engine) as s:
            s.add(Lote(codigo="01", nome="Alta"))
            s.add(Animal(numero="101", grupo_primario="01 - Alta", raca="Girolando", del_dias=50, ativo=True))
            s.add(Animal(numero="102", grupo_primario="01 - Alta", raca="Holandês", del_dias=80, ativo=True))
            s.commit()
        yield c

    main.app.dependency_overrides.clear()


class TestModeloExcel:
    def test_modelo_controle_leiteiro_animal(self, client):
        r = client.get("/producao/controle-leiteiro/modelo-excel", params={"modo": "animal"})
        assert r.status_code == 200
        assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
        wb = openpyxl.load_workbook(io.BytesIO(r.content))
        cabecalho = [c.value for c in wb.active[1]]
        assert "Número" in cabecalho

    def test_modelo_controle_leiteiro_lote(self, client):
        r = client.get("/producao/controle-leiteiro/modelo-excel", params={"modo": "lote"})
        assert r.status_code == 200
        wb = openpyxl.load_workbook(io.BytesIO(r.content))
        cabecalho = [c.value for c in wb.active[1]]
        assert "Lote" in cabecalho

    def test_modelo_qualidade_leite(self, client):
        r = client.get("/producao/qualidade-leite/modelo-excel")
        assert r.status_code == 200
        wb = openpyxl.load_workbook(io.BytesIO(r.content))
        cabecalho = [c.value for c in wb.active[1]]
        assert "CCS" in cabecalho


class TestImportarControleLeiteiroPorAnimal:
    def test_planilha_por_animal_cria_controle_com_del_e_raca(self, client):
        conteudo = _xlsx_bytes(
            ["Número", "Data", "Primeira ordenha (kg)", "Segunda ordenha (kg)", "Terceira ordenha (kg)", "Total"],
            [["101", "05/07/2026", "14,5", "13,0", "", "27,5"]],
        )
        r = client.post(
            "/producao/controle-leiteiro/importar",
            files={"file": ("controle.xlsx", conteudo, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert r.status_code == 200, r.text
        d = r.json()
        assert d["criados"] == 1 and d["erros"] == [] and d["modo"] == "animal"
        registro = next(c for c in client.get("/producao/controles").json()["controles"] if c["numero"] == "101")
        assert registro["producao_kg"] == 27.5
        assert registro["del"] == 50
        assert registro["raca"] == "Girolando"

    def test_linha_sem_numero_ou_ordenha_vira_erro_mas_nao_derruba_as_outras(self, client):
        conteudo = _xlsx_bytes(
            ["Número", "Data", "Primeira ordenha (kg)", "Segunda ordenha (kg)"],
            [["101", "05/07/2026", "14,5", "13,0"], ["", "05/07/2026", "10", "9"]],
        )
        r = client.post(
            "/producao/controle-leiteiro/importar",
            files={"file": ("controle.xlsx", conteudo, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert r.status_code == 200
        d = r.json()
        assert d["criados"] == 1
        assert len(d["erros"]) == 1


class TestPreVisualizarEConfirmarControleLeiteiro:
    """Fluxo novo: enviar a planilha só devolve as linhas normalizadas (sem
    gravar nada) para o usuário revisar/editar; a gravação de verdade só
    acontece em /confirmar, com as linhas (já editadas ou não)."""

    def test_pre_visualizar_nao_grava_nada(self, client):
        conteudo = _xlsx_bytes(
            ["Número", "Data", "Primeira ordenha (kg)", "Segunda ordenha (kg)"],
            [["101", "05/07/2026", "14,5", "13,0"]],
        )
        r = client.post(
            "/producao/controle-leiteiro/pre-visualizar",
            files={"file": ("controle.xlsx", conteudo, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert r.status_code == 200, r.text
        d = r.json()
        assert d["modo"] == "animal" and d["erros"] == []
        assert d["linhas"] == [{
            "numero_matriz": "101", "data_controle": "2026-07-05",
            "ordenha1_kg": 14.5, "ordenha2_kg": 13.0, "ordenha3_kg": None, "total_kg": None,
        }]
        assert client.get("/producao/controles").json()["controles"] == []

    def test_pre_visualizar_por_lote_ja_distribui_por_animal(self, client):
        conteudo = _xlsx_bytes(
            ["Lote", "Data", "Total"],
            [["01 - Alta", "05/07/2026", "36"]],
        )
        r = client.post(
            "/producao/controle-leiteiro/pre-visualizar",
            files={"file": ("controle.xlsx", conteudo, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert r.status_code == 200, r.text
        d = r.json()
        assert d["modo"] == "lote"
        assert {l["numero_matriz"]: l["total_kg"] for l in d["linhas"]} == {"101": 18.0, "102": 18.0}

    def test_confirmar_grava_as_linhas_editadas(self, client):
        conteudo = _xlsx_bytes(
            ["Número", "Data", "Primeira ordenha (kg)", "Segunda ordenha (kg)"],
            [["101", "05/07/2026", "14,5", "13,0"]],
        )
        preview = client.post(
            "/producao/controle-leiteiro/pre-visualizar",
            files={"file": ("controle.xlsx", conteudo, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        ).json()
        linhas = preview["linhas"]
        # Usuário corrige um valor digitado errado na planilha antes de confirmar.
        linhas[0]["ordenha1_kg"] = 15.0

        r = client.post("/producao/controle-leiteiro/confirmar", json={"linhas": linhas})
        assert r.status_code == 200, r.text
        assert r.json()["criados"] == 1

        registro = next(c for c in client.get("/producao/controles").json()["controles"] if c["numero"] == "101")
        assert registro["producao_kg"] == 28.0  # 15 + 13, valor corrigido

    def test_confirmar_agrupa_por_data_mesmo_com_datas_diferentes_na_planilha(self, client):
        linhas = [
            {"numero_matriz": "101", "data_controle": "2026-07-05", "ordenha1_kg": 14.0, "ordenha2_kg": None, "ordenha3_kg": None, "total_kg": None},
            {"numero_matriz": "102", "data_controle": "2026-07-06", "ordenha1_kg": 12.0, "ordenha2_kg": None, "ordenha3_kg": None, "total_kg": None},
        ]
        r = client.post("/producao/controle-leiteiro/confirmar", json={"linhas": linhas})
        assert r.status_code == 200, r.text
        assert r.json()["criados"] == 2
        controles = client.get("/producao/controles").json()["controles"]
        assert next(c for c in controles if c["numero"] == "101")["data"] == "2026-07-05"
        assert next(c for c in controles if c["numero"] == "102")["data"] == "2026-07-06"


class TestImportarControleLeiteiroPorLote:
    def test_planilha_por_lote_distribui_igualmente_entre_animais(self, client):
        conteudo = _xlsx_bytes(
            ["Lote", "Data", "Primeira ordenha (kg)", "Segunda ordenha (kg)", "Terceira ordenha (kg)", "Total"],
            [["01 - Alta", "05/07/2026", "20", "16", "", "36"]],
        )
        r = client.post(
            "/producao/controle-leiteiro/importar",
            files={"file": ("controle.xlsx", conteudo, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert r.status_code == 200, r.text
        d = r.json()
        assert d["criados"] == 2 and d["erros"] == [] and d["modo"] == "lote"
        controles = client.get("/producao/controles").json()["controles"]
        r101 = next(c for c in controles if c["numero"] == "101")
        r102 = next(c for c in controles if c["numero"] == "102")
        # 20/2=10, 16/2=8 -> total 18 por vaca; a soma bate com o total da planilha.
        assert r101["producao_kg"] == 18.0
        assert r102["producao_kg"] == 18.0
        assert r101["del"] == 50 and r102["del"] == 80  # DEL individual preservado

    def test_lote_por_so_o_codigo_tambem_resolve(self, client):
        conteudo = _xlsx_bytes(["Lote", "Data", "Primeira ordenha (kg)"], [["01", "05/07/2026", "10"]])
        r = client.post(
            "/producao/controle-leiteiro/importar",
            files={"file": ("controle.xlsx", conteudo, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert r.status_code == 200
        assert r.json()["criados"] == 2

    def test_lote_nao_cadastrado_vira_erro(self, client):
        conteudo = _xlsx_bytes(["Lote", "Data", "Primeira ordenha (kg)"], [["99 - Fantasma", "05/07/2026", "10"]])
        r = client.post(
            "/producao/controle-leiteiro/importar",
            files={"file": ("controle.xlsx", conteudo, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert r.status_code == 200
        d = r.json()
        assert d["criados"] == 0
        assert len(d["erros"]) == 1 and "não encontrado" in d["erros"][0]


class TestImportarControleLeiteiroSoTotal:
    def test_planilha_por_animal_so_com_total_grava_producao_sem_quebrar_ordenha(self, client):
        conteudo = _xlsx_bytes(["Número", "Data", "Total"], [["101", "05/07/2026", "27,5"]])
        r = client.post(
            "/producao/controle-leiteiro/importar",
            files={"file": ("controle.xlsx", conteudo, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert r.status_code == 200, r.text
        d = r.json()
        assert d["criados"] == 1 and d["erros"] == []
        registro = next(c for c in client.get("/producao/controles").json()["controles"] if c["numero"] == "101")
        assert registro["producao_kg"] == 27.5
        assert registro["ordenha1_kg"] is None
        assert registro["ordenha2_kg"] is None
        assert registro["ordenha3_kg"] is None

    def test_planilha_por_lote_so_com_total_distribui_sem_quebrar_ordenha(self, client):
        conteudo = _xlsx_bytes(["Lote", "Data", "Total"], [["01 - Alta", "05/07/2026", "36"]])
        r = client.post(
            "/producao/controle-leiteiro/importar",
            files={"file": ("controle.xlsx", conteudo, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert r.status_code == 200, r.text
        d = r.json()
        assert d["criados"] == 2 and d["erros"] == []
        controles = client.get("/producao/controles").json()["controles"]
        r101 = next(c for c in controles if c["numero"] == "101")
        r102 = next(c for c in controles if c["numero"] == "102")
        assert r101["producao_kg"] == 18.0 and r102["producao_kg"] == 18.0
        assert r101["ordenha1_kg"] is None and r102["ordenha1_kg"] is None

    def test_linha_sem_ordenha_nem_total_vira_erro(self, client):
        conteudo = _xlsx_bytes(["Número", "Data"], [["101", "05/07/2026"]])
        r = client.post(
            "/producao/controle-leiteiro/importar",
            files={"file": ("controle.xlsx", conteudo, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert r.status_code == 200
        d = r.json()
        assert d["criados"] == 0 and len(d["erros"]) == 1


class TestImportarControleLeiteiroValidacao:
    def test_planilha_sem_coluna_reconhecida_da_400(self, client):
        conteudo = _xlsx_bytes(["Foo", "Bar"], [["x", "y"]])
        r = client.post(
            "/producao/controle-leiteiro/importar",
            files={"file": ("controle.xlsx", conteudo, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert r.status_code == 400
        assert "identifiquei" in r.json()["detail"].lower()


class TestImportarQualidadeLeite:
    def test_planilha_com_apelidos_de_coluna_cria_registro(self, client):
        conteudo = _xlsx_bytes(
            ["Vaca", "Data", "Gordura", "Proteina", "CCS", "CBT"],
            [["101", "05/07/2026", "3,69", "3,42", "181", "11"]],
        )
        r = client.post(
            "/producao/qualidade-leite/importar",
            files={"file": ("qualidade.xlsx", conteudo, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert r.status_code == 200, r.text
        d = r.json()
        assert d["criados"] == 1 and d["erros"] == []
        registro = client.get("/producao/qualidade-leite").json()["registros"][0]
        assert registro["numero_matriz"] == "101"
        assert registro["gordura_pct"] == 3.69
        assert registro["ccs"] == 181

    def test_tanque_sem_numero_matriz(self, client):
        conteudo = _xlsx_bytes(["Data", "Gordura"], [["05/07/2026", "3,5"]])
        r = client.post(
            "/producao/qualidade-leite/importar",
            files={"file": ("qualidade.xlsx", conteudo, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert r.status_code == 200
        assert r.json()["criados"] == 1
        registro = client.get("/producao/qualidade-leite").json()["registros"][0]
        assert registro["numero_matriz"] is None

    def test_linha_sem_data_vira_erro(self, client):
        conteudo = _xlsx_bytes(["Data", "Gordura"], [["", "3,5"]])
        r = client.post(
            "/producao/qualidade-leite/importar",
            files={"file": ("qualidade.xlsx", conteudo, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert r.status_code == 200
        d = r.json()
        assert d["criados"] == 0 and len(d["erros"]) == 1


class TestModeloExcelPesagemCorporal:
    def test_modelo_pesagem_corporal(self, client):
        r = client.get("/producao/pesagens/modelo-excel")
        assert r.status_code == 200
        assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
        wb = openpyxl.load_workbook(io.BytesIO(r.content))
        cabecalho = [c.value for c in wb.active[1]]
        assert "Peso (kg)" in cabecalho


class TestImportarPesagemCorporal:
    def test_planilha_cria_pesagem_com_del_e_lote_do_cadastro(self, client):
        conteudo = _xlsx_bytes(["Número do animal", "Data", "Peso (kg)"], [["101", "05/07/2026", "420,5"]])
        r = client.post(
            "/producao/pesagens/importar",
            files={"file": ("pesagem.xlsx", conteudo, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert r.status_code == 200, r.text
        d = r.json()
        assert d["criados"] == 1 and d["erros"] == []
        registro = client.get("/producao/pesagens/relatorio", params={"numero_matriz": "101"}).json()["linhas"][0]
        assert registro["ultima_peso"] == 420.5
        assert registro["grupo_primario"] == "01 - Alta"

    def test_linha_sem_peso_vira_erro_mas_nao_derruba_as_outras(self, client):
        conteudo = _xlsx_bytes(
            ["Número do animal", "Data", "Peso (kg)"],
            [["101", "05/07/2026", "420"], ["102", "05/07/2026", ""]],
        )
        r = client.post(
            "/producao/pesagens/importar",
            files={"file": ("pesagem.xlsx", conteudo, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert r.status_code == 200
        d = r.json()
        assert d["criados"] == 1
        assert len(d["erros"]) == 1

    def test_planilha_vazia_retorna_zero_criados(self, client):
        conteudo = _xlsx_bytes(["Número do animal", "Data", "Peso (kg)"], [])
        r = client.post(
            "/producao/pesagens/importar",
            files={"file": ("pesagem.xlsx", conteudo, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert r.status_code == 200
        assert r.json()["criados"] == 0
